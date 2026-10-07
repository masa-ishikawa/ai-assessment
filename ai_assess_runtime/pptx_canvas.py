"""ReportLab互換座標を編集可能なPowerPoint図形へ変換する基盤。"""

import re
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader

from ai_assess_runtime.rendering_profile import (
    COMMON_CHROME_FONT_PT,
    MIN_EDITABLE_BODY_FONT_PT,
)

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt
except ImportError:
    Presentation = None  # type: ignore[assignment]

# 顧客向け成果物の編集可能テキストに適用する唯一の下限定数。
# 個別スライド側でこれより小さい値を指定しても、描画時と最終保存時の
# 二段階で12ptへ正規化する。
MIN_PPTX_FONT_SIZE = MIN_EDITABLE_BODY_FONT_PT
# 共通ヘッダー／フッターと出典注記に適用する明示的な例外下限。
# 図形名へ役割を記録し、保存時の全体監査でも該当図形だけを例外扱いする。
PPTX_CHROME_FONT_SIZE = COMMON_CHROME_FONT_PT
PPTX_SOURCE_NOTE_FONT_SIZE = 9.0
PPTX_HEADER_SHAPE_NAME = "AI_ASSESS_COMMON_HEADER"
PPTX_FOOTER_SHAPE_NAME = "AI_ASSESS_COMMON_FOOTER"
PPTX_SOURCE_NOTE_SHAPE_NAME = "AI_ASSESS_SOURCE_NOTE"

class PptCanvas:
    """ReportLab座標系をPowerPointの編集可能な図形へ写す最小キャンバス。

    既存の描画ヘルパーを共用し、座標・色・画像の定義を二重管理せずに、
    PPTXでは文字、線、矩形、円を個別に編集できるようにする。
    """

    POINT_TO_EMU = 12700
    # 生成する全ての編集可能テキストをMeiryo UIへ統一する。既存の英字向け指定も
    # PPTXでは同じフォントへ正規化し、スライド内のフォント混在を防ぐ。
    FONT_MAP = {
        "AssessmentJapanese": "Meiryo UI",
        "AssessmentJapaneseBold": "Meiryo UI",
        "Helvetica": "Meiryo UI",
        "Helvetica-Bold": "Meiryo UI",
    }

    def __init__(self, page_width: float, page_height: float) -> None:
        if Presentation is None:
            raise RuntimeError("PPTX出力には python-pptx が必要です。requirements.txt をインストールしてください。")
        self.page_width, self.page_height = page_width, page_height
        self.presentation = Presentation()
        self.presentation.slide_width = Emu(round(page_width * self.POINT_TO_EMU))
        self.presentation.slide_height = Emu(round(page_height * self.POINT_TO_EMU))
        self.slide = None
        self.fill_color = colors.black
        self.stroke_color = colors.black
        self.line_width = 1.0
        self.font_name = "AssessmentJapanese"
        self.font_size = MIN_PPTX_FONT_SIZE
        self._chrome_role_once: str | None = None
        self._transform = (0.0, 0.0, 1.0, 1.0)
        self._transform_stack: list[tuple[float, float, float, float]] = []
        self._new_slide()

    def _new_slide(self) -> None:
        self._chrome_role_once = None
        self.slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
        # 背景は編集対象の白い矩形ではなくスライド背景で管理する。これにより
        # 利用者がPPTXを編集する際、全ページに不要な全画面図形が残らない。
        self._set_slide_background(colors.white)

    def _set_slide_background(self, value: colors.Color) -> None:
        background_fill = self.slide.background.fill
        background_fill.solid()
        background_fill.fore_color.rgb = self._rgb(value)

    @staticmethod
    def _rgb(value: colors.Color) -> RGBColor:
        return RGBColor(round(value.red * 255), round(value.green * 255), round(value.blue * 255))

    def _x(self, value: float) -> Emu:
        tx, _ty, sx, _sy = self._transform
        return Emu(round((tx + value * sx) * self.POINT_TO_EMU))

    def _y(self, value: float) -> Emu:
        _tx, ty, _sx, sy = self._transform
        return Emu(round((self.page_height - (ty + value * sy)) * self.POINT_TO_EMU))

    def _height(self, value: float) -> Emu:
        return Emu(round(value * self._transform[3] * self.POINT_TO_EMU))

    def _width(self, value: float) -> Emu:
        return Emu(round(value * self._transform[2] * self.POINT_TO_EMU))

    def setTitle(self, _title: str) -> None:
        return None

    def setFillColor(self, value: colors.Color) -> None:
        self.fill_color = value

    def setStrokeColor(self, value: colors.Color) -> None:
        self.stroke_color = value

    def setLineWidth(self, value: float) -> None:
        self.line_width = value

    def setFont(self, name: str, size: float) -> None:
        self.font_name = name
        self.font_size = max(MIN_PPTX_FONT_SIZE, float(size))
        # 通常本文の指定が入った時点で、未使用のheader/footer例外は破棄する。
        self._chrome_role_once = None

    def setFooterFont(self, name: str, size: float) -> None:
        """次に描くテキストを、10ptを許す共通フッターとして明示する。"""
        self.font_name = name
        self.font_size = max(PPTX_CHROME_FONT_SIZE, float(size))
        self._chrome_role_once = "footer"

    def setHeaderFont(self, name: str, size: float) -> None:
        """次に描くテキストを、10ptを許す共通ヘッダーとして明示する。"""
        self.font_name = name
        self.font_size = max(PPTX_CHROME_FONT_SIZE, float(size))
        self._chrome_role_once = "header"

    def _shape_style(self, shape, stroke: int, fill: int) -> None:
        if fill:
            shape.fill.solid(); shape.fill.fore_color.rgb = self._rgb(self.fill_color)
        else:
            shape.fill.background()
        if stroke:
            shape.line.color.rgb = self._rgb(self.stroke_color)
            shape.line.width = Pt(self.line_width)
        else:
            shape.line.fill.background()

    def rect(self, x: float, y: float, width: float, height: float, stroke: int = 1, fill: int = 0) -> None:
        # ページ全面の塗りはスライド背景へ正規化する。ReportLab互換の呼び出しを
        # 保ったまま、PPTX内の不要な全画面図形を除去する。
        is_full_slide = (
            abs(x) < 0.01 and abs(y) < 0.01
            and abs(width - self.page_width) < 0.01 and abs(height - self.page_height) < 0.01
            and self._transform == (0.0, 0.0, 1.0, 1.0)
        )
        if fill and not stroke and is_full_slide:
            self._set_slide_background(self.fill_color)
            return
        shape = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, self._x(x), self._y(y + height), self._width(width), self._height(height))
        self._shape_style(shape, stroke, fill)

    def roundRect(self, x: float, y: float, width: float, height: float, _radius: float, stroke: int = 1, fill: int = 0) -> None:
        shape = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, self._x(x), self._y(y + height), self._width(width), self._height(height))
        self._shape_style(shape, stroke, fill)

    def circle(self, x: float, y: float, radius: float, stroke: int = 1, fill: int = 0) -> None:
        shape = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, self._x(x - radius), self._y(y + radius), self._width(radius * 2), self._height(radius * 2))
        self._shape_style(shape, stroke, fill)

    def ellipse(self, x1: float, y1: float, x2: float, y2: float, stroke: int = 1, fill: int = 0) -> None:
        shape = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, self._x(x1), self._y(y2), self._width(x2 - x1), self._height(y2 - y1))
        self._shape_style(shape, stroke, fill)

    def line(self, x1: float, y1: float, x2: float, y2: float) -> None:
        shape = self.slide.shapes.add_connector(1, self._x(x1), self._y(y1), self._x(x2), self._y(y2))
        shape.line.color.rgb = self._rgb(self.stroke_color); shape.line.width = Pt(self.line_width)

    def stringWidth(self, value: str, _font_name: str, font_size: float) -> float:
        # 文字量調整側も実描画と同じ下限（本文12pt、共通chrome 10pt）で測る。
        # 小さい指定値のまま幅を見積もると、保存時補正後にはみ出すため。
        minimum_size = (
            PPTX_CHROME_FONT_SIZE
            if getattr(self, "_chrome_role_once", None)
            else MIN_PPTX_FONT_SIZE
        )
        effective_font_size = max(minimum_size, float(font_size))
        return sum(
            effective_font_size * (1 if ord(char) < 128 else 1.85) * 0.52
            for char in str(value)
        )

    def _text(self, text: str, x: float, baseline_y: float, width: float, align=PP_ALIGN.LEFT) -> None:
        chrome_role = getattr(self, "_chrome_role_once", None)
        minimum_size = PPTX_CHROME_FONT_SIZE if chrome_role else MIN_PPTX_FONT_SIZE
        font_size = max(minimum_size, float(self.font_size))
        self._chrome_role_once = None
        top = baseline_y - font_size * 0.86
        height = max(font_size * 1.45, 14)
        box = self.slide.shapes.add_textbox(self._x(x), self._y(top + height), self._width(max(width, 2)), self._height(height))
        frame = box.text_frame; frame.clear(); frame.word_wrap = False; frame.margin_left = frame.margin_right = 0; frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        paragraph = frame.paragraphs[0]; paragraph.alignment = align
        run = paragraph.add_run(); run.text = str(text)
        run.font.name = self.FONT_MAP.get(self.font_name, self.font_name); run.font.size = Pt(font_size); run.font.bold = self.font_name.endswith("Bold"); run.font.color.rgb = self._rgb(self.fill_color)
        if chrome_role == "header":
            box.name = PPTX_HEADER_SHAPE_NAME
        elif chrome_role == "footer":
            box.name = PPTX_FOOTER_SHAPE_NAME

    def drawString(self, x: float, y: float, text: str) -> None:
        self._text(text, x, y, self.page_width - x)

    def drawRightString(self, x: float, y: float, text: str) -> None:
        width = self.stringWidth(text, self.font_name, self.font_size) + 4
        self._text(text, x - width, y, width, PP_ALIGN.RIGHT)

    def drawCentredString(self, x: float, y: float, text: str) -> None:
        width = self.stringWidth(text, self.font_name, self.font_size) + 4
        self._text(text, x - width / 2, y, width, PP_ALIGN.CENTER)

    def drawImage(self, path, x: float, y: float, width: float, height: float, **_kwargs) -> None:
        image_path = str(path.fileName) if isinstance(path, ImageReader) else str(path)
        self.slide.shapes.add_picture(image_path, self._x(x), self._y(y + height), self._width(width), self._height(height))

    def add_table(self, rows: list[list[str]], x: float, y: float, widths: list[float], heights: list[float],
                  header_color: colors.Color, body_colors: list[colors.Color],
                  font_size: float = MIN_PPTX_FONT_SIZE,
                  cell_margin_mm: float = 2.2,
                  cell_vertical_margin_mm: float = 0.7,
                  line_spacing_pt: float | None = None) -> None:
        effective_font_size = max(MIN_PPTX_FONT_SIZE, float(font_size))
        shape = self.slide.shapes.add_table(len(rows), len(widths), self._x(x), self._y(y + sum(heights)), self._width(sum(widths)), self._height(sum(heights)))
        table = shape.table
        for index, width in enumerate(widths): table.columns[index].width = self._width(width)
        for index, height in enumerate(heights): table.rows[index].height = self._height(height)
        for row_index, values in enumerate(rows):
            background = header_color if row_index == 0 else body_colors[min(row_index - 1, len(body_colors) - 1)]
            for column_index, value in enumerate(values):
                cell = table.cell(row_index, column_index); cell.fill.solid(); cell.fill.fore_color.rgb = self._rgb(background)
                cell.margin_left = cell.margin_right = Emu(round(cell_margin_mm * mm * self.POINT_TO_EMU))
                cell.margin_top = cell.margin_bottom = Emu(round(cell_vertical_margin_mm * mm * self.POINT_TO_EMU))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                frame = cell.text_frame; frame.clear(); frame.word_wrap = True
                paragraph = frame.paragraphs[0]
                paragraph.space_before = Pt(0); paragraph.space_after = Pt(0)
                if line_spacing_pt is not None:
                    paragraph.line_spacing = Pt(max(effective_font_size, float(line_spacing_pt)))
                run = paragraph.add_run(); run.text = _plain_text(value)
                run.font.name = "Meiryo UI"; run.font.size = Pt(effective_font_size); run.font.bold = row_index == 0
                run.font.color.rgb = self._rgb(colors.white if row_index == 0 else colors.HexColor("#1D252C"))

    def showPage(self) -> None:
        self._new_slide()

    def saveState(self) -> None:
        self._transform_stack.append(self._transform)

    def restoreState(self) -> None:
        self._transform = self._transform_stack.pop() if self._transform_stack else (0.0, 0.0, 1.0, 1.0)

    def translate(self, x: float, y: float) -> None:
        tx, ty, sx, sy = self._transform
        self._transform = (tx + x * sx, ty + y * sy, sx, sy)

    def scale(self, x: float, y: float) -> None:
        tx, ty, sx, sy = self._transform
        self._transform = (tx, ty, sx * x, sy * y)

    def save(self) -> None:
        return None

    def _enforce_minimum_font_size(self) -> None:
        """保存直前に本文12pt、共通chrome 10pt、出典注記9ptを保証する。

        通常の描画ヘルパーは ``setFont`` / ``draw_paragraph`` / ``add_table`` で
        下限を守る。一方、個別レイアウトが ``run.font.size`` を直接指定した場合や
        サイズをテーマ継承にした場合も漏らさないよう、全スライドを最後に監査する。
        例外は、共通描画関数が図形名に明示した役割だけを対象とする。
        """
        for slide in self.presentation.slides:
            shapes = list(slide.shapes)
            while shapes:
                shape = shapes.pop()
                child_shapes = getattr(shape, "shapes", None)
                if child_shapes is not None:
                    shapes.extend(list(child_shapes))

                frames = [shape.text_frame] if getattr(shape, "has_text_frame", False) else []
                if getattr(shape, "has_table", False):
                    frames.extend(
                        cell.text_frame
                        for row in shape.table.rows
                        for cell in row.cells
                    )
                is_common_chrome = getattr(shape, "name", "") in {
                    PPTX_HEADER_SHAPE_NAME,
                    PPTX_FOOTER_SHAPE_NAME,
                }
                is_source_note = getattr(shape, "name", "") == PPTX_SOURCE_NOTE_SHAPE_NAME
                minimum_size = (
                    PPTX_CHROME_FONT_SIZE if is_common_chrome
                    else PPTX_SOURCE_NOTE_FONT_SIZE if is_source_note
                    else MIN_PPTX_FONT_SIZE
                )
                for frame in frames:
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            if not run.text:
                                continue
                            # 保存時に入力テキストを書き換えない。表示領域が不足する
                            # 場合は、呼び出し側が折返し・行高・ページ分割で吸収する。
                            # 原文の省略記号も含め、文字列はそのまま編集可能に残す。
                            if run.font.size is None or run.font.size.pt < minimum_size:
                                run.font.size = Pt(minimum_size)

    def write(self, output_path: Path) -> None:
        self._enforce_minimum_font_size()
        self.presentation.save(str(output_path))


def _plain_text(value: str) -> str:
    return re.sub(r"<[^>]+>", "", str(value)).replace("&nbsp;", " ")
