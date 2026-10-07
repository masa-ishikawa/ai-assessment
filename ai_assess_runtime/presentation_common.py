"""PowerPoint描画層で共有する定数とプリミティブ。

このモジュールはスライド固有の構成を持たず、16:9座標系、文字組み、
共通ヘッダー／フッターの責務だけを担う。公開互換性のため、従来の
``ai_assess_runtime.presentation`` から各名前を再エクスポートする。
"""

from __future__ import annotations

import math
from pathlib import Path
import re

from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm

try:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt
except ImportError:
    MSO_ANCHOR = MSO_AUTO_SIZE = PP_ALIGN = Pt = None  # type: ignore[assignment]

from ai_assess_runtime.paths import PRIORITY_INSIGHT_IMAGE
from ai_assess_runtime.pptx_canvas import (
    MIN_PPTX_FONT_SIZE,
    PPTX_CHROME_FONT_SIZE,
    PPTX_SOURCE_NOTE_SHAPE_NAME,
    PptCanvas,
    _plain_text,
)


# Compatibility-only exports for callers that imported the former PDF font
# discovery constants. PPTX rendering does not inspect filesystem font paths.
FONT_PATH_CANDIDATES: tuple[Path, ...] = ()
FONT_BOLD_PATH_CANDIDATES: tuple[Path, ...] = ()

STANDARD_TITLE_X = 18 * mm
STANDARD_TITLE_Y_OFFSET = 20 * mm
STANDARD_TITLE_SIZE = 20
STANDARD_HEADER_HEIGHT = 14 * mm

PPTX_WIDESCREEN_WIDTH = 13.3333333333 * 72
PPTX_WIDESCREEN_HEIGHT = 7.5 * 72

WIDE_CONTENT_LEFT = 18 * mm
WIDE_CONTENT_RIGHT = 18 * mm
WIDE_EXPLANATORY_LEAD_TOP = 156 * mm
WIDE_TABLE_MIN_ROW_HEIGHT = 11 * mm

USE_CASES_PER_LIST_SLIDE = 15
# 15行一覧の表下端。動的機密ラベルを避けて上げた共通フッターと
# 9mm以上の間隔を保ち、最下行とCopyrightを会議室投影でも分離する。
USE_CASE_LIST_TABLE_BOTTOM = 20.5 * mm
# 見出しテキストボックスの下端と重ねず、一覧表が使ってよい上端。
# ページ分割判定と実際の描画で同じ領域を使う。
USE_CASE_LIST_TABLE_TOP = 30.0 * mm
USE_CASE_LIST_HEADER_HEIGHT = 9.0 * mm
USE_CASE_LIST_ROW_HEIGHT = 8.65 * mm
# No. / ユースケース / 主なAI技術。残りを「業務価値・処理」に充てる。
# 一覧表は列幅を判定用と描画用で共用し、余分な安全係数で文字列を削らない。
USE_CASE_LIST_COLUMN_WIDTHS = (10 * mm, 74 * mm, 65 * mm)
USE_CASE_LIST_CELL_MARGIN_MM = 1.0
# Meiryo UI / PowerPointの実描画は、単純な「全角1文字=font size」見積りより
# おおむね数%だけ狭い。境界付近の1行セルを不要に2行扱いしない許容率。
USE_CASE_LIST_WRAP_TOLERANCE = 1.08
# 外部呼び出しとの互換性のため名前は残す。本文を切る目的には使用しない。
USE_CASE_LIST_TEXT_SAFETY_FACTOR = 1.0
USE_CASE_LIST_BODY_COLOR = colors.HexColor("#F5F7F8")
USE_CASE_LIST_PRIORITY_ICON_SIZE = 5.5 * mm
USE_CASE_LIST_PRIORITY_ICON_GAP = 2.0 * mm
USE_CASE_LIST_PRIORITY_ICON_IMAGE = PRIORITY_INSIGHT_IMAGE

# 共通ヘッダー／フッターは固定10pt、本文は12pt以上とする。
PPTX_HEADER_FONT_SIZE = PPTX_CHROME_FONT_SIZE
PPTX_FOOTER_FONT_SIZE = PPTX_CHROME_FONT_SIZE
# PowerPointがテナントの機密ラベルをスライド下端へ動的に描画しても、
# Copyrightとページ番号が重ならない10ptフッターの従来ベースライン。
PPTX_FOOTER_BASELINE_Y = 9.0 * mm

BUSINESS_VALUE_CARD_ACCENT_HEX = ("#C74634", "#367A9B", "#467653")

ASSESSMENT_SUBSECTION_LABELS = {
    "overview": "1.2 アセスメント概要",
    "industry": "2.1 業界・サービスの変化",
    "value_model": "2.2 導入先・提供者の価値モデル",
    "plan_alignment": "2.3 中期経営計画とAI実装の整合性",
    "decision": "2.4 AI活用の価値仮説・判断領域",
    "current_state": "3.1 価値連鎖・利用者・業務フロー",
    "data_issues": "3.2 データ資産と構造的課題",
    "evidence_target": "3.3 根拠台帳・将来業務フロー",
    "prioritization": "4.1 AIユースケース候補の整理・評価",
    "poc_selection": "4.2 優先PoC 3テーマの選定",
    "delivery": "4.3 PoC開始条件・実行計画・推進体制",
    "use_cases": "4.1 AIユースケース候補の整理・評価",
    "poc_design": "5.1 優先PoCの実現設計",
    "technical": "5.4 OCI技術提案・概算費用・支援",
}


def assessment_subsection_label(key: str) -> str:
    """目次と共通化した小項目ラベルを返す。"""
    return ASSESSMENT_SUBSECTION_LABELS[key]


def poc_design_subsection_label(priority: int) -> str:
    """選定済みPoCの詳細頁を、第5章の個別小項目として表示する。"""
    return f"5.{max(1, priority)} 優先PoC {max(1, priority)}の実現設計"


def register_japanese_font() -> None:
    """旧呼び出し元との互換性を保つPPTX専用のno-op。

    PPTXのフォントは :class:`PptCanvas` がMeiryo UIへ正規化するため、
    ReportLabのPDFフォント登録は不要である。
    """
    return None


def draw_footer(deck: PptCanvas, page_width: float, page_number: int) -> None:
    """共通フッターを編集可能なPPTXテキストとして描画する。"""
    # 座標は固定し、動的機密ラベル用の下端安全域を全ページで保つ。
    deck.setFillColor(colors.HexColor("#777777"))
    deck.setFooterFont("AssessmentJapanese", PPTX_FOOTER_FONT_SIZE)
    deck.drawString(
        18 * mm, PPTX_FOOTER_BASELINE_Y,
        "Copyright © 2026, Oracle and/or its affiliates",
    )
    deck.setFooterFont("AssessmentJapanese", PPTX_FOOTER_FONT_SIZE)
    deck.drawRightString(page_width - 18 * mm, PPTX_FOOTER_BASELINE_Y, str(page_number))


def estimate_pptx_paragraph_height(text: str, style: ParagraphStyle, width: float, *,
                                   minimum_font_size: float | None = None) -> float:
    """``draw_paragraph`` と同じ折返し規則で、必要な本文高さを返す。"""
    plain = re.sub(r"<br\s*/?>", "\n", str(text), flags=re.IGNORECASE)
    plain = _plain_text(plain)
    font_floor = MIN_PPTX_FONT_SIZE if minimum_font_size is None else float(minimum_font_size)
    font_size = max(font_floor, float(style.fontSize))
    leading = max(font_size * 1.22, float(style.leading or font_size * 1.25))
    usable_chars = max(8, int(width / (font_size * 0.92)))
    lines = sum(max(1, (len(line) + usable_chars - 1) // usable_chars) for line in plain.splitlines() or [""])
    return max(leading, lines * leading) + 1.5


def draw_paragraph(deck: PptCanvas, text: str, style: ParagraphStyle,
                   x: float, y: float, width: float, *,
                   minimum_font_size: float | None = None,
                   shape_name: str | None = None) -> float:
    """PPTXの編集可能なテキストボックスへ段落を描画する。"""
    plain = re.sub(r"<br\s*/?>", "\n", str(text), flags=re.IGNORECASE)
    plain = _plain_text(plain)
    font_floor = MIN_PPTX_FONT_SIZE if minimum_font_size is None else float(minimum_font_size)
    font_size = max(font_floor, float(style.fontSize))
    height = estimate_pptx_paragraph_height(
        plain, style, width, minimum_font_size=minimum_font_size,
    )
    box = deck.slide.shapes.add_textbox(deck._x(x), deck._y(y), deck._width(width), deck._height(height + 2))
    if shape_name:
        box.name = shape_name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    # 文字を切らず、必要な高さへ伸ばす。固定のカード境界を越える場合でも、
    # 原文の消失より全文の可視性を優先する。
    frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = plain
    run.font.name = PptCanvas.FONT_MAP.get(style.fontName, style.fontName)
    run.font.size = Pt(font_size)
    run.font.bold = style.fontName.endswith("Bold")
    run.font.color.rgb = deck._rgb(style.textColor or deck.fill_color)
    return y - height


def draw_number_badge(deck: PptCanvas, x: float, y: float, radius: float,
                      number: object, color: colors.Color,
                      font_size: float = MIN_PPTX_FONT_SIZE) -> None:
    """PPTXでも数字を円の中心へ収める、編集可能な丸数字を描画する。"""
    font_name = "AssessmentJapaneseBold"
    deck.setFillColor(color)
    deck.circle(x, y, radius, stroke=0, fill=1)
    box = deck.slide.shapes.add_textbox(
        deck._x(x - radius), deck._y(y + radius),
        deck._width(radius * 2), deck._height(radius * 2),
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(number)
    run.font.name = PptCanvas.FONT_MAP[font_name]
    run.font.size = Pt(max(MIN_PPTX_FONT_SIZE, font_size))
    run.font.bold = True
    run.font.color.rgb = deck._rgb(colors.white)


def wide_content_width(page_width: float) -> float:
    """16:9 PPTXで本文・表に使う横幅を一元化する。"""
    return page_width - WIDE_CONTENT_LEFT - WIDE_CONTENT_RIGHT


def estimate_wide_table_lines(value: object, width: float,
                              *, font_size: float = MIN_PPTX_FONT_SIZE,
                              cell_margin_mm: float = 2.2) -> int:
    """PowerPoint表セルの余白を考慮して、12pt本文に必要な行数を見積もる。"""
    text = _plain_text(str(value or ""))
    usable_width = max(20.0, width - 2 * max(0.0, cell_margin_mm) * mm)
    character_units = lambda line: sum(1.0 if ord(char) >= 128 else 0.56 for char in line)
    units_per_line = max(4.0, usable_width / max(font_size, MIN_PPTX_FONT_SIZE))
    return max(1, sum(max(1, math.ceil(character_units(line) / units_per_line)) for line in text.splitlines() or [""]))


def wide_table_row_height(row: list[object], widths: list[float],
                          *, font_size: float = MIN_PPTX_FONT_SIZE,
                          minimum: float = WIDE_TABLE_MIN_ROW_HEIGHT,
                          cell_margin_mm: float = 2.2) -> float:
    """内容に応じた行高を返し、12ptテキストがセルから切れないようにする。"""
    lines = max(
        estimate_wide_table_lines(
            value, widths[index], font_size=font_size, cell_margin_mm=cell_margin_mm,
        )
        for index, value in enumerate(row)
    ) if row else 1
    leading = max(font_size * 1.30, 13.0)
    return max(minimum, lines * leading + 1.8 * mm)


def basis_label(basis: str) -> str:
    return {"顧客入力": "入力", "公開資料": "公開資料", "分析仮説": "仮説", "要確認": "要確認"}.get(
        basis, basis or "仮説"
    )


def draw_story_source_note(deck: PptCanvas, page_width: float,
                           front_matter: dict, page_number: int,
                           source_urls: list[str] | None = None) -> None:
    """業界動向の参照URLを小さな出典欄で示し、共通フッターも描画する。"""
    urls = [str(url).strip() for url in (source_urls or []) if str(url).strip()][:2]
    if urls:
        source_text = "参考：" + "   ".join(
            f"※{index} {url}" for index, url in enumerate(urls, 1)
        )
        draw_paragraph(
            deck, source_text,
            ParagraphStyle(
                "industry_source_urls", fontName="AssessmentJapanese", fontSize=8.0,
                leading=9.4, textColor=colors.HexColor("#4C5961"), wordWrap="CJK",
            ),
            WIDE_CONTENT_LEFT, 16.2 * mm, page_width - WIDE_CONTENT_LEFT - WIDE_CONTENT_RIGHT,
        )
    draw_footer(deck, page_width, page_number)


def draw_pptx_base_header(deck: PptCanvas, page_width: float, page_height: float,
                          title: str, section: str) -> None:
    """標準PPTXの本文スライドで共有するヘッダー／見出しレイアウト。"""
    dark, green = colors.HexColor("#1D252C"), colors.HexColor("#467653")
    deck.setFillColor(colors.white)
    deck.rect(0, 0, page_width, page_height, stroke=0, fill=1)
    deck.setFillColor(green)
    deck.rect(0, page_height - STANDARD_HEADER_HEIGHT, page_width, STANDARD_HEADER_HEIGHT, stroke=0, fill=1)
    deck.setFillColor(colors.white)
    deck.setHeaderFont("AssessmentJapanese", PPTX_HEADER_FONT_SIZE)
    deck.drawString(WIDE_CONTENT_LEFT, page_height - 8.8 * mm, "OCI AI USE CASE ASSESSMENT")
    title_style = ParagraphStyle(
        "wide_story_title", fontName="AssessmentJapaneseBold", fontSize=20.0,
        leading=24.0, textColor=dark, wordWrap="CJK",
    )
    draw_paragraph(deck, title, title_style, WIDE_CONTENT_LEFT, page_height - 20 * mm,
                   wide_content_width(page_width))


def _executive_evidence_header(deck: PptCanvas, page_width: float, page_height: float,
                               title: str, section_label: str) -> None:
    """全説明ページで共用するPPTXヘッダーを描画する。"""
    draw_pptx_base_header(deck, page_width, page_height, title, section_label)


__all__ = [
    "FONT_PATH_CANDIDATES", "FONT_BOLD_PATH_CANDIDATES",
    "STANDARD_TITLE_X", "STANDARD_TITLE_Y_OFFSET", "STANDARD_TITLE_SIZE", "STANDARD_HEADER_HEIGHT",
    "PPTX_WIDESCREEN_WIDTH", "PPTX_WIDESCREEN_HEIGHT",
    "WIDE_CONTENT_LEFT", "WIDE_CONTENT_RIGHT", "WIDE_EXPLANATORY_LEAD_TOP", "WIDE_TABLE_MIN_ROW_HEIGHT",
    "USE_CASES_PER_LIST_SLIDE", "USE_CASE_LIST_TABLE_BOTTOM", "USE_CASE_LIST_TABLE_TOP",
    "USE_CASE_LIST_HEADER_HEIGHT",
    "USE_CASE_LIST_ROW_HEIGHT", "USE_CASE_LIST_COLUMN_WIDTHS", "USE_CASE_LIST_CELL_MARGIN_MM",
    "USE_CASE_LIST_WRAP_TOLERANCE",
    "USE_CASE_LIST_TEXT_SAFETY_FACTOR",
    "USE_CASE_LIST_BODY_COLOR", "USE_CASE_LIST_PRIORITY_ICON_SIZE", "USE_CASE_LIST_PRIORITY_ICON_GAP",
    "USE_CASE_LIST_PRIORITY_ICON_IMAGE", "PPTX_HEADER_FONT_SIZE", "PPTX_FOOTER_FONT_SIZE",
    "PPTX_FOOTER_BASELINE_Y",
    "BUSINESS_VALUE_CARD_ACCENT_HEX",
    "ASSESSMENT_SUBSECTION_LABELS", "assessment_subsection_label", "poc_design_subsection_label",
    "register_japanese_font", "draw_footer", "draw_paragraph", "estimate_pptx_paragraph_height", "draw_number_badge",
    "wide_content_width", "estimate_wide_table_lines", "wide_table_row_height", "basis_label",
    "draw_story_source_note", "draw_pptx_base_header", "_executive_evidence_header",
]
