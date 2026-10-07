import copy
import io
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm

import generate_assessment as assessment_generator
from test import create_isv_test_input


TEST_COST_ESTIMATE = {
    "currency": "JPY",
    "lines": [{"name": "test", "monthly_jpy": 1, "assumption": "test", "source": "test"}],
    "total_monthly_jpy": 1,
    "notes": [],
    "api_priced_count": 0,
}


def v4_quantitative_item(index: int, theme: str, metric: str, *, use_case_id: str | None = None,
                         outcome: str | None = None, kpi: str | None = None,
                         display_kpi_layer: str = "product_kpi",
                         product_metric: str | None = None,
                         product_kpi_name: str | None = None) -> dict:
    """Build one generic schema-v4 P1-P3 measurement item for unit tests."""
    priority = f"P{index}"
    use_case_id = use_case_id or f"UC{index:02d}"
    outcome = outcome or f"顧客業務成果{index}"
    kpi = kpi or f"製品KPI{index}"
    product_owner = {"scope": "provider", "role": "製品責任者", "status": "confirmed"}
    customer_owner = {"scope": "customer", "role": "業務責任者", "status": "confirm"}
    provider_owner = {"scope": "provider", "role": "事業責任者", "status": "confirm"}
    item = {
        "priority": priority, "use_case_id": use_case_id, "theme": theme,
        "business_outcome": outcome, "headline_metric": metric, "kpi": kpi,
        "baseline_definition": "対象機能の直近12週間実績値",
        "comparison_condition": "代表対象で12週間、現行機能とAI機能を比較する。",
        "detail": "製品KPIを直接測定し、顧客業務成果と提供者事業成果への寄与を確認する。",
        "formula": f"製品KPI改善率を{metric}として算定する",
        "assumption": "代表範囲と既存ログを用いて直接測定できるため設定する。",
        "target_rationale": "代表範囲と既存ログを用いて直接測定できるため設定する。",
        "display_kpi_layer": "product_kpi",
        "beneficiary": "shared", "metric_owner": product_owner,
        "attribution_level": "direct",
        "causal_link": "製品KPIの改善が顧客の業務成果へ寄与し、利用定着を通じて提供者事業成果へ接続する。",
        "product_kpi": {
            "kpi": kpi, "target": metric,
            "baseline_definition": "対象機能の直近12週間実績値",
            "comparison_condition": "代表対象で12週間、現行機能とAI機能を比較する。",
            "formula": f"製品KPI改善率を{metric}として算定する",
            "metric_owner": product_owner, "attribution_level": "direct",
            "causal_link": "AI機能の利用有無が製品KPIを直接変える。",
        },
        "customer_outcome_kpi": {
            "kpi": outcome, "target": "PoCで確定",
            "baseline_definition": "対象業務の現状実績をPoC開始前に確定する。",
            "comparison_condition": "製品KPI達成時の対象業務を現行運用と比較する。",
            "formula": "顧客実績値を用いてPoCで確定する。",
            "metric_owner": customer_owner, "attribution_level": "contributory",
            "causal_link": "製品KPI改善が顧客業務成果へ寄与する。",
        },
        "provider_business_kpi": {
            "kpi": "利用定着・継続利用", "target": "PoCで確定",
            "baseline_definition": "提供者側の現状実績をPoC開始前に確定する。",
            "comparison_condition": "PoC利用群と現行提供条件を比較する。",
            "formula": "提供者実績値を用いてPoCで確定する。",
            "metric_owner": provider_owner, "attribution_level": "enabling",
            "causal_link": "顧客成果と利用定着を提供者事業成果へ接続する。",
        },
        "poc_gate": {
            "measurement_period": "12週間", "go_condition": f"同一条件で{metric}を達成する。",
            "stop_condition": "品質・安全性または運用負荷が許容条件を外れた場合は停止する。",
            "decision_owner": "共同判定会議",
        },
    }
    if display_kpi_layer in {"customer_outcome_kpi", "provider_business_kpi"}:
        product_metric = product_metric or "65〜80%"
        product_kpi_name = product_kpi_name or f"製品先行KPI{index}"
        item["display_kpi_layer"] = display_kpi_layer
        item["product_kpi"].update({
            "kpi": product_kpi_name,
            "target": product_metric,
            "formula": f"{product_kpi_name}を同一条件で比較し、{product_metric}を判定する",
        })
        selected = item[display_kpi_layer]
        selected.update({
            "kpi": kpi,
            "target": metric,
            "formula": f"{kpi}を同一条件で比較し、{metric}を判定する",
        })
        if display_kpi_layer == "customer_outcome_kpi":
            item.update({
                "beneficiary": "shared",
                "metric_owner": customer_owner,
                "attribution_level": "contributory",
            })
        else:
            item.update({
                "beneficiary": "provider",
                "metric_owner": provider_owner,
                "attribution_level": "enabling",
            })
        item["formula"] = f"{kpi}を同一条件で比較し、{metric}を判定する"
        item["poc_gate"].update({
            "business_go_condition": f"事業KPIで{metric}を達成する。",
            "leading_kpi_condition": f"製品先行KPIで{product_metric}を達成する。",
            "go_condition": f"事業効果{metric}、製品先行KPI{product_metric}をともに達成する。",
        })
    return item


def v4_hypothesis_item(index: int, theme: str, metric: str, *, use_case_id: str | None = None,
                       outcome: str | None = None, kpi: str | None = None,
                       display_kpi_layer: str = "product_kpi",
                       product_metric: str | None = None,
                       product_kpi_name: str | None = None) -> dict:
    item = v4_quantitative_item(
        index, theme, metric, use_case_id=use_case_id, outcome=outcome, kpi=kpi,
        display_kpi_layer=display_kpi_layer, product_metric=product_metric,
        product_kpi_name=product_kpi_name,
    )
    item.pop("target_rationale", None)
    return item


class ResearchLoopTests(unittest.TestCase):
    def test_isv_template_extracts_required_questions_from_cde_columns(self) -> None:
        extracted = assessment_generator.extract_isv_assessment_input(
            assessment_generator.INPUT_DIR / "ISV_AI_Use_Case_Assessment_Input_Template.xlsx"
        )
        self.assertEqual("入力フォーム", extracted["sheet_name"])
        self.assertEqual(107, len(extracted["answers"]))
        self.assertEqual(24, len(extracted["required_answers"]))
        self.assertGreater(len(extracted["required_answers"]), 5)
        first = extracted["required_answers"][0]
        self.assertTrue(first["required"])
        self.assertTrue(first["question"])
        self.assertIn(first["question"], extracted["missing_required"])

    def test_archived_input_filename_is_not_resolved_implicitly(self) -> None:
        resolved = assessment_generator.resolve_input_file(Path("service_input_wms.xlsx"))
        self.assertEqual(Path("service_input_wms.xlsx"), resolved)

    def test_current_isv_template_resolves_from_input_directory(self) -> None:
        resolved = assessment_generator.resolve_input_file(Path(assessment_generator.CURRENT_INPUT_TEMPLATE))
        self.assertEqual(assessment_generator.INPUT_DIR / assessment_generator.CURRENT_INPUT_TEMPLATE, resolved)

    def test_isv_context_normalizes_assessment_fields_and_excludes_personal_data(self) -> None:
        extracted = assessment_generator.extract_isv_assessment_input(
            Path(__file__).parent / "fixtures" / "ISV_AI_Use_Case_Assessment_Input_ekusu.xlsx"
        )
        context = assessment_generator.normalize_isv_assessment_context(extracted)
        source_text = assessment_generator.build_isv_source_text(context)

        self.assertEqual("株式会社エクス", context["company"]["name"])
        self.assertEqual("Factory-ONE 電脳工場", context["service"]["name"])
        self.assertEqual("需要予測・在庫最適化", context["priority_use_cases"][0]["name"])
        self.assertIn("顧客指定の優先ユースケース", source_text)
        self.assertIn("会社URL: https://www.xeex.co.jp/", source_text)
        self.assertIn("提供形態: パッケージ;クラウドサービス", source_text)
        self.assertGreater(len(context["additional_answers"]), 5)
        self.assertNotIn("山田", source_text)
        self.assertNotIn("jiro@yamada.com", source_text)
        self.assertNotIn("F107", str(context))
        prompt = assessment_generator.build_prompt(source_text)
        self.assertIn("優先度「高」のテーマはpoc_recommendationsの3件にも必ず含める", prompt)
        self.assertIn("提案と共同検討の余地が伝わる穏やかな文体", prompt)
        self.assertIn("〜を参考情報として扱います", prompt)

    def test_isv_optional_context_excludes_inputter_role_and_consent(self) -> None:
        extracted = {
            "sheet_name": "入力フォーム",
            "answers": [
                {"question_id": "F101", "question": "会社名", "answer": "テスト株式会社", "required": True},
                {"question_id": "F109", "question": "入力者所属/役職", "answer": "AI推進部 部長", "required": False},
                {"question_id": "F406", "question": "主な競合製品・競合企業", "answer": "顧客入力の競合A", "required": False},
                {"question_id": "F1010", "question": "入力内容の共有可否", "answer": "共有可", "required": True},
            ],
        }
        context = assessment_generator.normalize_isv_assessment_context(extracted)
        source_text = assessment_generator.build_isv_source_text(context)
        self.assertIn("主な競合製品・競合企業: 顧客入力の競合A", source_text)
        self.assertNotIn("AI推進部 部長", source_text)
        self.assertNotIn("共有可", source_text)

    def test_test_input_script_fills_all_required_template_answers(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "test_input.xlsx"
            answers = create_isv_test_input.build_required_answers("テスト株式会社", "テスト製品", [])
            written = create_isv_test_input.fill_template(
                assessment_generator.INPUT_DIR / assessment_generator.CURRENT_INPUT_TEMPLATE,
                output_path, answers,
            )
            extracted = assessment_generator.extract_isv_assessment_input(output_path)
            with zipfile.ZipFile(assessment_generator.INPUT_DIR / assessment_generator.CURRENT_INPUT_TEMPLATE) as template_zip, \
                 zipfile.ZipFile(output_path) as output_zip:
                self.assertEqual(template_zip.read("xl/styles.xml"), output_zip.read("xl/styles.xml"))
                self.assertEqual(template_zip.read("xl/workbook.xml"), output_zip.read("xl/workbook.xml"))
        self.assertEqual(24, written)
        self.assertEqual([], extracted["missing_required"])

    def test_midterm_landing_page_prefers_same_site_detail_pdf(self) -> None:
        html = '''<a href="/ir/managementplan/MidtermManagementPlan202504.pdf">中期経営計画（2025年4月～2028年3月）</a>'''
        self.assertEqual(
            "https://www.nsw.co.jp/ir/managementplan/MidtermManagementPlan202504.pdf",
            assessment_generator.linked_midterm_pdf_url(html, "https://www.nsw.co.jp/ir/managementplan.html"),
        )

    def test_midterm_landing_page_prefers_latest_main_plan_over_qa(self) -> None:
        html = '''
        <a href="/ir/management/vision2031-qa-20240913.pdf">中期経営計画 Q&amp;A</a>
        <a href="/ir/management/vision2031-20260313.pdf">事業計画及び成長可能性</a>
        <a href="/ir/management/vision2031-20260612.pdf">事業計画及び成長可能性（最新版）</a>
        '''
        self.assertEqual(
            "https://corp.example.jp/ir/management/vision2031-20260612.pdf",
            assessment_generator.linked_midterm_pdf_url(
                html, "https://corp.example.jp/ir/management/growth-potential.php",
            ),
        )
        self.assertEqual(
            20260612,
            assessment_generator._document_date_rank("vision2031-20260612.pdf"),
        )

    def test_pdf_extraction_uses_the_more_readable_engine(self) -> None:
        response = type("Response", (), {
            "headers": {"content-type": "application/pdf"},
            "content": b"%PDF-fake",
            "text": "",
        })()
        pypdf_reader = type("Reader", (), {
            "pages": [type("Page", (), {"extract_text": lambda self: "Ϗδϣϯ Ϗδϣϯ 2026"})()],
        })()
        pymupdf_page = type("Page", (), {
            "get_text": lambda self: "株式会社サンプル 第3次中期経営計画 売上高と営業利益の目標",
        })()
        with patch.object(assessment_generator, "PdfReader", return_value=pypdf_reader), patch.object(
            assessment_generator.pymupdf, "open", return_value=[pymupdf_page],
        ):
            extracted = assessment_generator._extract_web_document_text(response)
        self.assertIn("中期経営計画", extracted)
        self.assertIn("株式会社サンプル", extracted)

    def test_web_document_title_prefers_original_ir_capitalization(self) -> None:
        response = type("Response", (), {
            "headers": {"content-type": "text/html; charset=utf-8"},
            "content": b"",
            "text": (
                "<html><head><title>  中期経営計画│IR情報│"
                "ロジザード株式会社  </title></head><body></body></html>"
            ),
        })()

        self.assertEqual(
            "中期経営計画│IR情報│ロジザード株式会社",
            assessment_generator._extract_web_document_title(response),
        )

    def test_management_targets_are_extracted_deterministically_from_the_primary_plan(self) -> None:
        plan = {
            "status": "found", "title": "中期経営計画（2025-2027）",
            "url": "https://example.com/plan",
            "excerpt": (
                "重要経営指標 2028年 売上高 120億円 2025年比+20%。"
                "2028年 営業利益 18億円 2025年比+25%。"
                "2028年 月次経常収益 2.5億円 2025年比+30%。"
            ),
        }
        with patch.object(assessment_generator, "_response_json", side_effect=AssertionError("LLMを呼ばない")):
            result = assessment_generator.extract_verified_management_targets_from_midterm_plan(plan)
        self.assertEqual(["重要経営指標 売上高", "営業利益", "月次経常収益"], [item["label"] for item in result])
        self.assertTrue(all(item["claim_status"] == "verified_external" for item in result))
        self.assertTrue(all(assessment_generator._management_target_is_verifiable(item, {"midterm_plan": plan}) for item in result))

    def test_assessment_json_contract_requires_reviewable_core_fields(self) -> None:
        valid = {
            "format": assessment_generator.ASSESSMENT_JSON_FORMAT,
            "assessment": {"company_name": "Example株式会社", "service_name": "Example Service",
                           "use_cases": [{"title": "u"}], "poc_recommendations": [{"theme": "p"}]},
        }
        self.assertEqual([], assessment_generator.validate_assessment_payload(valid))
        valid["assessment"]["service_name"] = ""
        self.assertIn("assessment.service_nameがありません。", assessment_generator.validate_assessment_payload(valid))

    def test_midterm_plan_rejects_unrelated_company_document(self) -> None:
        other_company = {
            "title": "PDF 中期経営計画説明資料",
            "url": "https://columbiaworks.example/plan.pdf",
            "snippet": "コロンビア・ワークス株式会社の中期経営計画",
        }

        def fake_search(_query: str, limit: int = 8) -> tuple[list[dict[str, str]], bool]:
            self.assertEqual(12, limit)
            return [other_company], False

        def fake_fetch(_item: dict[str, str]) -> str:
            return "コロンビア・ワークス株式会社 中期経営計画 2025年から2027年の目標"

        plan = assessment_generator.collect_midterm_plan(
            "株式会社フレームワークス", search_fn=fake_search, fetch_fn=fake_fetch,
        )

        self.assertEqual("not_found", plan["status"])
        self.assertNotIn("コロンビア", plan.get("excerpt", ""))
        self.assertEqual(
            [{"url": other_company["url"], "reason": "対象企業名を資料内で確認できないため除外しました。"}],
            plan["excluded_candidates"],
        )

    def test_midterm_plan_does_not_mark_blank_extraction_as_found(self) -> None:
        candidate = {
            "title": "Example株式会社 中期経営計画",
            "url": "https://example.co.jp/ir/management-plan.pdf",
            "snippet": "Example株式会社の中期経営計画",
        }
        plan = assessment_generator.collect_midterm_plan(
            "Example株式会社",
            search_fn=lambda _query, limit=12: ([candidate], False),
            fetch_fn=lambda _item: " ",
        )
        self.assertEqual("unavailable", plan["status"])
        self.assertIn("本文", plan["reason"])

    def test_midterm_plan_rejects_company_name_prefix_match(self) -> None:
        other_company = {
            "title": "株式会社エクスモーション 中期経営計画 — Edinet Db",
            "url": "https://disclosure2.edinet-fsa.go.jp/example.pdf",
            "snippet": "株式会社エクスモーションの中期経営計画",
        }
        plan = assessment_generator.collect_midterm_plan(
            "株式会社エクス",
            search_fn=lambda _query, limit=12: ([other_company], False),
            fetch_fn=lambda _item: "株式会社エクスモーション 中期経営計画 2025年から2027年の目標",
        )

        self.assertFalse(assessment_generator.document_matches_target_company(
            "株式会社エクス", other_company["title"], other_company["snippet"],
        ))
        self.assertEqual("not_found", plan["status"])
        self.assertEqual(other_company["url"], plan["excluded_candidates"][0]["url"])

    def test_midterm_plan_accepts_target_company_without_corporate_designator(self) -> None:
        target_company = {
            "title": "フレームワークス 中期経営計画",
            "url": "https://frame-wx.example/plan.pdf",
            "snippet": "物流事業の中期経営計画",
        }

        plan = assessment_generator.collect_midterm_plan(
            "株式会社フレームワークス",
            search_fn=lambda _query, limit=8: ([target_company], False),
            fetch_fn=lambda _item: "フレームワークスの中期経営計画。物流基盤の強化を進めます。",
        )

        self.assertEqual("found", plan["status"])
        self.assertEqual(target_company["url"], plan["url"])

    def test_midterm_plan_normalizes_ir_information_title(self) -> None:
        source = {
            "title": "中期経営計画│Ir情報│ロジザード株式会社",
            "url": "https://www.logizard.co.jp/ir/management-plan/",
            "snippet": "ロジザード株式会社の中期経営計画",
        }

        plan = assessment_generator.collect_midterm_plan(
            "ロジザード株式会社",
            search_fn=lambda _query, limit=12: ([source], False),
            fetch_fn=lambda _item: (
                "ロジザード株式会社 中期経営計画 2026年から2028年 "
                "売上高31億円、営業利益5億円を目標とする。"
            ),
        )

        self.assertEqual("found", plan["status"])
        self.assertEqual("中期経営計画│IR情報│ロジザード株式会社", plan["title"])
        self.assertNotIn("Ir情報", plan["title"])

    def test_midterm_plan_rejects_target_company_service_page(self) -> None:
        service_page = {
            "title": "コンサルティング - 株式会社フレームワークス",
            "url": "https://www.frame-wx.example/consulting-service/",
            "snippet": "中期経営計画の策定を支援します。",
        }

        plan = assessment_generator.collect_midterm_plan(
            "株式会社フレームワークス",
            search_fn=lambda _query, limit=8: ([service_page], False),
            fetch_fn=lambda _item: "株式会社フレームワークスは中期経営計画の策定を支援します。",
        )

        self.assertEqual("not_found", plan["status"])

    def test_midterm_plan_accepts_official_ir_document_with_generic_title(self) -> None:
        ir_document = {
            "title": "2025年3月期 決算説明資料",
            "url": "https://www.example.co.jp/ir/plan/2025.pdf",
            "snippet": "株式会社フレームワークスのIR資料",
        }

        plan = assessment_generator.collect_midterm_plan(
            "株式会社フレームワークス",
            search_fn=lambda _query, limit=12: ([ir_document], False),
            fetch_fn=lambda _item: (
                "株式会社フレームワークス 2025年3月期 決算説明資料 "
                "中期経営計画 2025年から2027年 売上高100億円、営業利益10億円を目標とする。"
            ),
        )

        self.assertEqual("found", plan["status"])
        self.assertEqual(ir_document["url"], plan["url"])

    def test_midterm_plan_accepts_official_plan_path_when_pdf_metadata_omits_company_name(self) -> None:
        official_plan = {
            "title": "PDF document",
            "url": "https://nsw.co.jp/ir/managementplan/MidtermManagementPlan202504.pdf",
            "snippet": "",
        }
        plan = assessment_generator.collect_midterm_plan(
            "NSW株式会社", search_fn=lambda _query, limit=12: ([official_plan], False),
            fetch_fn=lambda _item: "2025年度から2027年度までの成長方針と数値目標を掲載。",
        )
        self.assertEqual("found", plan["status"])
        self.assertEqual(official_plan["url"], plan["url"])

    def test_midterm_plan_prefers_official_ir_over_aggregated_plan_page(self) -> None:
        aggregated = {
            "title": "Example株式会社 中期経営計画 — Edinet Db",
            "url": "https://edinetdb.jp/company/example/midterm-plans/",
            "snippet": "Example株式会社の中期経営計画を構造化したページ",
        }
        official = {
            "title": "事業計画及び成長可能性に関する資料",
            "url": "https://example.co.jp/ir/management-plan.pdf",
            "snippet": "Example株式会社のIR資料",
        }
        plan = assessment_generator.collect_midterm_plan(
            "Example株式会社", search_fn=lambda _query, limit=12: ([aggregated, official], False),
            fetch_fn=lambda item: "Example株式会社 中期経営計画 2025年から2028年の売上高と営業利益の目標を掲載。",
        )
        self.assertEqual("found", plan["status"])
        self.assertEqual(official["url"], plan["url"])

    def test_quantitative_hypothesis_does_not_invent_company_targets_without_plan(self) -> None:
        hypothesis = {
            "plan_evidence_summary": "LLMが作った計画要約",
            "strategic_premises": [
                {"label": "利用者", "headline": "現場判断を支援", "detail": "サービス利用者の確認業務を対象にする。", "basis": "顧客入力"},
                {"label": "データ", "headline": "既存データを活用", "detail": "取引データを判断に用いる。", "basis": "顧客入力"},
                {"label": "検証", "headline": "実測で判断", "detail": "効果はPoCで確認する。", "basis": "PoC仮説"},
            ],
            "ai_necessity_analysis": "PoCで対象業務の効果を測定する。",
            "strategic_logic": [
                {"label": "要請", "headline": "価値", "fact": "サービスは取引データを扱う。", "decision_implication": "AI投資を優先する。"},
                {"label": "基盤", "headline": "価値", "current_constraint": "目視確認に時間が掛かる。", "ai_decision_change": "取引データから優先度を提示する。"},
                {"label": "展開", "headline": "価値", "delay_risk": "対応遅れで機会を逃す。", "proof_conditions": "工数と採用率を確認する。"},
            ],
            "management_targets": [{"label": "架空の会社目標", "source_id": "M1"}],
            "hypotheses": [
                v4_hypothesis_item(
                    i, f"テーマ{i}", "5〜10%" if i == 1 else f"{i * 10}%",
                    outcome=outcome, kpi=f"提供者事業KPI{i}",
                    display_kpi_layer="provider_business_kpi",
                    product_metric=f"{60 + i * 5}〜{75 + i * 5}%",
                    product_kpi_name=f"製品先行KPI{i}",
                )
                for i, outcome in enumerate(("回答提示時間", "異常候補提示時間", "順位採用率"), 1)
            ],
            "decision_message": "PoCを開始する。",
            "caveat": "PoC仮説である。",
        }
        with patch.object(assessment_generator, "_response_json", return_value=hypothesis):
            result = assessment_generator.build_quantitative_hypothesis(
                "会社名: 株式会社フレームワークス",
                {"poc_recommendations": [{"theme": f"テーマ{i}"} for i in range(1, 4)]},
                {"status": "not_found", "reason": "未確認"}, client=object(), model_id="test-model",
            )

        self.assertEqual("llm_estimate", result["evidence_mode"])
        self.assertEqual([], result["management_targets"])
        self.assertIn("対象サービスの利用者、業務フロー、扱うデータ", result["plan_evidence_summary"])
        self.assertEqual(3, len(result["strategic_premises"]))
        self.assertEqual("対象機能の直近12週間実績値", result["benchmarks"][0]["baseline_definition"])
        self.assertEqual("5〜10%", result["benchmarks"][0]["headline_metric"])
        self.assertIn("5〜10%", result["value_scenarios"][0]["formula"])

    def test_quantitative_hypothesis_repairs_incomplete_first_response(self) -> None:
        incomplete = {
            "hypotheses": [{
                "business_outcome": "業務生産性",
                "headline_metric": "12.5%",
                "kpi": "対応工数費",
                "baseline_definition": "対象業務の直近12か月工数費",
                "comparison_condition": "代表対象で12週間比較",
                "detail": "AI支援で判断時間を短縮する。",
                "formula": "対象工数費×12.5%",
                "poc_gate": "12週間で12.5%達成",
                "assumption": "代表対象で確認する。",
            }],
        }
        repaired = {
            "hypotheses": [
                v4_hypothesis_item(
                    index, f"テーマ{index}", metric, outcome=outcome,
                    kpi=f"提供者事業KPI{index}",
                    display_kpi_layer="provider_business_kpi",
                    product_metric=f"{60 + index * 5}〜{75 + index * 5}%",
                    product_kpi_name=f"製品先行KPI{index}",
                )
                for index, (outcome, metric) in enumerate((
                    ("回答提示時間", "7.3〜11.8%"),
                    ("異常候補提示時間", "18.6〜23.4%"),
                    ("順位採用率", "12.2〜16.9%"),
                ), 1)
            ],
            "management_targets": [],
        }
        analysis_log: list[dict] = []
        with patch.object(
            assessment_generator, "_response_json", side_effect=[incomplete, repaired],
        ):
            result = assessment_generator.build_quantitative_hypothesis(
                "会社名: Example株式会社\nサービス名: Example Service",
                {"poc_recommendations": [{"theme": "テーマ1"}, {"theme": "テーマ2"}, {"theme": "テーマ3"}]},
                {"status": "not_found"}, client=object(), model_id="test-model",
                analysis_log=analysis_log,
            )

        self.assertEqual(2, len(analysis_log))
        self.assertTrue(analysis_log[0]["issues"])
        self.assertEqual([], analysis_log[1]["issues"])
        self.assertEqual(3, len(result["benchmarks"]))
        self.assertEqual("7.3〜11.8%", result["benchmarks"][0]["headline_metric"])

    def test_quantitative_hypothesis_stops_after_three_invalid_attempts_with_audit(self) -> None:
        analysis_log: list[dict] = []
        with patch.object(
            assessment_generator, "_response_json",
            return_value={"hypotheses": [{"headline_metric": "10%"}]},
        ):
            result = assessment_generator.build_quantitative_hypothesis(
                "会社名: Example株式会社\nサービス名: Example Service",
                {"poc_recommendations": [{"theme": f"テーマ{i}"} for i in range(1, 4)]},
                {"status": "not_found"}, client=object(), model_id="test-model",
                analysis_log=analysis_log,
            )

        self.assertEqual({}, result)
        self.assertEqual(3, len(analysis_log))
        self.assertTrue(all(item["status"] == "repair_required" for item in analysis_log))
        self.assertTrue(all(item["issue_count"] > 0 for item in analysis_log))

    def test_quantitative_hypothesis_contract_rejects_mismatched_numbers(self) -> None:
        value = {"hypotheses": [
            v4_hypothesis_item(index, f"テーマ{index}", "13.7%", kpi=f"製品KPI{index}")
            for index in range(1, 4)
        ]}
        for item in value["hypotheses"]:
            item["poc_gate"]["go_condition"] = "12週間で10%を達成"
        issues = assessment_generator.quantitative_hypothesis_contract_issues(value)
        self.assertGreaterEqual(sum("poc_gate" in issue for issue in issues), 3)

    def test_quantitative_hypothesis_requires_decision_owner_and_rejects_downstream_fixed_numbers(self) -> None:
        value = {"hypotheses": [
            v4_hypothesis_item(index, f"テーマ{index}", "13.7%", kpi=f"製品KPI{index}")
            for index in range(1, 4)
        ]}
        value["hypotheses"][0]["poc_gate"].pop("decision_owner")
        value["hypotheses"][1]["customer_outcome_kpi"].update({
            "target": "顧客工数を20%削減",
            "formula": "顧客実績値×20%",
        })
        value["hypotheses"][2]["provider_business_kpi"].update({
            "target": "継続率を10%向上",
            "formula": "提供者実績値×10%",
        })

        issues = assessment_generator.quantitative_hypothesis_contract_issues(value)
        self.assertTrue(any("判定責任者" in issue for issue in issues))
        self.assertTrue(any("customer_outcome_kpi.target" in issue for issue in issues))
        self.assertTrue(any("provider_business_kpi.target" in issue for issue in issues))

    def test_quantitative_hypothesis_accepts_variable_percentage_formulas_using_times_100(self) -> None:
        """A percentage conversion constant is not an unsupported effect target.

        Real LLM responses commonly express a measured improvement rate as
        ``(baseline - measured) / baseline * 100%``.  The downstream target
        remains ``PoCで確定``; the 100% token only converts the measured ratio
        to a percentage and must not make an otherwise variable formula fail.
        """
        value = {"hypotheses": [
            v4_hypothesis_item(index, f"テーマ{index}", "13.7%", kpi=f"製品KPI{index}")
            for index in range(1, 4)
        ]}
        for item in value["hypotheses"]:
            item["customer_outcome_kpi"]["formula"] = (
                "（現状値－PoC測定値）÷現状値×100%"
            )
            item["provider_business_kpi"]["formula"] = (
                "（提供者実績値－PoC測定値）÷提供者実績値×100%"
            )

        issues = assessment_generator.quantitative_hypothesis_contract_issues(value)
        self.assertFalse(
            any("customer_outcome_kpi.formula" in issue for issue in issues), issues,
        )
        self.assertFalse(
            any("provider_business_kpi.formula" in issue for issue in issues), issues,
        )

    def test_quantitative_hypothesis_builder_repairs_downstream_fixed_claims_without_new_targets(self) -> None:
        candidate = {
            "hypotheses": [
                v4_hypothesis_item(
                    index, f"テーマ{index}", f"約{index * 10}%改善",
                    kpi=f"提供者事業KPI{index}",
                    display_kpi_layer="provider_business_kpi",
                    product_metric=f"{60 + index * 5}〜{75 + index * 5}%",
                    product_kpi_name=f"製品先行KPI{index}",
                )
                for index in range(1, 4)
            ],
            "management_targets": [],
        }
        for item in candidate["hypotheses"]:
            item["customer_outcome_kpi"].update({
                "target": "顧客工数を20%削減", "formula": "顧客実績値×20%",
            })
        analysis_log: list[dict] = []
        with patch.object(assessment_generator, "_response_json", return_value=candidate):
            result = assessment_generator.build_quantitative_hypothesis(
                "会社名: Example株式会社\nサービス名: Example Service",
                {"poc_recommendations": [{"theme": f"テーマ{i}"} for i in range(1, 4)]},
                {"status": "not_found"}, client=object(), model_id="test-model",
                analysis_log=analysis_log,
            )

        self.assertEqual(["10%", "20%", "30%"], [
            item["headline_metric"] for item in result["benchmarks"]
        ])
        self.assertEqual(1, len(analysis_log))
        self.assertEqual("accepted", analysis_log[0]["status"])
        self.assertTrue(analysis_log[0]["normalization_actions"])
        for index, item in enumerate(result["benchmarks"], 1):
            self.assertEqual("PoCで確定", item["customer_outcome_kpi"]["target"])
            self.assertEqual(f"{index * 10}%", item["provider_business_kpi"]["target"])
            self.assertNotIn("20%", item["customer_outcome_kpi"]["formula"])

    def test_quantitative_hypothesis_accepts_descriptive_percentage_targets_from_real_llm_style(self) -> None:
        """KPI prose around one LLM-authored range must not discard the range.

        The production model commonly returns ``問い合わせ解決時間20〜30%削減``
        even when the schema requests a bare ``20〜30%``.  Normalisation may
        extract that existing range, but it must not create a new number.
        """
        descriptive_targets = (
            "問い合わせ解決時間を20〜30%削減",
            "異常検知後の初動対応時間15〜25%削減",
            "誤出荷事象率10〜20%削減",
        )
        expected_metrics = ("20〜30%", "15〜25%", "10〜20%")
        candidate = {"hypotheses": [
            v4_hypothesis_item(
                index, f"テーマ{index}", descriptive_targets[index - 1],
                kpi=f"事業効果KPI{index}",
                display_kpi_layer="customer_outcome_kpi",
                product_metric=f"{60 + index * 5}〜{75 + index * 5}%",
                product_kpi_name=f"製品先行KPI{index}",
            )
            for index in range(1, 4)
        ]}
        normalized, changes = assessment_generator.normalize_quantitative_hypothesis_candidate(
            candidate,
            [{"priority": f"P{i}", "use_case_id": f"UC{i:02d}", "theme": f"テーマ{i}"}
             for i in range(1, 4)],
        )

        self.assertEqual(list(expected_metrics), [
            item["headline_metric"] for item in normalized["hypotheses"]
        ])
        self.assertEqual(list(expected_metrics), [
            item["customer_outcome_kpi"]["target"] for item in normalized["hypotheses"]
        ])
        self.assertFalse(
            assessment_generator.quantitative_hypothesis_contract_issues(
                normalized,
                [{"priority": f"P{i}", "use_case_id": f"UC{i:02d}", "theme": f"テーマ{i}"}
                 for i in range(1, 4)],
            )
        )
        self.assertTrue(any("headline_metric:format_normalized" in item for item in changes))

    def test_percentage_target_normalizer_rejects_ambiguous_or_out_of_range_prose(self) -> None:
        self.assertEqual(
            "低位10%・高位20%",
            assessment_generator._normalize_percentage_target("低位10%・高位20%"),
        )
        self.assertEqual(
            "生産性120%向上",
            assessment_generator._normalize_percentage_target("生産性120%向上"),
        )
        self.assertEqual(
            "改善率-10%",
            assessment_generator._normalize_percentage_target("改善率-10%"),
        )

    def test_real_log_percentage_headline_variants_all_normalize(self) -> None:
        real_log_variants = {
            "問い合わせ解決時間削減率20〜30%": "20〜30%",
            "異常起因の当日出荷遅延件数削減率15〜25%": "15〜25%",
            "誤出荷率削減率20〜30%": "20〜30%",
            "問い合わせ解決までの担当者工数を20〜30%削減": "20〜30%",
            "異常見逃し起因の緊急調査工数を15〜25%削減": "15〜25%",
            "誤出荷件数を10〜20%削減": "10〜20%",
            "問い合わせ解決時間20〜30%削減": "20〜30%",
            "異常検知後の初動対応時間15〜25%削減": "15〜25%",
            "誤出荷事象率10〜20%削減": "10〜20%",
        }
        self.assertEqual(
            list(real_log_variants.values()),
            [assessment_generator._normalize_percentage_target(value)
             for value in real_log_variants],
        )

    def test_quantitative_hypothesis_normalizer_does_not_create_a_missing_numeric_target(self) -> None:
        candidate = {"hypotheses": [
            v4_hypothesis_item(index, f"テーマ{index}", "要確認", kpi=f"製品KPI{index}")
            for index in range(1, 4)
        ]}
        normalized, _changes = assessment_generator.normalize_quantitative_hypothesis_candidate(
            candidate,
            [{"priority": f"P{i}", "use_case_id": f"UC{i:02d}", "theme": f"テーマ{i}"}
             for i in range(1, 4)],
        )
        self.assertEqual(["要確認", "要確認", "要確認"], [
            item["headline_metric"] for item in normalized["hypotheses"]
        ])
        self.assertTrue(assessment_generator.quantitative_hypothesis_contract_issues(normalized))

    def test_research_retries_when_evidence_is_insufficient(self) -> None:
        replies = iter([
            {"queries": ["q1", "q2", "q3", "q4"]},
            {"sufficient": False, "missing_evidence": ["定量事例不足"],
             "additional_queries": ["q5", "q6"]},
            {"sufficient": True, "missing_evidence": [], "additional_queries": []},
        ])

        def fake_search(query: str) -> list[dict[str, str]]:
            return [{"title": f"title-{query}", "url": f"https://example.com/{query}",
                     "snippet": f"{query} improves KPI by 20%."}]

        audit_log: list[dict] = []
        with patch.object(assessment_generator, "_response_json", side_effect=lambda *args, **kwargs: next(replies)), \
             patch.object(assessment_generator, "httpx", None), \
             patch.object(assessment_generator, "BeautifulSoup", None):
            sources = assessment_generator.collect_industry_research(
                "会社名: Example\nサービス名: Example Service",
                client=object(), model_id="test-model",
                assessment={"poc_recommendations": [
                    {"theme": "theme-a"}, {"theme": "theme-b"}, {"theme": "theme-c"},
                ]},
                max_rounds=3, audit_log=audit_log, search_fn=fake_search,
            )

        self.assertEqual(6, len(sources))
        self.assertEqual([f"R{index}" for index in range(1, 7)], [source["id"] for source in sources])
        self.assertEqual(2, len(audit_log))
        self.assertFalse(audit_log[0]["audit"]["sufficient"])
        self.assertTrue(audit_log[1]["audit"]["sufficient"])

    def test_research_url_guard_rejects_private_and_non_https_destinations(self) -> None:
        self.assertTrue(assessment_generator.is_safe_public_https_url("https://www.example.com/report.pdf"))
        for url in (
            "http://www.example.com/report.pdf", "https://localhost/report.pdf",
            "https://127.0.0.1/report.pdf", "https://10.0.0.1/report.pdf",
            "https://[::1]/report.pdf", "https://user:pass@example.com/report.pdf",
        ):
            self.assertFalse(assessment_generator.is_safe_public_https_url(url), url)

    def test_executive_analysis_repairs_invalid_first_response(self) -> None:
        invalid = {"benchmarks": [], "value_scenarios": []}
        valid = {
            "plan_evidence_summary": "公開計画に記載された成長目標を、対象サービスのAI投資と接続する。",
            "strategic_premises": [
                {"label": "事業機会", "headline": "業務品質を高める", "detail": "公開事例の効果を同じKPIで検証する。", "basis": "公開資料"},
                {"label": "優先テーマ", "headline": "判断業務を優先", "detail": "既存データを使う判断支援を優先する。", "basis": "顧客入力"},
                {"label": "事業化条件", "headline": "代表範囲で実測", "detail": "同一条件比較で本番化を判断する。", "basis": "PoC仮説"},
            ],
            "ai_necessity_analysis": "対象サービスの業務データを活用し、AIで判断支援と自動化を段階的に検証する。",
            "strategic_logic": [
                {"label": "経営上の要請", "headline": "成長を支援", "fact": "公開計画は成長を掲げる。", "decision_implication": "AI投資を段階的に検証する。"},
                {"label": "既存基盤 × AI", "headline": "判断を高度化", "current_constraint": "人手確認が残る。", "ai_decision_change": "業務データから担当者へ提案を行う。"},
                {"label": "PoCから事業化へ", "headline": "効果を測定", "delay_risk": "対応遅れで競争力を失う。", "proof_conditions": "代表データで経営効果と連携負荷を確認する。"},
            ],
            "management_targets": [
                {"label": "成長目標", "baseline": "100億円", "target": "120億円", "baseline_value": 100,
                 "target_value": 120, "unit": "億円", "direction": "increase", "delta": "+20億円",
                 "implication": "AI施策で成長を支援", "source_id": "M1", "source_locator": "数値目標",
                 "source_metric": "100億円から120億円", "evidence_excerpt": "成長目標は100億円から120億円"},
                {"label": "利益目標", "baseline": "10%", "target": "12%", "baseline_value": 10,
                 "target_value": 12, "unit": "%", "direction": "increase", "delta": "+2ポイント",
                 "implication": "AI施策で収益性を支援", "source_id": "M1", "source_locator": "数値目標",
                 "source_metric": "利益率は10%から12%", "evidence_excerpt": "利益率は10%から12%"},
            ],
            "benchmarks": [{
                "use_case": f"施策{i}", "headline_metric": f"{i * 10}%", "kpi": "工数",
                "detail": "公開事例に記載された改善値", "source_id": f"R{i}", "source_title": f"事例{i}",
                "source_locator": "導入効果", "source_metric": f"工数を{i * 10}%削減",
                "evidence_excerpt": f"工数を{i * 10}%削減",
            } for i in range(1, 4)],
            "value_scenarios": [{
                "use_case": f"施策{i}", "benchmark": f"{i * 10}%", "formula": f"年間工数×{i * 10}%",
                "example": "", "poc_gate": f"現行比{i * 10}%を確認", "source_ids": [f"R{i}"],
            } for i in range(1, 4)],
            "decision_message": "代表データでPoCを開始する。",
            "caveat": "効果はPoCで検証する。",
        }
        for index, (benchmark, scenario) in enumerate(zip(valid["benchmarks"], valid["value_scenarios"]), 1):
            layer = v4_quantitative_item(index, f"施策{index}", f"{index * 10}%", kpi="工数")
            benchmark.update({
                "priority": layer["priority"], "use_case_id": layer["use_case_id"],
                "theme": layer["theme"], "display_kpi_layer": "product_kpi",
                "beneficiary": layer["beneficiary"], "metric_owner": layer["metric_owner"],
                "attribution_level": layer["attribution_level"], "causal_link": layer["causal_link"],
                "product_kpi": layer["product_kpi"],
                "customer_outcome_kpi": layer["customer_outcome_kpi"],
                "provider_business_kpi": layer["provider_business_kpi"],
            })
            scenario.update({
                "priority": layer["priority"], "use_case_id": layer["use_case_id"],
                "theme": layer["theme"], "poc_gate": layer["poc_gate"],
            })
        assessment = {"poc_recommendations": [{"theme": f"施策{i}"} for i in range(1, 4)]}
        research_sources = [
            {"id": f"R{i}", "title": f"事例{i}", "url": f"https://example.com/case{i}",
             "excerpt": f"工数を{i * 10}%削減"}
            for i in range(1, 4)
        ]
        with patch.object(assessment_generator, "_response_json", side_effect=[
            invalid, valid, {"all_supported": True, "issues": []},
        ]) as mocked:
            result = assessment_generator.build_executive_evidence_analysis(
                "顧客情報", assessment, research_sources,
                {"status": "not_found", "reason": "未確認"},
                client=object(), model_id="test-model",
            )
        self.assertEqual(3, mocked.call_count)
        self.assertEqual(3, len(result["benchmarks"]))
        self.assertEqual(3, len(result["value_scenarios"]))
        self.assertEqual(["Q01", "Q02", "Q03"], [item["claim_id"] for item in result["benchmarks"]])
        self.assertTrue(all(item["claim_status"] == "verified_external" for item in result["benchmarks"]))
        self.assertTrue(all(item["source_url"].startswith("https://") for item in result["benchmarks"]))
        self.assertEqual(["R1", "R2", "R3"], result["generation_audit"]["calibration_source_ids"])
        self.assertRegex(result["generation_audit"]["base_prompt_input_sha256"], r"^[0-9a-f]{64}$")
        research = {
            "midterm_plan": {"status": "disabled"}, "industry_sources": research_sources,
            "quantitative_evidence": {"approved_source_ids": ["R1", "R2", "R3"]},
        }
        design = assessment_generator.poc_measurement_design_from_verified_evidence(
            result, research, assessment,
        )
        self.assertEqual("external_verified", design["status"])
        self.assertEqual(
            assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
            design["schema_version"],
        )
        self.assertEqual(
            [(f"P{i}", f"UC{i:02d}", f"施策{i}") for i in range(1, 4)],
            [(item["priority"], item["use_case_id"], item["theme"]) for item in design["items"]],
        )

    def test_quantitative_source_selection_keeps_partial_verified_evidence(self) -> None:
        """One usable public benchmark must not be discarded solely for lacking three cases."""
        sources = [
            {"id": "R1", "title": "効果事例", "url": "https://example.com/case",
             "excerpt": "問い合わせ対応時間を25%削減した。"},
            {"id": "R2", "title": "一般資料", "url": "https://example.com/overview",
             "excerpt": "製品の機能を紹介する。"},
        ]
        with patch.object(assessment_generator, "_response_json", return_value={
            "approved_source_ids": ["R1"], "reason": "対応時間の削減率を本文で確認できる。",
        }):
            selected = assessment_generator.select_verified_quantitative_sources(
                "サービス情報", {"poc_recommendations": [{"theme": "問い合わせ回答支援"}]},
                sources, client=object(), model_id="test-model",
            )
        self.assertEqual(["R1"], selected["approved_source_ids"])

    def test_quantitative_evidence_window_finds_effect_beyond_page_intro(self) -> None:
        source = "概要説明。" + ("一般的な説明文。" * 260) + "導入後、問い合わせ対応工数を25%削減した。" + ("補足。" * 80)
        excerpt = assessment_generator.quantitative_evidence_window(source, max_chars=900)
        self.assertIn("問い合わせ対応工数を25%削減", excerpt)
        self.assertLessEqual(len(excerpt), 900)

    def test_independent_verifier_uses_customer_input_and_only_claim_sources(self) -> None:
        valid = {
            "plan_evidence_summary": "公開方針と顧客入力を接続する。",
            "strategic_premises": [
                {"label": "事業機会", "headline": "品質を高める", "detail": "公開事例と同じKPIを測る。", "basis": "公開資料"},
                {"label": "優先テーマ", "headline": "判断を支援", "detail": "顧客入力の業務データを使う。", "basis": "顧客入力"},
                {"label": "事業化条件", "headline": "実測で判断", "detail": "同一条件で本番化を判断する。", "basis": "PoC仮説"},
            ],
            "ai_necessity_analysis": "顧客入力に記載された対象サービスの業務データを使い、AIで判断支援を行う。",
            "strategic_logic": [
                {"label": "経営上の要請", "headline": "成長を支援", "fact": "顧客入力に対象業務がある。", "decision_implication": "段階的にAIを検証する。"},
                {"label": "既存基盤 × AI", "headline": "判断を高度化", "current_constraint": "人手確認が残る。", "ai_decision_change": "業務データから提案する。"},
                {"label": "PoCから事業化へ", "headline": "実測で判断", "delay_risk": "対応遅れで価値を失う。", "proof_conditions": "代表データで効果を確認する。"},
            ],
            "management_targets": [],
            "benchmarks": [{
                "use_case": "問い合わせ回答支援", "headline_metric": "25%", "kpi": "対応工数",
                "detail": "公開事例で対応工数を25%削減した。", "source_id": "R1", "source_title": "公開導入事例",
                "source_locator": "導入効果", "source_metric": "対応工数を25%削減",
                "evidence_excerpt": "対応工数を25%削減",
            }],
            "value_scenarios": [{
                "use_case": "問い合わせ回答支援", "benchmark": "25%", "formula": "現行対応工数×25%を参照",
                "example": "", "poc_gate": "同一KPIを測定し、公開参照水準25%との差を評価", "source_ids": ["R1"],
            }],
            "decision_message": "代表データでPoCを開始する。", "caveat": "公開事例は参照水準とする。",
        }
        layer = v4_quantitative_item(1, "問い合わせ回答支援", "25%", kpi="対応工数")
        valid["benchmarks"][0].update({
            "priority": layer["priority"], "use_case_id": layer["use_case_id"],
            "theme": layer["theme"], "display_kpi_layer": "product_kpi",
            "beneficiary": layer["beneficiary"], "metric_owner": layer["metric_owner"],
            "attribution_level": layer["attribution_level"], "causal_link": layer["causal_link"],
            "product_kpi": layer["product_kpi"],
            "customer_outcome_kpi": layer["customer_outcome_kpi"],
            "provider_business_kpi": layer["provider_business_kpi"],
        })
        valid["value_scenarios"][0].update({
            "priority": layer["priority"], "use_case_id": layer["use_case_id"],
            "theme": layer["theme"], "poc_gate": layer["poc_gate"],
        })
        with patch.object(assessment_generator, "_response_json", side_effect=[
            valid, {"all_supported": True, "issues": []},
        ]) as mocked:
            result = assessment_generator.build_executive_evidence_analysis(
                "会社名: Example株式会社\nサービス名: Example Service\n問い合わせ履歴を扱う。",
                {"poc_recommendations": [
                    {"theme": "問い合わせ回答支援"},
                    {"theme": "異常検知"}, {"theme": "作業優先度支援"},
                ]},
                [
                    {"id": "R1", "title": "採用事例", "url": "https://example.com/case1", "excerpt": "対応工数を25%削減"},
                    {"id": "R2", "title": "未使用事例", "url": "https://example.com/case2", "excerpt": "在庫を10%削減"},
                ],
                {"status": "not_found", "reason": "未確認"}, client=object(), model_id="test-model",
            )
        self.assertEqual(1, len(result["benchmarks"]))
        verification_prompt = mocked.call_args_list[1].kwargs["prompt"]
        self.assertIn("顧客入力（会社・サービス固有の事実I1）", verification_prompt)
        self.assertIn("問い合わせ履歴を扱う", verification_prompt)
        self.assertIn('"id": "R1"', verification_prompt)
        self.assertNotIn('"id": "R2"', verification_prompt)


class DynamicPptxPageTests(unittest.TestCase):
    @staticmethod
    def verified_quantitative_fixture() -> tuple[dict, dict]:
        benchmarks = []
        scenarios = []
        sources = []
        for index, metric in enumerate(("10%", "20%", "30%"), 1):
            source_id = f"R{index}"
            use_case = f"ユースケース{index}"
            item = v4_quantitative_item(index, use_case, metric, outcome=f"顧客業務成果{index}")
            excerpt = f"対象業務の処理コストを{metric}削減"
            sources.append({
                "id": source_id, "title": f"公開事例{index}",
                "url": f"https://example.com/case{index}", "excerpt": excerpt,
                "fetch_status": "fetched",
            })
            benchmarks.append({
                "claim_id": f"Q0{index}", "claim_status": "verified_external",
                "priority": item["priority"], "use_case_id": item["use_case_id"], "theme": item["theme"],
                "use_case": use_case, "headline_metric": metric, "kpi": "処理コスト",
                "detail": f"公開事例では対象業務の処理コストを{metric}削減した。",
                "source_id": source_id, "source_url": f"https://example.com/case{index}",
                "source_title": f"公開事例{index}", "source_locator": "導入効果",
                "source_metric": excerpt, "evidence_excerpt": excerpt,
                "beneficiary": item["beneficiary"], "metric_owner": item["metric_owner"],
                "attribution_level": item["attribution_level"], "causal_link": item["causal_link"],
                "product_kpi": {**item["product_kpi"], "kpi": "処理コスト"},
                "customer_outcome_kpi": item["customer_outcome_kpi"],
                "provider_business_kpi": item["provider_business_kpi"],
            })
            scenarios.append({
                "priority": item["priority"], "use_case_id": item["use_case_id"], "theme": item["theme"],
                "use_case": use_case, "benchmark": metric,
                "formula": f"年間対象コスト×{metric}", "example": "",
                "poc_gate": item["poc_gate"], "source_ids": [source_id],
            })
        evidence = {
            "evidence_mode": "external_verified",
            "plan_evidence_summary": "公開方針と対象業務を接続する。",
            "ai_necessity_analysis": "対象業務の判断をAIで支援し、公開事例と同じKPIで事業効果を確認する。",
            "strategic_premises": [
                {"label": "事業機会", "headline": "品質を高める", "detail": "公開事例の効果を同じKPIで検証する。", "basis": "公開資料"},
                {"label": "優先テーマ", "headline": "判断を支援", "detail": "既存データを使う判断支援を優先する。", "basis": "顧客入力"},
                {"label": "事業化条件", "headline": "実測で判断", "detail": "同一条件比較で本番化を判断する。", "basis": "PoC仮説"},
            ],
            "strategic_logic": [{"label": "a"}, {"label": "b"}, {"label": "c"}],
            "management_targets": [], "benchmarks": benchmarks,
            "value_scenarios": scenarios, "sources": [],
        }
        research = {
            "midterm_plan": {"status": "disabled"}, "industry_sources": sources,
            "quantitative_evidence": {
                "schema_version": assessment_generator.QUANTITATIVE_EVIDENCE_SCHEMA_VERSION,
                "status": "sufficient", "approved_source_ids": ["R1", "R2", "R3"],
            },
        }
        return evidence, research

    @staticmethod
    def llm_quantitative_estimate_fixture() -> dict:
        metrics = (
            ("E01", "売上・粗利機会", "7.3〜11.8%", "追加粗利額", 7.3, 11.8),
            ("E02", "業務生産性・コスト", "18.6〜23.4%", "対応工数費", 18.6, 23.4),
            ("E03", "損失・品質リスク", "12.2〜16.9%", "実績損失額", 12.2, 16.9),
        )
        benchmarks = []
        scenarios = []
        for index, (estimate_id, outcome, metric, kpi, low, high) in enumerate(metrics, 1):
            display_layer = "provider_business_kpi" if index in {1, 2} else "customer_outcome_kpi"
            item = v4_quantitative_item(
                index, f"ユースケース{index}", metric, outcome=outcome,
                kpi=kpi, display_kpi_layer=display_layer,
                product_metric=("70〜85%", "15〜25%", "60〜75%")[index - 1],
                product_kpi_name=f"製品先行KPI{index}",
            )
            benchmarks.append({
                "estimate_id": estimate_id,
                "estimate_low": low,
                "estimate_high": high,
                "estimate_unit": "%",
                "priority": item["priority"],
                "use_case_id": item["use_case_id"],
                "theme": item["theme"],
                "use_case": outcome,
                "headline_metric": metric,
                "kpi": item["kpi"],
                "detail": "対象サービスの業務データと優先PoCから、AI導入効果を設定する。",
                "baseline_definition": item["baseline_definition"],
                "comparison_condition": item["comparison_condition"],
                "source_title": "対象範囲とデータ利用可能性を踏まえたPoC目標レンジ",
                "display_kpi_layer": item["display_kpi_layer"],
                "beneficiary": item["beneficiary"], "metric_owner": item["metric_owner"],
                "attribution_level": item["attribution_level"], "causal_link": item["causal_link"],
                "product_kpi": item["product_kpi"],
                "customer_outcome_kpi": item["customer_outcome_kpi"],
                "provider_business_kpi": item["provider_business_kpi"],
            })
            scenarios.append({
                "estimate_id": estimate_id,
                "priority": item["priority"], "use_case_id": item["use_case_id"], "theme": item["theme"],
                "use_case": outcome,
                "benchmark": metric,
                "formula": item["formula"],
                "example": "",
                "poc_gate": item["poc_gate"],
            })
        return {
            "evidence_mode": "llm_estimate",
            "management_targets": [],
            "benchmarks": benchmarks,
            "value_scenarios": scenarios,
        }

    @staticmethod
    def detailed_assessment() -> dict:
        assessment = {
            "company_name": "Example株式会社", "service_name": "Example Service",
            "executive_summary": "対象サービスの利用者、業務データ、優先判断を基に、段階的なAI実装の開始条件を整理する。",
            "assessment_points": ["利用者が必要な情報を適時に得られるようにする。", "担当者の確認・優先順位付けを支援する。", "PoCの評価結果を標準化判断へ接続する。"],
            "business_value": {"impact_areas": [
                {"area": "顧客価値", "business_rationale": "利用者の判断を支援する。", "ai_enabled": "情報を整理する。", "expected_impact": "利用体験を評価する。"},
                {"area": "業務生産性", "business_rationale": "確認を支援する。", "ai_enabled": "優先度を提示する。", "expected_impact": "確認工数を評価する。"},
                {"area": "品質・統制", "business_rationale": "例外を把握する。", "ai_enabled": "根拠を表示する。", "expected_impact": "品質を評価する。"},
            ]},
            "poc_recommendations": [
                {"theme": f"優先テーマ{index}", "reason": "代表業務で事業価値と実現性を確認する。", "first_step": "対象データと比較条件を合意する。"}
                for index in range(1, 4)
            ],
            "use_cases": [
                {"no": index, "use_case": f"ユースケース{index}", "ai_technology": "生成AI", "description": "業務データを整理する。判断を支援する。"}
                for index in range(1, 16)
            ],
        }
        assessment["consulting_front_matter"] = assessment_generator.fallback_consulting_front_matter(assessment)
        return assessment

    @staticmethod
    def pptx_slide_text(slide) -> str:
        """PPTXの通常テキストと表セルを同じ検証対象として収集する。"""
        fragments: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                fragments.append(shape.text)
            if shape.has_table:
                fragments.extend(cell.text for row in shape.table.rows for cell in row.cells)
        return "\n".join(fragments)

    @staticmethod
    def is_ten_point_chrome_shape(shape, slide_height: int) -> bool:
        """12pt本文契約の例外となる共通ヘッダー・フッターだけを識別する。"""
        if not shape.has_text_frame:
            return False
        text = shape.text.strip()
        return (
            (text == "OCI AI USE CASE ASSESSMENT" and shape.top < slide_height * 0.12)
            or (text.startswith("Copyright ©") and shape.top > slide_height * 0.85)
            or (text.isdigit() and shape.top > slide_height * 0.9)
        )

    @staticmethod
    def is_nine_point_source_shape(shape) -> bool:
        """本文下限の明示的な例外となる業界動向ページの出典注記。"""
        return getattr(shape, "name", "") == "AI_ASSESS_SOURCE_NOTE"

    @staticmethod
    def pptx_text_font_sizes(
        presentation: Presentation, *, include_footer: bool = True,
        excluded_slide_indexes: set[int] | None = None,
    ) -> list[float | None]:
        """編集可能テキストのサイズを返す。Falseなら共通クロームと出典を除く。"""
        sizes: list[float | None] = []
        excluded_slide_indexes = excluded_slide_indexes or set()
        for slide_index, slide in enumerate(presentation.slides):
            if slide_index in excluded_slide_indexes:
                continue
            for shape in slide.shapes:
                is_small_chrome = DynamicPptxPageTests.is_ten_point_chrome_shape(
                    shape, presentation.slide_height,
                )
                is_source_note = DynamicPptxPageTests.is_nine_point_source_shape(shape)
                if not include_footer and (is_small_chrome or is_source_note):
                    continue
                frames = [shape.text_frame] if shape.has_text_frame else []
                if shape.has_table:
                    frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
                for frame in frames:
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text.strip():
                                sizes.append(run.font.size.pt if run.font.size else None)
        return sizes

    @staticmethod
    def pptx_text_font_names(
        presentation: Presentation, *, excluded_slide_indexes: set[int] | None = None,
    ) -> list[str | None]:
        """画像以外の編集可能テキストに設定されたフォント名を返す。"""
        names: list[str | None] = []
        excluded_slide_indexes = excluded_slide_indexes or set()
        for slide_index, slide in enumerate(presentation.slides):
            if slide_index in excluded_slide_indexes:
                continue
            for shape in slide.shapes:
                frames = [shape.text_frame] if shape.has_text_frame else []
                if shape.has_table:
                    frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
                for frame in frames:
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text.strip():
                                names.append(run.font.name)
        return names

    def test_assessment_story_builds_reusable_front_matter_from_assessment(self) -> None:
        story = assessment_generator.assessment_story_for({
            "company_name": "Example株式会社", "service_name": "Example Service", "executive_summary": "要約",
            "assessment_points": ["論点1", "論点2", "論点3"],
            "business_value": {"impact_areas": [
                {"area": "価値1", "business_rationale": "理由", "ai_enabled": "判断", "expected_impact": "効果"},
            ]},
            "poc_recommendations": [{"theme": "PoC", "reason": "理由", "first_step": "次の手順"}],
            "use_cases": [{"no": 1, "use_case": "候補"}],
        })
        self.assertEqual("Example Service", story["service_name"])
        self.assertEqual(["論点1", "論点2", "論点3"], story["points"])
        self.assertEqual("PoC", story["pocs"][0]["theme"])

    def test_industry_value_story_is_derived_from_input_specific_front_matter(self) -> None:
        assessment = self.detailed_assessment()
        story = assessment_generator.assessment_story_for(assessment)["industry_value_story"]
        self.assertEqual(3, len(story["drivers"]))
        self.assertEqual(4, len(story["flow"]))
        self.assertEqual(3, len(story["outcomes"]))
        self.assertIn("Example Service", story["decision_question"])
        self.assertEqual("顧客価値", story["outcomes"][0]["label"])

    def test_business_value_model_separates_customer_outcomes_from_provider_kpis(self) -> None:
        assessment = self.detailed_assessment()
        assessment["business_model_role"] = "provider"
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        model = assessment_generator.business_value_model_for(
            assessment, assessment["consulting_front_matter"],
        )

        self.assertEqual("provider", model["organization_role"])
        self.assertTrue(model["provider"]["applicable"])
        self.assertEqual(["品質", "時間", "リスク", "定着"], [
            item["dimension"] for item in model["customer_operation"]["outcomes"]
        ])
        self.assertEqual(["MRR・ARPA", "継続利用・定着", "サポート原価", "横展開・スケール"], [
            item["metric"] for item in model["provider"]["metrics"]
        ])
        self.assertEqual([item["theme"] for item in assessment["poc_recommendations"]], [
            item["theme"] for item in model["provider"]["packaging"]
        ])
        self.assertTrue(assessment_generator.normalize_business_value_model(model))
        assessment_generator.materialize_business_value_model(assessment, assessment["consulting_front_matter"])
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertIn("business_value_model", assessment)
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))

    def test_business_value_model_can_suppress_provider_layer_for_operator_only_assessment(self) -> None:
        assessment = self.detailed_assessment()
        assessment["business_model_role"] = "operator"
        model = assessment_generator.fallback_business_value_model(
            assessment, assessment["consulting_front_matter"],
        )

        self.assertEqual("operator", model["organization_role"])
        self.assertFalse(model["provider"]["applicable"])
        self.assertEqual("導入先側：業務・利用価値", model["customer_operation"]["label"])

    def test_business_value_model_rejects_malformed_optional_payload(self) -> None:
        assessment = self.detailed_assessment()
        model = assessment_generator.fallback_business_value_model(
            assessment, assessment["consulting_front_matter"],
        )
        model["customer_operation"]["outcomes"][0]["dimension"] = "別の軸"
        assessment["business_value_model"] = model
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}

        errors = assessment_generator.validate_assessment_payload(payload)
        self.assertTrue(any("business_value_model" in error for error in errors))




    def test_consultative_page_lead_turns_confirmation_notes_into_customer_guidance(self) -> None:
        lead = assessment_generator.consultative_page_lead(
            "データ・評価・責任者が未取得の項目は「要確認」として残し、PoC開始前に確定します。"
            "数値や責任者を仮置きで事実化しません。",
            title="PoCチャーター：対象・データ条件",
        )
        self.assertIn("必要となるデータ、評価方法、役割分担", lead)
        self.assertIn("次の投資判断", lead)
        self.assertNotIn("事実化しません", lead)

    def test_consultative_page_lead_does_not_append_the_same_generic_sentence(self) -> None:
        source = "対象業務の判断点と利用データを整理し、PoCで確認する評価条件を具体化します。"
        lead = assessment_generator.consultative_page_lead(source, title="評価条件")

        self.assertEqual(source, lead)
        self.assertNotIn("現状から読み取れる論点と判断に必要な材料", lead)
        self.assertNotIn("優先順位と次のアクションを具体的に合意", lead)

    def test_display_text_helpers_preserve_full_source_despite_legacy_limits(self) -> None:
        """版面都合の旧limitでは、PPTXへ渡る原文を失わない。"""
        source = "複数テナントの就業規則・設定・FAQ・対応履歴を権限別に参照し、根拠付き回答と監査記録を提供する。"
        service_name = "勤怠・工数・人事手続を統合する「RocoTime Enterprise Platform」"

        self.assertEqual(source, assessment_generator._front_text(source, 8))
        self.assertEqual(source, assessment_generator.consultative_page_lead(source, title="評価条件", limit=8))
        self.assertEqual(service_name, assessment_generator.compact_service_name_for_heading(service_name))

    def test_detailed_front_matter_has_valid_counts_and_basis_labels(self) -> None:
        assessment = self.detailed_assessment()
        front = assessment["consulting_front_matter"]
        self.assertTrue(assessment_generator.normalize_consulting_front_matter(front, assessment, {"I1"}))
        self.assertEqual(5, len(front["service_model"]["value_chain"]))
        self.assertEqual(15, len(front["use_case_prioritization"]["candidates"]))
        self.assertEqual([str(index) for index in range(1, 16)], [
            item["use_case_no"] for item in front["use_case_prioritization"]["candidates"]
        ])
        invalid = assessment_generator.normalize_adb_terminology(front.copy())
        invalid["operating_diagnosis"]["data_assets"][0]["readiness"] = "unsupported"  # type: ignore[index]
        self.assertEqual({}, assessment_generator.normalize_consulting_front_matter(invalid, assessment, {"I1"}))



    def test_pptx_draw_paragraph_preserves_html_break_as_newline(self) -> None:
        """PPTX描画でも`<br/>`を文字列のまま残さず、明示改行として扱う。"""
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        style = assessment_generator.ParagraphStyle(
            "html_break", fontName="AssessmentJapanese", fontSize=10.0,
            leading=12.0, textColor=assessment_generator.colors.HexColor("#1D252C"),
            wordWrap="CJK",
        )
        assessment_generator.draw_paragraph(
            ppt, "見出し<br/>本文", style, 20 * assessment_generator.mm,
            120 * assessment_generator.mm, 80 * assessment_generator.mm,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "html-break.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        shape = next(shape for shape in presentation.slides[0].shapes if shape.has_text_frame)
        self.assertEqual("見出し\n本文", shape.text)

    def test_payload_validation_rejects_invalid_detailed_front_matter(self) -> None:
        assessment = self.detailed_assessment()
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))
        assessment["consulting_front_matter"]["use_case_prioritization"]["candidates"][0]["basis"] = "未定義"
        self.assertTrue(any("consulting_front_matter" in error for error in assessment_generator.validate_assessment_payload(payload)))

    def test_front_matter_builder_preserves_a_valid_modelled_structure(self) -> None:
        assessment = self.detailed_assessment()
        expected = assessment["consulting_front_matter"]
        with patch.object(assessment_generator, "_response_json", return_value=expected) as mocked:
            result = assessment_generator.build_consulting_front_matter(
                "会社名: Example株式会社\nサービス名: Example Service", assessment,
                midterm_plan={"status": "disabled"}, research_sources=[], client=object(), model_id="test-model",
            )
        self.assertEqual(1, mocked.call_count)
        self.assertEqual(expected["decision_context"], result["decision_context"])
        self.assertEqual("利用・入力", result["use_case_prioritization"]["candidates"][0]["workflow_stage"])

    def test_front_matter_repair_keeps_valid_customer_specific_rows(self) -> None:
        assessment = self.detailed_assessment()
        raw = json.loads(json.dumps(assessment["consulting_front_matter"], ensure_ascii=False))
        raw["service_model"]["value_chain"][0]["stage"] = "顧客固有の利用工程"
        raw["operating_diagnosis"]["issue_tree"][0]["issue"] = "顧客固有の構造的課題"
        raw["operating_diagnosis"]["data_assets"][0]["readiness"] = "高"
        raw["use_case_prioritization"]["candidates"][0]["data_readiness_reason"] = (
            "対象データは特定済みだが、期間・件数・欠損・権限・正解条件を確認する必要がある。"
        )
        raw["use_case_prioritization"]["candidates"][0]["data_next_action"] = (
            "代表期間を抽出し、件数・欠損・権限・正解条件を業務責任者と確認する。"
        )
        for candidate in raw["use_case_prioritization"]["candidates"]:
            candidate.pop("workflow_stage")
        repaired = assessment_generator.repair_consulting_front_matter(raw, assessment, {"I1"})
        normalized = assessment_generator.normalize_consulting_front_matter(repaired, assessment, {"I1"})
        self.assertTrue(normalized)
        self.assertEqual("顧客固有の利用工程", normalized["service_model"]["value_chain"][0]["stage"])
        self.assertEqual("顧客固有の構造的課題", normalized["operating_diagnosis"]["issue_tree"][0]["issue"])
        self.assertEqual("high", normalized["operating_diagnosis"]["data_assets"][0]["readiness"])
        self.assertEqual("顧客固有の利用工程", normalized["use_case_prioritization"]["candidates"][0]["workflow_stage"])
        self.assertIn("期間・件数・欠損・権限", normalized["use_case_prioritization"]["candidates"][0]["data_readiness_reason"])
        self.assertIn("代表期間を抽出", normalized["use_case_prioritization"]["candidates"][0]["data_next_action"])

    def test_front_matter_normalizes_duplicate_priority_labels_to_three_pocs(self) -> None:
        assessment = self.detailed_assessment()
        raw = json.loads(json.dumps(assessment["consulting_front_matter"], ensure_ascii=False))
        for item in raw["use_case_prioritization"]["candidates"]:
            item["priority"] = "P1"
        normalized = assessment_generator.normalize_consulting_front_matter(raw, assessment, {"I1"})
        self.assertEqual(["P1", "P2", "P3"], [
            item["priority"] for item in normalized["use_case_prioritization"]["candidates"][:3]
        ])
        self.assertTrue(all(item["priority"] == "Watch" for item in normalized["use_case_prioritization"]["candidates"][3:]))

    def test_poc_portfolio_materializes_one_canonical_priority_set(self) -> None:
        assessment = self.detailed_assessment()
        portfolio = assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )

        self.assertEqual(["UC01", "UC02", "UC03"], [item["use_case_id"] for item in portfolio["items"]])
        self.assertEqual(["ユースケース1", "ユースケース2", "ユースケース3"], [
            item["theme"] for item in assessment["poc_recommendations"]
        ])
        self.assertEqual(["ユースケース1", "ユースケース2", "ユースケース3"], [
            item["theme"] for item in assessment["consulting_front_matter"]["use_case_prioritization"]["selection_logic"]
        ])
        self.assertEqual(["UC01", "UC02", "UC03"], [
            item["use_case_id"] for item in assessment["poc_logic_details"]
        ])
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))

    def test_poc_portfolio_rejects_cross_theme_priority_reference(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        assessment["poc_portfolio"]["items"][1]["use_case_id"] = "UC01"
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}

        errors = assessment_generator.validate_assessment_payload(payload)
        self.assertTrue(any("poc_portfolio" in error for error in errors))

    def test_poc_decision_data_materializes_confirm_values_without_inventing_measurements(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        decision_data = assessment_generator.materialize_poc_decision_data(
            assessment, assessment["consulting_front_matter"],
        )

        charters = decision_data["poc_charters"]["charters"]
        self.assertEqual(["P1", "P2", "P3"], [item["priority"] for item in charters])
        self.assertEqual(["UC01", "UC02", "UC03"], [item["use_case_id"] for item in charters])
        for charter in charters:
            self.assertEqual("confirm", charter["data_period"])
            self.assertEqual("confirm", charter["data_volume"])
            self.assertEqual("confirm", charter["missingness"])
            self.assertEqual("confirm", charter["ground_truth"])
            self.assertEqual("confirm", charter["evaluator"])
            self.assertEqual("confirm", charter["comparator"])
            self.assertEqual("confirm", charter["baseline"])
            self.assertEqual("confirm", charter["success_criteria"])
            self.assertEqual("confirm", charter["stop_criteria"])
            self.assertTrue(all(value == "confirm" for value in charter["owners"].values()))
        self.assertEqual("confirm", decision_data["multitenant_governance"]["rag_documents"]["prompt_injection"])
        self.assertEqual("confirm", decision_data["multitenant_governance"]["answer_controls"]["rollback"])
        self.assertEqual("confirm", decision_data["multitenant_governance"]["model_operations"]["drift"])

        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))

    def test_poc_decision_data_rejects_cross_theme_and_incomplete_governance(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        assessment_generator.materialize_poc_decision_data(assessment, assessment["consulting_front_matter"])
        assessment["poc_charters"]["charters"][1]["use_case_id"] = "UC01"
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertTrue(any("poc_charters" in error for error in assessment_generator.validate_assessment_payload(payload)))

        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        assessment_generator.materialize_poc_decision_data(assessment, assessment["consulting_front_matter"])
        assessment["multitenant_governance"]["answer_controls"].pop("rollback")
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertTrue(any("multitenant_governance" in error for error in assessment_generator.validate_assessment_payload(payload)))

    def test_poc_selection_scorecard_keeps_unknown_start_gates_as_confirm(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        scorecard = assessment_generator.materialize_poc_selection_scorecard(
            assessment, assessment["consulting_front_matter"],
        )

        self.assertEqual(["UC01", "UC02", "UC03"], [
            item["use_case_id"] for item in scorecard["items"]
        ])
        self.assertTrue(all(
            item["scores"]["business_value"] in {1, 3, 5, "confirm"}
            for item in scorecard["items"]
        ))
        self.assertTrue(all(value == "confirm" for value in scorecard["start_gates"].values()))
        self.assertEqual("confirm", scorecard["items"][0]["scores"]["scale_readiness"])
        self.assertEqual("2", scorecard["schema_version"])
        self.assertTrue(all(
            all(item["score_reasons"].get(field) for field in assessment_generator.POC_SELECTION_SCORE_FIELDS)
            and item["data_next_action"]
            for item in scorecard["items"]
        ))
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))

        assessment["poc_selection_scorecard"]["items"][1]["use_case_id"] = "UC01"
        self.assertTrue(any(
            "poc_selection_scorecard" in error
            for error in assessment_generator.validate_assessment_payload(payload)
        ))

    def test_weighted_priority_is_input_order_independent_and_keeps_confirm_coverage(self) -> None:
        candidates = [
            {"use_case_no": str(index), "value": "low", "feasibility": "low",
             "data_readiness": "low", "scale": "low"}
            for index in range(1, 16)
        ]
        candidates[8].update({"value": "high", "feasibility": "high", "data_readiness": "high", "scale": "high"})
        candidates[4].update({"value": "high", "feasibility": "high", "data_readiness": "high", "scale": "medium"})
        candidates[11].update({"value": "high", "feasibility": "high", "data_readiness": "medium", "scale": "medium"})
        ranked = assessment_generator.normalize_front_candidate_priorities([dict(item) for item in candidates])
        reversed_ranked = assessment_generator.normalize_front_candidate_priorities([dict(item) for item in reversed(candidates)])
        self.assertEqual(
            {item["use_case_no"]: item["priority"] for item in ranked},
            {item["use_case_no"]: item["priority"] for item in reversed_ranked},
        )
        self.assertEqual({"9": "P1", "5": "P2", "12": "P3"}, {
            item["use_case_no"]: item["priority"] for item in ranked if item["priority"] != "Watch"
        })
        partial = assessment_generator.candidate_priority_evaluation({
            "value": "high", "feasibility": "confirm", "data_readiness": "confirm", "scale": "confirm",
        })
        self.assertEqual(5.0, partial["score"])
        self.assertEqual(0.4, partial["evaluation_coverage"])
        self.assertEqual(2.0, partial["ranking_index"])
        self.assertFalse(partial["eligible_for_priority"])

        complete = assessment_generator.candidate_priority_evaluation({
            "value": "high", "feasibility": "high", "data_readiness": "medium", "scale": "high",
        })
        self.assertEqual(4.6, complete["ranking_index"])
        self.assertTrue(complete["eligible_for_priority"])
        self.assertLess(partial["ranking_index"], complete["ranking_index"])

    def test_representative_themes_prefer_distinct_technical_modalities_after_scoring(self) -> None:
        candidates = [
            {"use_case_no": str(index), "value": "medium", "feasibility": "medium",
             "data_readiness": "low", "scale": "medium"}
            for index in range(1, 16)
        ]
        # 得点上位を予測系だけで埋められる状況でも、適格な別方式があれば
        # 代表3テーマには最適化・RAGを含める。
        candidates[0].update(value="high", feasibility="high", data_readiness="high", scale="high")
        candidates[1].update(value="high", feasibility="high", data_readiness="high", scale="high")
        candidates[2].update(value="high", feasibility="high", data_readiness="high", scale="medium")
        candidates[4].update(data_readiness="medium")
        candidates[12].update(data_readiness="medium")
        catalog = {
            str(index): {
                "use_case_no": str(index),
                "theme": (
                    "入出荷量予測" if index <= 3 else
                    "作業順の優先提案" if index == 5 else
                    "現場ナレッジ検索" if index == 13 else
                    f"業務支援{index}"
                ),
                "ai_technology": (
                    "機械学習・時系列予測" if index <= 3 else
                    "数理最適化" if index == 5 else
                    "RAG・生成AI" if index == 13 else
                    "生成AI"
                ),
            }
            for index in range(1, 16)
        }
        ranked = assessment_generator.normalize_front_candidate_priorities(
            candidates, catalog_by_no=catalog,
        )
        selected = sorted(
            (item for item in ranked if item["priority"] in {"P1", "P2", "P3"}),
            key=lambda item: item["priority"],
        )
        self.assertEqual(["1", "5", "13"], [item["use_case_no"] for item in selected])
        self.assertEqual(
            ["anomaly_ml", "optimization", "rag"],
            [item["technical_modality"] for item in selected],
        )
        self.assertTrue(all(
            item["priority_selection_method"] == "technology_portfolio_v1"
            for item in selected
        ))

    def test_priority_eligibility_requires_coverage_required_axes_and_data_three(self) -> None:
        scale_pending = assessment_generator.candidate_priority_evaluation({
            "value": "high", "feasibility": "high", "data_readiness": "medium", "scale": "confirm",
        })
        self.assertEqual(0.9, scale_pending["evaluation_coverage"])
        self.assertEqual(4.1, scale_pending["ranking_index"])
        self.assertTrue(scale_pending["eligible_for_priority"])

        data_missing = assessment_generator.candidate_priority_evaluation({
            "value": "high", "feasibility": "high", "data_readiness": "confirm", "scale": "high",
        })
        self.assertFalse(data_missing["eligible_for_priority"])
        self.assertTrue(any("data_readiness" in reason for reason in data_missing["eligibility_reasons"]))

        data_low = assessment_generator.candidate_priority_evaluation({
            "value": "high", "feasibility": "high", "data_readiness": "low", "scale": "high",
        })
        self.assertFalse(data_low["eligible_for_priority"])
        self.assertIn("data_readinessが3未満", data_low["eligibility_reasons"])

        candidates = [
            {"use_case_no": str(index), "value": "low", "feasibility": "low",
             "data_readiness": "low", "scale": "low"}
            for index in range(1, 16)
        ]
        candidates[0].update({
            "value": "high", "feasibility": "high", "data_readiness": "medium", "scale": "high",
        })
        ranked = assessment_generator.normalize_front_candidate_priorities(candidates)
        selected = [item for item in ranked if item["priority"] in {"P1", "P2", "P3"}]
        self.assertEqual(3, len(selected))
        self.assertEqual("eligible", next(item for item in selected if item["priority"] == "P1")["priority_selection_status"])
        self.assertTrue(any(item["priority_selection_status"] == "provisional" for item in selected))

    def test_v1_priority_and_scorecard_are_upgraded_without_breaking_review_json(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        decision_v2 = assessment_generator.priority_decision_for(
            assessment, assessment["consulting_front_matter"],
        )
        legacy_decision = {
            "schema_version": "1", "selection_method": "weighted_model",
            "weights": decision_v2["weights"], "tie_break_rule": "legacy",
            "items": [
                {
                    "use_case_id": item["use_case_id"], "use_case_no": item["use_case_no"],
                    "priority": item["priority"], "weighted_score": item["weighted_score"],
                    "evaluation_coverage": item["evaluation_coverage"],
                    "dimensions": {
                        field: {key: item["dimensions"][field][key] for key in ("level", "score", "weight")}
                        for field in assessment_generator.POC_SELECTION_SCORE_FIELDS
                    },
                    "basis": item["basis"],
                }
                for item in decision_v2["items"]
            ],
        }
        upgraded_decision = assessment_generator.normalize_poc_priority_decision(
            legacy_decision, assessment, assessment["consulting_front_matter"],
        )
        self.assertEqual("3", upgraded_decision["schema_version"])
        self.assertEqual("technology_portfolio_v1", upgraded_decision["selection_method"])
        self.assertTrue(all("ranking_index" in item for item in upgraded_decision["items"]))

        scorecard_v2 = assessment_generator.fallback_poc_selection_scorecard(
            assessment, assessment["consulting_front_matter"],
        )
        legacy_scorecard = json.loads(json.dumps(scorecard_v2, ensure_ascii=False))
        legacy_scorecard["schema_version"] = "1"
        for item in legacy_scorecard["items"]:
            item.pop("score_reasons")
            item.pop("data_next_action")
        upgraded_scorecard = assessment_generator.normalize_poc_selection_scorecard(
            legacy_scorecard, assessment, assessment["consulting_front_matter"],
        )
        self.assertEqual("2", upgraded_scorecard["schema_version"])
        self.assertTrue(all(item["data_next_action"] for item in upgraded_scorecard["items"]))

    def test_decision_contract_materializes_the_same_review_data_for_legacy_json(self) -> None:
        assessment = self.detailed_assessment()
        assessment.pop("poc_portfolio", None)
        assessment.pop("technical_proposal", None)
        contract = assessment_generator.materialize_assessment_decision_contract(assessment)

        self.assertIn("poc_selection_scorecard", contract)
        self.assertIn("poc_priority_decision", contract)
        self.assertIn("poc_start_readiness", contract)
        self.assertIn("poc_charters", contract)
        self.assertIn("multitenant_governance", contract)
        self.assertIn("business_value_model", contract)
        self.assertIn("technical_proposal", contract)
        self.assertEqual(["blocked", "blocked", "blocked"], [
            item["status"] for item in contract["poc_start_readiness"]["items"]
        ])
        self.assertTrue(all(item["open_gate_count"] > 0 for item in contract["poc_start_readiness"]["items"]))
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))

    def test_poc_decision_builder_preserves_structured_charters_and_governance(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        expected = {
            "poc_charters": assessment_generator.fallback_poc_charters(
                assessment, assessment["consulting_front_matter"],
            ),
            "multitenant_governance": assessment_generator.fallback_multitenant_governance(),
        }
        with patch.object(assessment_generator, "_response_json", return_value=expected) as mocked:
            result = assessment_generator.build_poc_decision_data(
                "会社名: Example株式会社\nサービス名: Example Service", assessment,
                front_matter=assessment["consulting_front_matter"], midterm_plan={"status": "disabled"},
                research_sources=[], client=object(), model_id="test-model",
            )
        self.assertEqual(1, mocked.call_count)
        assessment["poc_charters"] = result["poc_charters"]
        assessment["multitenant_governance"] = result["multitenant_governance"]
        self.assertEqual("modelled", result["generation_mode"])
        self.assertEqual(expected["poc_charters"], result["poc_charters"])
        self.assertEqual(3, len(result["poc_charters"]["charters"]))
        self.assertTrue(all(
            item["success_criteria"] == assessment_generator.POC_CONFIRM_VALUE
            for item in result["poc_charters"]["charters"]
        ))
        self.assertEqual(expected["multitenant_governance"], result["multitenant_governance"])

    def test_poc_logic_details_do_not_reuse_a_different_theme_by_position(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        assessment["poc_logic_details"][1]["theme"] = "別テーマ"

        details = assessment_generator.poc_logic_details_for(assessment)
        self.assertEqual(["UC01", "UC02", "UC03"], [item["use_case_id"] for item in details])
        self.assertEqual(["ユースケース1", "ユースケース2", "ユースケース3"], [item["theme"] for item in details])

    def test_unverified_quantitative_hypotheses_do_not_enable_benchmark_pages(self) -> None:
        assessment = self.detailed_assessment()
        assessment["executive_evidence"] = {
            "evidence_mode": "hypothesis",
            "benchmarks": [{"headline_metric": "5〜10%", "source_id": "H1"}],
            "value_scenarios": [{"benchmark": "5〜10%", "source_ids": ["H1"]}],
        }
        legacy_payload = {
            "format": assessment_generator.ASSESSMENT_JSON_FORMAT,
            "input": {"source_text": "会社名: Example株式会社\nサービス名: Example Service"},
            "assessment": assessment,
            "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []},
        }
        self.assertTrue(any("数値付きhypothesis" in error for error in assessment_generator.validate_assessment_payload(legacy_payload)))
        migrated = assessment_generator.migrate_legacy_payload_for_review(legacy_payload)
        frozen = assessment_generator.build_review_payload(
            migrated["assessment"], migrated["research"], source_file=None,
            source_text=legacy_payload["input"]["source_text"], preprocessing=None,
            cost_estimate=TEST_COST_ESTIMATE,
        )
        self.assertEqual("pre_poc", frozen["assessment"]["executive_evidence"]["evidence_mode"])
        self.assertEqual([], frozen["assessment"]["executive_evidence"]["benchmarks"])
        self.assertEqual([], frozen["assessment"]["executive_evidence"]["value_scenarios"])
        self.assertTrue(frozen["research"]["quantitative_claim_audit"])
        self.assertEqual([], assessment_generator.validate_assessment_payload(frozen, strict=True))

    def test_verified_quantitative_claim_requires_approved_source_and_excerpt(self) -> None:
        assessment = self.detailed_assessment()
        assessment["executive_evidence"] = {
            "evidence_mode": "external_verified",
            "strategic_premises": [
                {"label": "事業機会", "headline": "h1", "detail": "d1", "basis": "顧客入力"},
                {"label": "優先テーマ", "headline": "h2", "detail": "d2", "basis": "顧客入力"},
                {"label": "事業化条件", "headline": "h3", "detail": "d3", "basis": "PoC仮説"},
            ],
            "strategic_logic": [{"label": "a"}, {"label": "b"}, {"label": "c"}],
            "benchmarks": [{
                "claim_id": "Q01", "claim_status": "verified_external", "use_case": "問い合わせ支援",
                "headline_metric": "10%", "kpi": "工数", "detail": "公開資料の導入効果",
                "source_id": "R1", "source_url": "https://example.com/case", "source_title": "事例",
                "source_locator": "導入効果", "source_metric": "工数を10%削減",
                "evidence_excerpt": "工数を10%削減",
            }],
            "value_scenarios": [{"use_case": "問い合わせ支援", "benchmark": "10%", "formula": "工数×10%",
                                 "example": "", "poc_gate": "PoCで10%を確認", "source_ids": ["R1"]}],
        }
        payload = {
            "format": assessment_generator.ASSESSMENT_JSON_FORMAT,
            "assessment": assessment,
            "research": {
                "midterm_plan": {"status": "disabled"},
                "industry_sources": [{"id": "R1", "title": "事例", "url": "https://example.com/case", "excerpt": "工数を10%削減"}],
                "quantitative_evidence": {"approved_source_ids": ["R1"]},
            },
        }
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))
        payload["research"]["quantitative_evidence"]["approved_source_ids"] = []
        self.assertTrue(any("定量claim" in error for error in assessment_generator.validate_assessment_payload(payload)))

    def test_three_verified_effects_become_the_reproducible_quantitative_page_contract(self) -> None:
        evidence, research = self.verified_quantitative_fixture()
        design = assessment_generator.poc_measurement_design_from_verified_evidence(
            evidence, research, self.detailed_assessment(),
        )
        self.assertEqual("external_verified", design["status"])
        self.assertEqual(
            assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
            design["schema_version"],
        )
        self.assertEqual(["R1", "R2", "R3"], [item["source_id"] for item in design["items"]])
        research["poc_measurement_design_audit"] = {
            "status": "external_verified", "basis_type": "verified_external_benchmark",
            "claim_ids": ["Q01", "Q02", "Q03"],
            "approved_source_ids": ["R1", "R2", "R3"],
            "poc_bindings": [
                {field: str(item.get(field) or "") for field in ("priority", "use_case_id", "theme")}
                for item in design["items"]
            ],
            "reason": "3件の公開定量根拠を原文照合した。",
        }
        self.assertTrue(assessment_generator._poc_measurement_design_is_verifiable(design, research))

        assessment = self.detailed_assessment()
        assessment["executive_evidence"] = evidence
        assessment["poc_measurement_design"] = design
        payload = {
            "format": assessment_generator.ASSESSMENT_JSON_FORMAT,
            "assessment": assessment, "research": research,
        }
        self.assertFalse(any(
            "external_verifiedの定量効果" in error
            for error in assessment_generator.validate_assessment_payload(payload)
        ))
        assessment_generator.materialize_assessment_decision_contract(
            assessment, assessment["consulting_front_matter"],
        )
        frozen = assessment_generator.build_review_payload(
            assessment, research, source_file=None,
            source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None, cost_estimate=TEST_COST_ESTIMATE,
        )
        self.assertEqual([], assessment_generator.validate_assessment_payload(frozen, strict=True))

    def test_legacy_partial_effects_are_omitted_when_frozen_as_repro_v2(self) -> None:
        evidence, research = self.verified_quantitative_fixture()
        evidence["benchmarks"] = evidence["benchmarks"][:2]
        evidence["value_scenarios"] = evidence["value_scenarios"][:2]
        assessment = self.detailed_assessment()
        result = assessment_generator.apply_quantitative_research_outcome(
            assessment, research, evidence,
            llm_estimate=self.llm_quantitative_estimate_fixture(),
            llm_model_id="test-model",
            llm_generation_log=[{"attempt": 1, "status": "accepted"}],
            missing_evidence=["3件目の定量根拠が不足"],
        )
        self.assertEqual("llm_estimate", result["status"])
        self.assertEqual("decision_thresholds", assessment["poc_measurement_design"]["status"])
        self.assertEqual("partial", research["quantitative_evidence"]["status"])
        self.assertEqual(["R1", "R2"], research["quantitative_evidence"]["approved_source_ids"])
        self.assertFalse(research["quantitative_analysis"]["stopped_before_output"])
        self.assertEqual("llm_estimate", research["quantitative_analysis"]["final_display_mode"])
        self.assertEqual(3, research["quantitative_analysis"]["llm_estimate_count"])
        self.assertEqual("test-model", research["quantitative_analysis"]["llm_model_id"])
        self.assertEqual(["E01", "E02", "E03"], research["llm_quantitative_estimate"]["estimate_ids"])
        self.assertTrue(assessment_generator._llm_quantitative_estimate_is_consistent(
            assessment["poc_measurement_design"], research,
        ))
        self.assertEqual(2, len(assessment["executive_evidence"]["benchmarks"]))
        assessment_generator.materialize_assessment_decision_contract(
            assessment, assessment["consulting_front_matter"],
        )
        frozen = assessment_generator.build_review_payload(
            assessment, research, source_file=None,
            source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None, cost_estimate=TEST_COST_ESTIMATE,
        )
        self.assertEqual([], assessment_generator.validate_assessment_payload(frozen, strict=True))
        self.assertEqual("pre_poc", frozen["assessment"]["poc_measurement_design"]["status"])
        self.assertEqual(
            "measurement_design_only",
            frozen["research"]["quantitative_analysis"]["final_display_mode"],
        )
        self.assertTrue(
            frozen["research"]["quantitative_analysis"]["unsupported_claims_omitted"]
        )
        self.assertNotIn("llm_quantitative_estimate", frozen["research"])
        self.assertEqual(
            "unsupported_numeric_claims_omitted",
            frozen["research"]["legacy_quantitative_migration_audit"]["status"],
        )

    def test_zero_verified_effects_continue_with_llm_quantitative_estimates(self) -> None:
        assessment = self.detailed_assessment()
        research = {"midterm_plan": {"status": "disabled"}, "industry_sources": []}
        result = assessment_generator.apply_quantitative_research_outcome(
            assessment, research, {},
            llm_estimate=self.llm_quantitative_estimate_fixture(),
            llm_model_id="test-model",
            llm_generation_log=[{"attempt": 1, "status": "accepted"}],
            missing_evidence=["公開効果を確認できない"],
        )
        self.assertEqual("llm_estimate", result["status"])
        self.assertEqual("decision_thresholds", assessment["poc_measurement_design"]["status"])
        self.assertEqual([], research["quantitative_evidence"]["approved_source_ids"])
        self.assertFalse(research["quantitative_analysis"]["stopped_before_output"])
        self.assertEqual("llm_estimate", research["quantitative_analysis"]["final_display_mode"])
        self.assertEqual(3, research["quantitative_analysis"]["llm_estimate_count"])
        self.assertNotIn("executive_evidence", assessment)

    def test_invalid_poc_estimate_falls_back_to_measurement_design_without_stopping(self) -> None:
        assessment = self.detailed_assessment()
        research = {"midterm_plan": {"status": "disabled"}, "industry_sources": []}
        result = assessment_generator.apply_quantitative_research_outcome(
            assessment, research, {},
            llm_estimate={},
            llm_model_id="test-model",
            llm_generation_log=[{
                "attempt": 1, "status": "repair_required",
                "issues": ["3件の表示契約を満たさない"],
            }],
            missing_evidence=["公開効果を確認できない"],
        )

        self.assertEqual("measurement_design_only", result["status"])
        self.assertEqual("pre_poc", assessment["poc_measurement_design"]["status"])
        self.assertEqual(4, len(assessment["poc_measurement_design"]["items"]))
        self.assertEqual("pre_poc", research["quantitative_analysis"]["final_display_mode"])
        self.assertTrue(research["quantitative_analysis"]["unsupported_claims_omitted"])
        self.assertFalse(research["quantitative_analysis"]["stopped_before_output"])
        self.assertEqual(
            "unsupported_numeric_claims_omitted",
            research["poc_measurement_design_audit"]["generation_source"],
        )
        self.assertEqual("failed_omitted", research["llm_quantitative_estimate"]["status"])
        assessment_generator.materialize_assessment_decision_contract(
            assessment, assessment["consulting_front_matter"],
        )
        frozen = assessment_generator.build_review_payload(
            assessment, research, source_file=None,
            source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None, cost_estimate=TEST_COST_ESTIMATE,
        )
        self.assertEqual([], assessment_generator.validate_assessment_payload(frozen, strict=True))

    def test_three_verified_effects_skip_llm_estimates_and_log_external_origin(self) -> None:
        evidence, research = self.verified_quantitative_fixture()
        assessment = self.detailed_assessment()
        result = assessment_generator.apply_quantitative_research_outcome(
            assessment, research, evidence,
            llm_estimate=self.llm_quantitative_estimate_fixture(),
            llm_model_id="must-not-be-used",
        )
        self.assertEqual("external_verified", result["status"])
        self.assertEqual("external_verified", research["quantitative_analysis"]["final_display_mode"])
        self.assertEqual("not_used", research["llm_quantitative_estimate"]["status"])
        self.assertEqual("external_verified", research["poc_measurement_design_audit"]["generation_source"])

    def test_verified_effect_contract_rejects_partial_duplicate_and_mismatched_evidence(self) -> None:
        evidence, research = self.verified_quantitative_fixture()
        partial = json.loads(json.dumps(evidence, ensure_ascii=False))
        partial["benchmarks"].pop(); partial["value_scenarios"].pop()
        self.assertEqual({}, assessment_generator.poc_measurement_design_from_verified_evidence(
            partial, research, self.detailed_assessment(),
        ))

        duplicate = json.loads(json.dumps(evidence, ensure_ascii=False))
        duplicate["benchmarks"][2]["source_id"] = "R2"
        self.assertEqual({}, assessment_generator.poc_measurement_design_from_verified_evidence(
            duplicate, research, self.detailed_assessment(),
        ))

        mismatch = json.loads(json.dumps(evidence, ensure_ascii=False))
        mismatch["value_scenarios"][1]["benchmark"] = "15%"
        self.assertEqual({}, assessment_generator.poc_measurement_design_from_verified_evidence(
            mismatch, research, self.detailed_assessment(),
        ))

        duration_only = json.loads(json.dumps(evidence, ensure_ascii=False))
        duration_only["benchmarks"][0]["headline_metric"] = "12週間"
        duration_only["benchmarks"][0]["source_metric"] = "12週間"
        duration_only["benchmarks"][0]["evidence_excerpt"] = "12週間"
        duration_only["value_scenarios"][0]["benchmark"] = "12週間"
        duration_only["value_scenarios"][0]["formula"] = "12週間"
        duration_only["value_scenarios"][0]["poc_gate"] = "12週間"
        research_duration = json.loads(json.dumps(research, ensure_ascii=False))
        research_duration["industry_sources"][0]["excerpt"] = "12週間"
        self.assertEqual({}, assessment_generator.poc_measurement_design_from_verified_evidence(
            duration_only, research_duration, self.detailed_assessment(),
        ))

    def test_verified_effect_page_renders_sources_and_concrete_metrics(self) -> None:
        evidence, research = self.verified_quantitative_fixture()
        design = assessment_generator.poc_measurement_design_from_verified_evidence(
            evidence, research, self.detailed_assessment(),
        )
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_poc_quantitative_target_page(
            ppt, width, height, design, "Example Service", 6,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "verified-effects.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        text = self.pptx_slide_text(presentation.slides[0])
        for expected in ("期待できる定量効果", "10%", "20%", "30%", "公開事例の効果水準", "出典: 公開事例1 [R1]"):
            self.assertIn(expected, text)
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in self.pptx_text_font_sizes(presentation, include_footer=False)
        ))

    def test_frozen_v2_payload_detects_tampering(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_assessment_decision_contract(assessment)
        frozen = assessment_generator.build_review_payload(
            assessment, {"midterm_plan": {"status": "disabled"}, "industry_sources": []},
            source_file=None, source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None, cost_estimate=TEST_COST_ESTIMATE,
        )
        self.assertEqual([], assessment_generator.validate_assessment_payload(frozen, strict=True))
        frozen["assessment"]["company_name"] = "改ざん株式会社"
        self.assertTrue(any("provenance.assessment_sha256" in error
                            for error in assessment_generator.validate_assessment_payload(frozen, strict=True)))

    def test_from_frozen_json_does_not_reprice_or_rematerialize(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_assessment_decision_contract(assessment)
        frozen = assessment_generator.build_review_payload(
            assessment, {"midterm_plan": {"status": "disabled"}, "industry_sources": []},
            source_file=None, source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None, cost_estimate=TEST_COST_ESTIMATE,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory) / "frozen.json"
            output_path = Path(temporary_directory) / "copied.json"
            source_path.write_text(json.dumps(frozen, ensure_ascii=False), encoding="utf-8")
            argv = ["generate_assessment.py", "--from-json", str(source_path), "--json-only", "--output-json", str(output_path)]
            with patch.object(sys, "argv", argv), \
                 patch.object(assessment_generator, "build_poc_cost_estimate", side_effect=AssertionError("価格を再計算してはいけません")), \
                 patch.object(assessment_generator, "JSON_OUTPUT_DIR", Path(temporary_directory) / "json"), \
                 patch.object(assessment_generator, "RESEARCH_OUTPUT_DIR", Path(temporary_directory) / "research"):
                self.assertEqual(0, assessment_generator.main())
            self.assertEqual(frozen, json.loads(output_path.read_text(encoding="utf-8")))

    def test_llm_decision_thresholds_are_omitted_before_repro_v2_render(self) -> None:
        """Legacy LLM estimates may be read, but must not become v2 slide claims."""
        assessment = self.detailed_assessment()
        research = {"midterm_plan": {"status": "disabled"}, "industry_sources": []}
        outcome = assessment_generator.apply_quantitative_research_outcome(
            assessment,
            research,
            {},
            llm_estimate=self.llm_quantitative_estimate_fixture(),
            llm_model_id="test-model",
            llm_generation_log=[{"attempt": 1, "status": "accepted"}],
            missing_evidence=["公開効果を3件確認できない"],
        )
        self.assertEqual("llm_estimate", outcome["status"])
        self.assertEqual("decision_thresholds", assessment["poc_measurement_design"]["status"])
        assessment_generator.materialize_assessment_decision_contract(
            assessment, assessment["consulting_front_matter"],
        )
        frozen = assessment_generator.build_review_payload(
            assessment,
            research,
            source_file=None,
            source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None,
            cost_estimate=TEST_COST_ESTIMATE,
        )
        self.assertEqual([], assessment_generator.validate_assessment_payload(frozen, strict=True))
        self.assertEqual("pre_poc", frozen["assessment"]["poc_measurement_design"]["status"])
        self.assertNotIn("llm_quantitative_estimate", frozen["research"])

        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory) / "llm-estimate-frozen.json"
            output_path = Path(temporary_directory) / "llm-estimate-replayed.pptx"
            source_path.write_text(json.dumps(frozen, ensure_ascii=False), encoding="utf-8")
            with patch.object(sys, "argv", [
                "generate_assessment.py", "--from-json", str(source_path),
                "--output", str(output_path),
            ]), patch.object(
                assessment_generator, "analyze_assessment",
                side_effect=AssertionError("--from-jsonでAI分析を呼んではいけません"),
            ), patch.object(
                assessment_generator, "collect_midterm_plan",
                side_effect=AssertionError("--from-jsonでWeb調査を呼んではいけません"),
            ), patch.object(
                assessment_generator, "collect_industry_research",
                side_effect=AssertionError("--from-jsonで業界調査を呼んではいけません"),
            ), patch.object(
                assessment_generator, "build_poc_cost_estimate",
                side_effect=AssertionError("--from-jsonで費用を再計算してはいけません"),
            ):
                self.assertEqual(0, assessment_generator.main())
            presentation = Presentation(output_path)

        slide_texts = [self.pptx_slide_text(slide) for slide in presentation.slides]
        quantitative_page = next(
            text for text in slide_texts
            if "AI活用で期待できるビジネスインパクト" in text
        )
        self.assertNotIn("AI実装で目指す定量効果とPoC判定基準", "\n".join(slide_texts))
        self.assertIn("現時点では効果値を置かず", quantitative_page)
        for metric in ("7.3〜11.8%", "18.6〜23.4%", "12.2〜16.9%"): self.assertNotIn(metric, quantitative_page)

    def test_cli_is_pptx_only_and_rejects_pdf_output(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_assessment_decision_contract(assessment)
        frozen = assessment_generator.build_review_payload(
            assessment, {"midterm_plan": {"status": "disabled"}, "industry_sources": []},
            source_file=None, source_text="会社名: Example株式会社\nサービス名: Example Service",
            preprocessing=None, cost_estimate=TEST_COST_ESTIMATE,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory) / "frozen.json"
            source_path.write_text(json.dumps(frozen, ensure_ascii=False), encoding="utf-8")
            with patch.object(sys, "argv", [
                "generate_assessment.py", "--from-json", str(source_path),
                "--output", str(Path(temporary_directory) / "assessment.pdf"),
            ]), patch("sys.stderr", new_callable=io.StringIO):
                with self.assertRaises(SystemExit) as error:
                    assessment_generator.main()
            self.assertEqual(2, error.exception.code)

        for legacy_args in (("--format", "pdf"), ("--design", "graphical"), ("--detailed",)):
            with self.subTest(legacy_args=legacy_args), \
                 patch.object(sys, "argv", ["generate_assessment.py", *legacy_args]), \
                 patch("sys.stderr", new_callable=io.StringIO):
                with self.assertRaises(SystemExit) as error:
                    assessment_generator.main()
            self.assertEqual(2, error.exception.code)

    def test_technical_proposal_materializes_canonical_p1_with_stateful_nodes(self) -> None:
        assessment = self.detailed_assessment()
        assessment["use_cases"][0]["use_case"] = "在庫・出荷異常検知"
        assessment["poc_recommendations"][0]["theme"] = "在庫・出荷異常検知"
        assessment["consulting_front_matter"] = assessment_generator.fallback_consulting_front_matter(assessment)
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])

        proposal = assessment_generator.technical_proposal_for(assessment)
        self.assertEqual(assessment_generator.TECHNICAL_PROPOSAL_SCHEMA_VERSION, proposal["schema_version"])
        self.assertEqual("UC01", proposal["canonical_poc"]["use_case_id"])
        self.assertEqual("在庫・出荷異常検知", proposal["canonical_poc"]["theme"])
        self.assertEqual("anomaly_ml", proposal["modality"]["kind"])
        self.assertEqual(list(assessment_generator.TECHNICAL_PROPOSAL_NODE_IDS), [
            item["id"] for item in proposal["nodes"]
        ])
        self.assertEqual({"confirmed", "proposed", "verify"}, {
            item["state"] for item in proposal["nodes"] + proposal["controls"]
        })
        self.assertIn("イベント、API、定期連携", proposal["nodes"][1]["description"])
        self.assertIn("既存フローへ戻す", proposal["demo_scene"]["fallback_label"])

    def test_technical_modality_distinguishes_rag_ml_and_optimization(self) -> None:
        self.assertEqual("rag", assessment_generator.technical_modality_for({"theme": "問い合わせナレッジ検索"})["kind"])
        self.assertEqual("anomaly_ml", assessment_generator.technical_modality_for({"theme": "出荷異常検知"})["kind"])
        self.assertEqual("optimization", assessment_generator.technical_modality_for({"theme": "作業優先順位最適化"})["kind"])

    def test_technical_proposal_rejects_invalid_state_in_review_json(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(assessment, assessment["consulting_front_matter"])
        payload = {"format": assessment_generator.ASSESSMENT_JSON_FORMAT, "assessment": assessment,
                   "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []}}
        self.assertEqual([], assessment_generator.validate_assessment_payload(payload))
        assessment["technical_proposal"]["nodes"][1]["state"] = "unverified"
        self.assertTrue(any("technical_proposal" in error for error in assessment_generator.validate_assessment_payload(payload)))


    def test_ppt_canvas_writes_editable_text_and_shape_elements(self) -> None:
        page_width, page_height = landscape(A4)
        ppt = assessment_generator.PptCanvas(page_width, page_height)
        ppt.setFillColor(assessment_generator.colors.HexColor("#467653"))
        ppt.rect(0, 0, page_width, page_height, stroke=0, fill=1)
        ppt.setFillColor(assessment_generator.colors.white)
        ppt.setFont("AssessmentJapaneseBold", 22)
        ppt.drawString(18 * assessment_generator.mm, page_height - 34 * assessment_generator.mm, "編集可能な見出し")
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "editable.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        self.assertEqual(1, len(presentation.slides))
        self.assertTrue(any(shape.has_text_frame and "編集可能な見出し" in shape.text for shape in presentation.slides[0].shapes))

    def test_ppt_canvas_preserves_literal_text_before_saving(self) -> None:
        page_width, page_height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(page_width, page_height)
        ppt.setFillColor(assessment_generator.colors.black)
        ppt.setFont("AssessmentJapanese", 12)
        ppt.drawString(18 * assessment_generator.mm, 100 * assessment_generator.mm, "短文化前...途中…")
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "no-ellipsis.pptx"
            ppt.write(output_path)
            slide = Presentation(output_path).slides[0]
        text = self.pptx_slide_text(slide)
        self.assertEqual("短文化前...途中…", text)

    def test_pptx_number_badge_centers_a_plain_number_inside_its_circle(self) -> None:
        page_width, page_height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(page_width, page_height)
        assessment_generator.draw_number_badge(
            ppt, 30 * assessment_generator.mm, 80 * assessment_generator.mm, 5 * assessment_generator.mm,
            1, assessment_generator.colors.HexColor("#467653"),
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "number-badge.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        shapes = presentation.slides[0].shapes
        circle = next(shape for shape in shapes if shape.has_text_frame and not shape.text and shape.width == shape.height)
        label = next(shape for shape in shapes if shape.has_text_frame and shape.text == "1")
        self.assertLessEqual(circle.left, label.left)
        self.assertLessEqual(circle.top, label.top)
        self.assertLessEqual(label.left + label.width, circle.left + circle.width)
        self.assertLessEqual(label.top + label.height, circle.top + circle.height)
        self.assertEqual(assessment_generator.MIN_PPTX_FONT_SIZE, label.text_frame.paragraphs[0].runs[0].font.size.pt)

    def test_pptx_display_helpers_preserve_full_source_text(self) -> None:
        source_summary = "対象データを確認し、品質と権限を整備し、代表期間で比較可能な状態にします。"
        source_formula = "MRR改善率＝（AI有償機能付帯後12か月MRR－現状MRR）÷現状MRR＝8%"
        summary = assessment_generator._pptx_summary_text(source_summary, 22)
        formula = assessment_generator._pptx_formula_text(source_formula, "8%", 24)
        self.assertEqual(source_summary, summary)
        self.assertEqual(source_formula, formula)

    def test_pptx_reason_summary_preserves_full_reason(self) -> None:
        source = "顧客が優先する業務価値を代表データで検証するためです。"
        alignment_source = "物流課題を抱える企業への展開とWMSの運用効率化が接続するため。"
        self.assertEqual(source, assessment_generator._pptx_reason_summary(source, limit=24))
        self.assertEqual(alignment_source, assessment_generator._pptx_reason_summary(alignment_source, limit=28))

    def test_pptx_action_summary_preserves_full_action(self) -> None:
        source = "作業進捗と受注期限を対象倉庫単位で抽出し、利用者と担当者を確認する。"
        self.assertEqual(source, assessment_generator._pptx_action_summary(source, limit=18))

    def test_pptx_use_case_name_summary_preserves_source_copy(self) -> None:
        source = "問い合わせ・障害対応支援と状況説明・レポート作成"
        self.assertEqual(source, assessment_generator._pptx_use_case_name_summary(source))

    def test_pptx_business_challenge_summary_preserves_full_source(self) -> None:
        source = (
            "ロジザードZEROで同時に発生する作業・案件の優先順位を、締切・期限と処理能力を踏まえて"
            "一貫して提示し、担当者の判断時間と遅延リスクを抑えます。"
        )
        full = assessment_generator._pptx_business_challenge_summary(source, limit=68)
        compact = assessment_generator._pptx_business_challenge_summary(source, limit=30)
        self.assertEqual(source, full)
        self.assertEqual(source, compact)

        inquiry = assessment_generator._pptx_business_challenge_summary(
            "問い合わせに対して根拠を確認できる回答案を迅速に提示し、一次対応負荷を下げます。",
            limit=30,
        )
        self.assertEqual(
            "問い合わせに対して根拠を確認できる回答案を迅速に提示し、一次対応負荷を下げます。",
            inquiry,
        )

    def test_midterm_plan_number_badges_are_centered_in_their_circles(self) -> None:
        analysis = {
            "plan_summary": "公開計画とAI実装の接点を確認する。",
            "source": {"title": "中期経営計画", "url": "https://example.com/plan"},
            "ai_alignment": [{
                "plan_priority": "方針", "ai_role": "支援", "why_now": "検証", "related_use_case": "PoC",
            }],
            "caveat": "効果はPoCで検証する。",
        }
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_midterm_plan_alignment_page(ppt, width, height, analysis, 1)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "midterm-badge.pptx"
            ppt.write(output_path)
            slide = Presentation(output_path).slides[0]
        circle = next(shape for shape in slide.shapes if shape.has_text_frame and not shape.text and shape.width == shape.height)
        label = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == "1")
        self.assertLessEqual(circle.left, label.left)
        self.assertLessEqual(circle.top, label.top)
        self.assertLessEqual(label.left + label.width, circle.left + circle.width)
        self.assertLessEqual(label.top + label.height, circle.top + circle.height)

    def test_midterm_plan_source_title_uses_uppercase_ir_when_rendered(self) -> None:
        analysis = {
            "plan_summary": "公開計画とAI実装の接点を確認する。",
            "source": {
                "title": "中期経営計画│Ir情報│ロジザード株式会社",
                "url": "https://www.logizard.co.jp/ir/management-plan/",
            },
            "ai_alignment": [],
            "caveat": "効果はPoCで検証する。",
        }
        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)

        assessment_generator.draw_midterm_plan_alignment_page(
            ppt, width, height, analysis, 4,
        )

        slide_text = "\n".join(
            shape.text for shape in ppt.presentation.slides[0].shapes
            if shape.has_text_frame
        )
        self.assertIn(
            "出典：中期経営計画│IR情報│ロジザード株式会社",
            slide_text,
        )
        self.assertNotIn("Ir情報", slide_text)

    def test_midterm_management_target_text_is_vertically_centered_in_each_card(self) -> None:
        analysis = {
            "plan_summary": "公開計画とAI実装の接点を確認する。",
            "source": {"title": "中期経営計画", "url": "https://example.com/plan"},
            "management_targets": [
                {
                    "label": label, "target": target, "period": period,
                    "comparison": comparison,
                }
                for label, target, period, comparison in (
                    ("総売上高", "31.1億円", "2028年6月期", "2025年6月期比+43.0%"),
                    ("営業利益", "5.3億円", "2028年6月期", "2025年6月期比+31.8%"),
                    ("クラウドサービス 売上高", "23.9億円", "2028年6月期", "2025年6月期比+39.1%"),
                    ("月次経常収益", "2.09億円", "2028年6月", "2025年6月比+40.5%"),
                )
            ],
            "ai_alignment": [{
                "plan_priority": "方針", "ai_role": "支援", "why_now": "検証",
                "related_use_case": "PoC",
            }],
            "caveat": "効果はPoCで検証する。",
        }
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_midterm_plan_alignment_page(ppt, width, height, analysis, 4)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "midterm-target-cards.pptx"
            ppt.write(output_path)
            slide = Presentation(output_path).slides[0]

        target_text_shapes = sorted(
            (
                shape for shape in slide.shapes
                if shape.name.startswith("AI_ASSESS_MIDTERM_TARGET_TEXT_")
            ),
            key=lambda shape: shape.left,
        )
        self.assertEqual(4, len(target_text_shapes))
        cards = sorted(
            (
                shape for shape in slide.shapes
                if shape.name.startswith("AI_ASSESS_MIDTERM_TARGET_CARD_")
            ),
            key=lambda shape: shape.left,
        )
        self.assertEqual(4, len(cards))
        for index, (card, text_shape) in enumerate(zip(cards, target_text_shapes)):
            with self.subTest(index=index):
                self.assertEqual(MSO_ANCHOR.MIDDLE, text_shape.text_frame.vertical_anchor)
                self.assertGreaterEqual(text_shape.left, card.left)
                self.assertLessEqual(text_shape.left + text_shape.width, card.left + card.width)
                self.assertGreaterEqual(text_shape.top, card.top)
                self.assertLessEqual(text_shape.top + text_shape.height, card.top + card.height)
                top_margin = text_shape.top - card.top
                bottom_margin = card.top + card.height - text_shape.top - text_shape.height
                self.assertLessEqual(abs(top_margin - bottom_margin), 2)
                self.assertIn(analysis["management_targets"][index]["label"], text_shape.text)
                self.assertIn(analysis["management_targets"][index]["comparison"], text_shape.text)
                for paragraph in text_shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            self.assertIsNotNone(run.font.size)
                            self.assertGreaterEqual(run.font.size.pt, 12.0)
                            self.assertEqual("Meiryo UI", run.font.name)

    def test_midterm_plan_without_verified_targets_reclaims_the_empty_upper_band(self) -> None:
        """原典照合済み数値がない場合だけ、出典・方針・留意点を上へ詰める。"""
        base_analysis = {
            "plan_summary": "公開計画の重点方針とAI実装の接点を整理し、優先テーマと実行条件を確認する。",
            "source": {
                "title": "公開中期経営計画2028",
                "url": "https://example.com/investors/management-plan/2028/detail",
            },
            "ai_alignment": [
                {
                    "plan_priority": f"重点方針{index}",
                    "ai_role": f"AI支援{index}",
                    "why_now": f"実行条件{index}を確認する。",
                    "related_use_case": f"関連テーマ{index}",
                }
                for index in range(1, 4)
            ],
            "caveat": "数値効果は別途のPoCで同一条件比較する。",
        }

        def render(analysis: dict):
            width = assessment_generator.PPTX_WIDESCREEN_WIDTH
            height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
            ppt = assessment_generator.PptCanvas(width, height)
            assessment_generator.draw_midterm_plan_alignment_page(ppt, width, height, analysis, 4)
            return ppt.presentation.slides[0]

        without_targets = render(base_analysis)
        with_targets = render({
            **base_analysis,
            "management_targets": [{
                "label": "売上高", "target": "100億円", "period": "FY2028", "comparison": "FY2025比+20%",
            }],
        })

        def one(slide, text: str):
            return next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == text)

        no_source = one(without_targets, "参照した公開資料")
        target_source = one(with_targets, "経営目標（AI効果とは別）")
        no_headings = sorted(
            (shape for shape in without_targets.shapes
             if shape.has_text_frame and shape.text == "中期計画の重点方針"),
            key=lambda shape: shape.top,
        )
        target_headings = sorted(
            (shape for shape in with_targets.shapes
             if shape.has_text_frame and shape.text == "中期計画の重点方針"),
            key=lambda shape: shape.top,
        )
        no_caveat = one(without_targets, "判断上の留意点")
        target_caveat = one(with_targets, "判断上の留意点")

        self.assertEqual(3, len(no_headings))
        self.assertEqual(3, len(target_headings))
        self.assertLess(no_source.top, target_source.top)
        self.assertLess(no_headings[0].top, target_headings[0].top)
        self.assertIn(base_analysis["source"]["url"], self.pptx_slide_text(without_targets))
        self.assertNotRegex(self.pptx_slide_text(without_targets), r"(?:\.{3,}|…)")
        self.assertLess(no_caveat.top, target_caveat.top)
        self.assertGreaterEqual(
            target_source.top - no_source.top,
            round(20 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU),
        )
        for slide in (without_targets, with_targets):
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, slide.part.package.presentation_part.presentation.slide_width)
                self.assertLessEqual(shape.top + shape.height, slide.part.package.presentation_part.presentation.slide_height)
                if (
                    not shape.has_text_frame
                    or shape.name == "AI_ASSESS_SOURCE_NOTE"
                    or self.is_ten_point_chrome_shape(
                        shape, slide.part.package.presentation_part.presentation.slide_height,
                    )
                ):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size is not None:
                            self.assertGreaterEqual(
                                run.font.size.pt, assessment_generator.MIN_PPTX_FONT_SIZE,
                            )

    def test_default_pptx_is_a_compact_proposal_and_retains_each_priority_poc_page(self) -> None:
        assessment = self.detailed_assessment()
        assessment["poc_measurement_design"] = assessment_generator.normalize_poc_measurement_design({
            "schema_version": assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
            "status": "decision_thresholds",
            "items": [
                {
                    **v4_quantitative_item(
                        index, f"ユースケース{index}", metric,
                        outcome=outcome, kpi=kpi,
                    ),
                    "basis_type": "decision_threshold",
                }
                for index, (outcome, metric, kpi) in enumerate((
                    ("売上・粗利機会", "7.3〜11.8%", "候補提示採用率"),
                    ("業務生産性", "18.6〜23.4%", "判断リードタイム短縮率"),
                    ("損失・品質リスク", "12.2〜16.9%", "例外検知再現率改善"),
                ), 1)
            ],
        }, assessment)
        # 外部ベンチマークが存在しても、通常版では顧客固有の定量目標契約を優先する。
        assessment["executive_evidence"] = {
            "value_scenarios": [{"use_case": "外部事例", "benchmark": "50%"}],
        }
        rendering_cost_estimate = assessment_generator.cost_estimate_from_snapshot({
            "schema_version": "1", **TEST_COST_ESTIMATE,
        })
        self.assertIsNotNone(rendering_cost_estimate)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "default-assessment.pptx"
            assessment_generator.create_pptx(
                assessment, assessment_generator.DEFAULT_ARCHITECTURE_IMAGE,
                output_path, rendering_cost_estimate,
            )
            presentation = Presentation(output_path)
        slide_texts = [self.pptx_slide_text(slide) for slide in presentation.slides]
        self.assertEqual(
            assessment_generator.standard_document_page_count(assessment, include_cost_estimate=True),
            len(presentation.slides),
        )
        self.assertIn("OCI AI Use Case Assessment", slide_texts[1])
        self.assertIn("Example株式会社", slide_texts[1])
        self.assertIn("Example Service", slide_texts[1])
        self.assertIn("AIユースケースと成功条件を共同で設計", slide_texts[1])
        self.assertNotIn("価値仮説を定義", slide_texts[1])
        self.assertIn("最新動向とAI活用", slide_texts[2])
        self.assertNotIn("目次", slide_texts[1])
        self.assertFalse(any("アセスメント概要" in text for text in slide_texts))
        quantitative_page = next(text for text in slide_texts if "AI活用で期待できるビジネスインパクト" in text)
        for metric in ("7.3〜11.8%", "18.6〜23.4%", "12.2〜16.9%"): self.assertNotIn(metric, quantitative_page)
        self.assertTrue(any("15のAIユースケースから整理した AI技術アプローチの代表3テーマ" in text for text in slide_texts))
        self.assertEqual(
            3,
            sum(
                "AI技術テーマ" in text
                and "実装対象：データ準備から既存業務への返却まで" in text
                for text in slide_texts
            ),
        )
        self.assertTrue(any("PoC実施・支援概要" in text for text in slide_texts))
        font_sizes = self.pptx_text_font_sizes(presentation, include_footer=False)
        self.assertTrue(font_sizes)
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in font_sizes
        ))

    def test_default_intro_slide_is_native_dynamic_and_uses_the_reference_photo(self) -> None:
        assessment = self.detailed_assessment()
        assessment["service_use_case_groups"] = [{
            "service_name": "Example Service",
            "service_type": "業務支援SaaS",
            "use_cases": assessment["use_cases"],
        }]
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "native-intro.pptx"
            assessment_generator.create_pptx(
                assessment, assessment_generator.DEFAULT_ARCHITECTURE_IMAGE, output_path,
            )
            presentation = Presentation(output_path)

        slide = presentation.slides[1]
        slide_text = self.pptx_slide_text(slide)
        self.assertEqual("Blank", slide.slide_layout.name)
        self.assertEqual(1, len(presentation.slide_masters))
        self.assertIn("OCI AI Use Case Assessment", slide_text)
        self.assertIn("Example株式会社", slide_text)
        self.assertIn("Example Service", slide_text)
        self.assertIn("AIのビジネス価値を共同で構築", slide_text)
        self.assertIn("AIユースケースと成功条件を共同で設計", slide_text)
        self.assertIn("OCI環境でPoCを設計・検証", slide_text)
        self.assertIn("業務支援SaaSのAI活用を設計・実証", slide_text)
        self.assertNotIn("ISV様", slide_text)
        self.assertNotIn("仮説", slide_text)
        intro_copy = assessment_generator.assessment_intro_copy_for(assessment)
        implementation_sentence = (
            "OCI Generative AIを組み合わせ、実装方式と評価結果を次の判断へつなげます。"
        )
        self.assertEqual(1, intro_copy["blocks"][2]["description"].count(implementation_sentence))
        self.assertEqual(1, slide_text.count(implementation_sentence))

        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(1, len(pictures))
        self.assertEqual(presentation.slide_width // 2, pictures[0].left)
        self.assertEqual(0, pictures[0].top)
        self.assertAlmostEqual(presentation.slide_width / 2, pictures[0].width, delta=2)
        self.assertEqual(presentation.slide_height, pictures[0].height)

        footer = next(shape for shape in slide.shapes
                      if shape.has_text_frame and shape.text.startswith("Copyright ©"))
        page_number = next(shape for shape in slide.shapes
                           if shape.has_text_frame and shape.text.strip() == "2")
        self.assertEqual(
            [assessment_generator.PPTX_FOOTER_FONT_SIZE],
            [run.font.size.pt for paragraph in footer.text_frame.paragraphs
             for run in paragraph.runs if run.text.strip()],
        )
        self.assertEqual(
            [assessment_generator.PPTX_FOOTER_FONT_SIZE],
            [run.font.size.pt for paragraph in page_number.text_frame.paragraphs
             for run in paragraph.runs if run.text.strip()],
        )
        content_sizes = self.pptx_text_font_sizes(presentation, include_footer=False)
        self.assertTrue(content_sizes)
        self.assertTrue(all(size is not None and size >= 12 for size in content_sizes))

    def test_priority_poc_detail_pages_render_case_specific_oracle_implementation_designs(self) -> None:
        """前版の5ステップ＋3列表を保ち、P1〜P3の技術内容だけを差し替える。"""
        assessment = self.detailed_assessment()
        assessment_generator.materialize_assessment_decision_contract(assessment)
        details = [dict(item) for item in assessment_generator.poc_logic_details_for(assessment)]
        pocs = assessment_generator.priority_pocs_for(assessment)
        for index, detail in enumerate(details, 1):
            detail.update({
                "implementation_summary": f"固有データ{index}をOracle Databaseで処理し、固有結果{index}を既存画面へ返す。",
                "oracle_technologies": f"Oracle機能{index}／Autonomous AI Database",
                "database_objects": f"入力ビュー{index}／結果表{index}／監査ログ{index}",
                "output_interface": f"固有スコア{index}・根拠{index}をREST APIで返す。",
                "implementation_boundary": f"DBバージョンと遅延要件{index}に応じて既存DB内実行またはオフロードを選ぶ。",
                "control_design": f"権限{index}、例外フォールバック{index}、実行ログ{index}を管理する。",
            })

        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        for index, (detail, poc) in enumerate(zip(details, pocs), 1):
            assessment_generator.draw_poc_logic_detail_page(
                ppt, width, height, detail, poc, index, 8 + index,
            )
            if index < 3:
                ppt.showPage()

        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "priority-poc-technical-layout.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)

        self.assertEqual(3, len(presentation.slides))
        shared_labels = {
            "実装対象：データ準備から既存業務への返却まで",
            "技術観点", "Oracle Database／OCIでの実装", "既存システムとの接続・運用",
            "データ・AI処理", "DBオブジェクト", "配置・統制",
        }
        common_geometry = []
        for index, slide in enumerate(presentation.slides, 1):
            text = self.pptx_slide_text(slide)
            self.assertIn(f"AI技術テーマ {index}／3", text)
            for label in shared_labels:
                self.assertIn(label, text)
            for value in (
                f"固有データ{index}", f"Oracle機能{index}", f"入力ビュー{index}",
                f"固有スコア{index}", f"遅延要件{index}", f"フォールバック{index}",
            ):
                self.assertIn(value, text)
            self.assertIn("実装概要：", text)
            self.assertIn("対象データ：", text)
            self.assertIn("結果・根拠を", text)
            self.assertIn("担当者が確認し", text)
            self.assertNotIn("開始・見送り条件", text)
            footer = next(shape for shape in slide.shapes if shape.has_text_frame and "Copyright ©" in shape.text)
            self.assertTrue(all(
                shape.top + shape.height <= footer.top
                for shape in slide.shapes
                if not (shape.has_text_frame and (
                    "Copyright ©" in shape.text or shape.text.strip() == str(8 + index)
                ))
            ))
            common_geometry.append(tuple(
                (shape.left, shape.top, shape.width, shape.height)
                for label in shared_labels
                for shape in slide.shapes if shape.has_text_frame and shape.text == label
            ))
        self.assertEqual(1, len(set(common_geometry)))
        body_font_sizes = self.pptx_text_font_sizes(presentation, include_footer=False)
        self.assertTrue(body_font_sizes)
        self.assertTrue(all(size is not None and size >= 14 for size in body_font_sizes))

    def test_unverified_default_deck_omits_separate_business_impact_page(self) -> None:
        assessment = self.detailed_assessment()
        assessment["executive_evidence"] = {
            "evidence_mode": "pre_poc", "benchmarks": [], "value_scenarios": [],
        }
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "unverified-value-page.pptx"
            assessment_generator.create_pptx(
                assessment, assessment_generator.DEFAULT_ARCHITECTURE_IMAGE,
                output_path, None,
            )
            presentation = Presentation(output_path)

        slide_texts = [self.pptx_slide_text(slide) for slide in presentation.slides]
        kpi_page = next(text for text in slide_texts if "AI活用で期待できるビジネスインパクト" in text)
        self.assertIn("確認するKPI", kpi_page)
        self.assertIn("現時点では効果値を置かず", kpi_page)
        self.assertTrue(any("AIユースケース一覧" in text for text in slide_texts))

    def test_quantitative_hypothesis_is_frozen_as_poc_decision_thresholds(self) -> None:
        assessment = self.detailed_assessment()
        hypothesis_items = [
            v4_hypothesis_item(index, f"ユースケース{index}", metric)
            for index, metric in enumerate(("5〜10%", "18〜24%", "12〜17%"), 1)
        ]
        hypothesis = {
            "evidence_mode": "hypothesis",
            "management_targets": [
                {"label": "クラウドサービス売上", "target": "23.9億円（FY2028）",
                 "delta": "+39.1%", "source_id": "M1"},
            ],
            "benchmarks": [{
                **item,
                "use_case": item["business_outcome"],
                "source_title": "代表データによるPoC判定目標",
                "supporting_use_cases": [item["theme"]],
            } for item in hypothesis_items],
            "value_scenarios": [
                {
                    "priority": item["priority"], "use_case_id": item["use_case_id"],
                    "theme": item["theme"], "formula": item["formula"],
                    "poc_gate": item["poc_gate"],
                }
                for item in hypothesis_items
            ],
        }
        design = assessment_generator.poc_measurement_design_from_quantitative_hypothesis(
            hypothesis, assessment,
        )
        self.assertEqual("decision_thresholds", design["status"])
        self.assertEqual(3, len(design["items"]))
        self.assertTrue(all(item["basis_type"] == "decision_threshold" for item in design["items"]))
        self.assertEqual("5〜10%", design["items"][0]["headline_metric"])
        customer_text = json.dumps(design, ensure_ascii=False)
        self.assertNotIn("仮説", customer_text)
        self.assertNotIn("公開実績ではない", customer_text)

        assessment["poc_measurement_design"] = design
        assessment_generator.materialize_assessment_decision_contract(
            assessment, assessment["consulting_front_matter"],
        )
        self.assertEqual("decision_thresholds", assessment["poc_measurement_design"]["status"])
        self.assertEqual("5〜10%", assessment["poc_measurement_design"]["items"][0]["headline_metric"])

        payload = {
            "format": assessment_generator.ASSESSMENT_JSON_FORMAT,
            "assessment": assessment,
            "research": {"midterm_plan": {"status": "disabled"}, "industry_sources": []},
        }
        self.assertTrue(any(
            "poc_measurement_design_audit" in error
            for error in assessment_generator.validate_assessment_payload(payload)
        ))
        payload["research"]["poc_measurement_design_audit"] = {
            "status": "decision_thresholds",
            "basis_type": "decision_threshold",
            "reason": "顧客実績や公開導入効果ではなく、PoCの同一条件比較で採否を判断する目標値。",
        }
        self.assertFalse(any(
            "poc_measurement_design_audit" in error
            for error in assessment_generator.validate_assessment_payload(payload)
        ))

    def test_poc_decision_thresholds_reject_mismatched_values_or_missing_rationale(self) -> None:
        """大きく表示する数値は、算定式・合格条件・設定理由まで同じ契約にする。"""
        design = {
            "schema_version": assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
            "status": "decision_thresholds",
            "items": [{
                **v4_quantitative_item(index, f"ユースケース{index}", "18〜24%"),
                "basis_type": "decision_threshold",
            } for index in range(1, 4)],
        }
        self.assertEqual(
            "decision_thresholds",
            assessment_generator.normalize_poc_measurement_design(design)["status"],
        )
        mismatched = json.loads(json.dumps(design, ensure_ascii=False))
        mismatched["items"][1]["poc_gate"]["go_condition"] = "同一条件比較で10%を達成"
        self.assertEqual({}, assessment_generator.normalize_poc_measurement_design(mismatched))
        missing_rationale = json.loads(json.dumps(design, ensure_ascii=False))
        missing_rationale["items"][2]["target_rationale"] = ""
        self.assertEqual({}, assessment_generator.normalize_poc_measurement_design(missing_rationale))

    def test_poc_quantitative_target_page_renders_three_specific_targets_without_hypothesis_label(self) -> None:
        design = assessment_generator.normalize_poc_measurement_design({
            "schema_version": assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
            "status": "decision_thresholds",
            "lead": "対象サービスの優先PoCを、経営KPIへ接続して本番化を判断します。",
            "management_targets": [
                {"label": "クラウドサービス売上", "target": "23.9億円（FY2028）", "source_id": "M1"},
                {"label": "営業利益", "target": "5.3億円（FY2028）", "source_id": "M1"},
            ],
            "items": [{
                **v4_quantitative_item(
                    index, f"ユースケース{index}", metric, outcome=outcome, kpi=kpi,
                ),
                "basis_type": "decision_threshold",
            } for index, (outcome, metric, kpi) in enumerate((
                ("売上・粗利機会の拡大", "5〜10%", "候補提示採用率"),
                ("業務生産性の改善", "18〜24%", "判断リードタイム短縮率"),
                ("損失・品質リスクの抑制", "12〜17%", "例外検知再現率改善"),
            ), 1)],
        })
        # 本番契約では大数値を事業効果、別の数値を製品先行KPIとして表示する。
        for index, item in enumerate(design["items"], 1):
            business_layer = "provider_business_kpi" if index != 2 else "customer_outcome_kpi"
            business_owner = item[business_layer]["metric_owner"]
            item["display_kpi_layer"] = business_layer
            item[business_layer].update({
                "kpi": item["kpi"],
                "target": item["headline_metric"],
                "formula": item["formula"],
            })
            item["metric_owner"] = business_owner
            item["beneficiary"] = "provider" if business_layer == "provider_business_kpi" else "shared"
            item["attribution_level"] = item[business_layer]["attribution_level"]
            product_metric = ("65〜80%", "70〜85%", "60〜75%")[index - 1]
            item["product_kpi"].update({
                "kpi": f"製品先行KPI{index}",
                "target": product_metric,
                "formula": f"製品先行KPI{index}を同一条件で比較し、{product_metric}を判定する",
            })
            item["poc_gate"].update({
                "business_go_condition": f"事業効果で{item['headline_metric']}を達成",
                "leading_kpi_condition": f"製品先行KPIで{product_metric}を達成",
                "go_condition": f"事業効果{item['headline_metric']}、製品先行KPI{product_metric}をともに達成",
            })
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_poc_quantitative_target_page(
            ppt, width, height, design, "Example Service", 6,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "quantitative-targets.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        slide = presentation.slides[0]
        text = self.pptx_slide_text(slide)
        for expected in (
            "5〜10%", "18〜24%", "12〜17%", "事業効果Go", "23.9億円", "5.3億円",
            "P1｜ユースケース1", "顧客業務成果（AI寄与）", "提供者事業成果（AI寄与）",
            "製品先行KPI（提供者が直接管理）", "停止条件", "判定責任者", "要確認",
        ):
            self.assertIn(expected, text)
        self.assertNotIn("仮説", text)
        self.assertNotIn("補助確認", text)
        self.assertNotIn("…", text)
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in self.pptx_text_font_sizes(presentation, include_footer=False)
        ))
        card_width = round(
            ((assessment_generator.wide_content_width(width) - 12 * assessment_generator.mm) / 3)
            * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        card_height = round(113.5 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU)
        gate_height = round(31 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU)
        product_width = round(
            (((assessment_generator.wide_content_width(width) - 12 * assessment_generator.mm) / 3)
             - 8 * assessment_generator.mm)
            * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        product_height = round(19 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU)
        cards = [
            shape for shape in slide.shapes
            if abs(shape.width - card_width) <= 2 and abs(shape.height - card_height) <= 2
        ]
        gates = [
            shape for shape in slide.shapes
            if abs(shape.width - card_width) <= 2 and abs(shape.height - gate_height) <= 2
        ]
        product_bands = [
            shape for shape in slide.shapes
            if abs(shape.width - product_width) <= 2 and abs(shape.height - product_height) <= 2
        ]
        self.assertEqual(3, len(cards))
        self.assertEqual(3, len(gates))
        self.assertEqual(3, len(product_bands))
        cards.sort(key=lambda shape: shape.left)
        gates.sort(key=lambda shape: shape.left)
        product_bands.sort(key=lambda shape: shape.left)
        for card, gate, product_band in zip(cards, gates, product_bands):
            card_text = [
                shape for shape in slide.shapes
                if shape.has_text_frame and card.left <= shape.left < card.left + card.width
                and card.top <= shape.top < card.top + card.height
                and not (shape.text.startswith("Copyright") or shape.text.strip() == "6")
            ]
            self.assertTrue(card_text)
            self.assertTrue(all(shape.left + shape.width <= card.left + card.width for shape in card_text))
            product_contents = [
                next(shape for shape in card_text if "製品先行KPI（提供者が直接管理）" in shape.text),
                next(shape for shape in card_text if shape.text.strip() in {"65〜80%", "70〜85%", "60〜75%"}),
                next(shape for shape in card_text if shape.text.strip().startswith("製品先行KPI")
                     and "（提供者が直接管理）" not in shape.text),
            ]
            for content in product_contents:
                self.assertGreaterEqual(content.top, product_band.top)
                self.assertLessEqual(
                    content.top + content.height, product_band.top + product_band.height,
                    msg=f"product content overflow: {content.text!r}",
                )
            gate_contents = (
                next(shape for shape in card_text if "事業効果Go" in shape.text),
                next(shape for shape in card_text if "判定責任者" in shape.text),
                next(shape for shape in card_text if "停止条件" in shape.text),
            )
            for content in gate_contents:
                self.assertGreaterEqual(content.top, gate.top)
                self.assertLessEqual(content.top + content.height, gate.top + gate.height)
            ordered_gate_contents = sorted(gate_contents, key=lambda shape: shape.top)
            self.assertTrue(all(
                current.top + current.height <= following.top
                for current, following in zip(ordered_gate_contents, ordered_gate_contents[1:])
            ))
            business_kpi = next(
                shape for shape in card_text
                if shape.text.strip() in {"候補提示採用率", "判断リードタイム短縮率", "例外検知再現率改善"}
            )
            causal = next(shape for shape in card_text if shape.text.startswith("因果｜"))
            self.assertLessEqual(business_kpi.top + business_kpi.height, product_band.top)
            self.assertLessEqual(causal.top + causal.height, product_band.top)
            self.assertLessEqual(product_band.top + product_band.height, gate.top)
        footer = next(shape for shape in slide.shapes if shape.has_text_frame and "Copyright" in shape.text)
        self.assertTrue(all(
            shape.top + shape.height <= footer.top
            for shape in slide.shapes
            if not (
                shape.has_text_frame and (
                    "Copyright" in shape.text
                    or (shape.text.strip() == "6" and shape.top > presentation.slide_height * 0.9)
                )
            )
        ))

    def test_default_poc_selection_keeps_frozen_scores_internal_and_renders_qualitative_copy(self) -> None:
        assessment = self.detailed_assessment()
        candidates = assessment["consulting_front_matter"]["use_case_prioritization"]["candidates"]
        for candidate in candidates:
            candidate.update({
                "value": "low", "feasibility": "low", "data_readiness": "low", "scale": "low",
            })
        candidates[0].update({
            "value": "high", "feasibility": "high", "data_readiness": "medium", "scale": "high",
        })
        candidates[1].update({
            "value": "high", "feasibility": "medium", "data_readiness": "high", "scale": "medium",
        })
        candidates[2].update({
            "value": "high", "feasibility": "medium", "data_readiness": "medium", "scale": "medium",
        })
        assessment_generator.normalize_front_candidate_priorities(candidates)
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        frozen_decision = assessment_generator.materialize_poc_priority_decision(
            assessment, assessment["consulting_front_matter"],
        )
        self.assertEqual([4.6, 4.2, 3.8], [
            item["weighted_score"] for item in frozen_decision["items"]
            if item["priority"] in {"P1", "P2", "P3"}
        ])

        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_default_poc_selection_page(ppt, width, height, assessment, 8)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "priority-scorecard.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
            slide = presentation.slides[0]
        text = self.pptx_slide_text(slide)

        # P7の横並びカードは、後続の3枚のPoC詳細ページと同じテーマ対応色を使う。
        # この色は順位や評価の強弱ではなく、ページをまたぐ視覚的な対応付けである。
        card_bar_height = int(4.5 * assessment_generator.mm * 12700)
        theme_bars = sorted(
            (
                shape for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and not shape.text.strip()
                and shape.width > presentation.slide_width * 0.25
                and abs(shape.height - card_bar_height) <= 1
            ),
            key=lambda shape: shape.left,
        )
        self.assertEqual(3, len(theme_bars))
        self.assertEqual(
            [value.lstrip("#") for value in assessment_generator.POC_THEME_ACCENT_HEX],
            [str(shape.fill.fore_color.rgb).upper() for shape in theme_bars],
        )
        for heading in ("代表とする理由", "PoCで確かめること", "準備のポイント"):
            heading_shapes = sorted(
                (
                    shape for shape in slide.shapes
                    if shape.has_text_frame and shape.text.strip() == heading
                ),
                key=lambda shape: shape.left,
            )
            self.assertEqual(3, len(heading_shapes))
            heading_colors = []
            for shape in heading_shapes:
                runs = [
                    run for paragraph in shape.text_frame.paragraphs
                    for run in paragraph.runs if run.text.strip()
                ]
                self.assertTrue(runs)
                heading_colors.append(str(runs[0].font.color.rgb).upper())
            self.assertEqual(
                [value.lstrip("#") for value in assessment_generator.POC_THEME_ACCENT_HEX],
                heading_colors,
            )

        for priority_poc in assessment_generator.priority_pocs_for(
            assessment, assessment["consulting_front_matter"],
        ):
            self.assertIn(priority_poc["theme"], text)

        for phrase in (
            "15のAIユースケースから整理した AI技術アプローチの代表3テーマ", "整理観点", "業務価値",
            "既存機能との差分", "データ準備", "比較検証の成立性",
            "代表とする理由", "PoCで確かめること", "準備のポイント",
        ):
            self.assertIn(phrase, text)
        for phrase in (
            "4.6", "4.2", "3.8", "順位指数", "既知平均", "採点率",
            "評価：5=", "事業価値40%", "開始保留", "未充足",
            "責任者：", "データ準備：", "5へ上げる条件：", "P1｜", "順位",
            "5候補", "次段階", "保留",
        ):
            self.assertNotIn(phrase, text)
        sizes = self.pptx_text_font_sizes(presentation, include_footer=False)
        self.assertTrue(sizes)
        self.assertTrue(all(size is not None and size >= 14.0 for size in sizes))
        self.assertNotRegex(text, r"(?:\.{3,}|…)")

    def test_default_poc_selection_maximum_copy_preserves_full_text(self) -> None:
        """長文でも内容を削らず、選定ページの編集可能な14pt本文として残す。"""
        assessment = self.detailed_assessment()
        long_theme = "複数拠点の問い合わせ・障害状況・対応履歴を統合する根拠付き回答支援"
        long_reason = "事業価値と横展開性を両立し、代表業務の比較ログから判定品質と運用負荷を同時に測定できるため選定する。"
        long_data_reason = "対象履歴は複数テナントと権限区分を含むため、匿名化・欠損・期間の充足を比較条件ごとに確認する。"
        long_data_action = "代表範囲のデータ所有者、利用権限、欠損率、比較期間をPoC開始前に合意する。"
        pocs = [
            {
                "priority": f"P{index}", "use_case_id": f"UC{index:02d}",
                "theme": f"{long_theme}{index}", "reason": long_reason,
                "first_step": long_data_action,
            }
            for index in range(1, 4)
        ]
        dimensions = {
            key: {"score": score, "reason": long_data_reason}
            for key, score in (("business_value", 5), ("feasibility", 4),
                               ("data_readiness", 3), ("scale_readiness", 5))
        }
        decision = {
            "weights": {"business_value": .4, "feasibility": .3,
                        "data_readiness": .2, "scale_readiness": .1},
            "items": [
                {"priority": f"P{index}", "use_case_id": f"UC{index:02d}",
                 "weighted_score": 4.2, "ranking_index": 4.1,
                 "evaluation_coverage": 1.0, "eligible_for_priority": True,
                 "dimensions": dimensions}
                for index in range(1, 4)
            ],
        }
        scorecard = {"items": [
            {"use_case_id": f"UC{index:02d}",
             "score_reasons": {
                 "business_value": long_reason,
                 "feasibility": long_data_reason,
             },
             "data_next_action": long_data_action}
            for index in range(1, 4)
        ]}

        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        with patch("ai_assess_runtime.presentation.consulting_front_matter_for", return_value={}), \
             patch("ai_assess_runtime.presentation.normalize_poc_priority_decision", return_value=decision), \
             patch("ai_assess_runtime.presentation.poc_selection_scorecard_for", return_value=scorecard), \
             patch("ai_assess_runtime.presentation.priority_pocs_for", return_value=pocs):
            assessment_generator.draw_default_poc_selection_page(ppt, width, height, assessment, 8)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "priority-scorecard-maximum-copy.pptx"
            ppt.write(output_path)
            slide = Presentation(output_path).slides[0]

        slide_text = self.pptx_slide_text(slide)
        for index in range(1, 4):
            self.assertIn(f"{long_theme}{index}", slide_text)
        for source in (long_reason, long_data_reason, long_data_action):
            self.assertIn(source, slide_text)
        for phrase in ("順位指数", "採点率", "開始保留", "未充足", "P1｜"):
            self.assertNotIn(phrase, slide_text)
        sizes = [
            run.font.size.pt
            for shape in slide.shapes
            if shape.has_text_frame
            and not self.is_ten_point_chrome_shape(
                shape, slide.part.package.presentation_part.presentation.slide_height,
            )
            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text.strip()
        ]
        self.assertTrue(sizes)
        self.assertTrue(all(size >= 14.0 for size in sizes))

        # 各カードの長文本文は、次の見出しやカード下端へ侵入しない。
        cards = sorted(
            (
                shape for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and not shape.text.strip()
                and shape.width > slide.part.package.presentation_part.presentation.slide_width * 0.25
                and shape.height > slide.part.package.presentation_part.presentation.slide_height * 0.45
            ),
            key=lambda shape: shape.left,
        )
        self.assertEqual(3, len(cards))
        for card in cards:
            contained = [
                shape for shape in slide.shapes
                if shape.has_text_frame
                and card.left <= shape.left < card.left + card.width
            ]
            bodies = {
                value: next(shape for shape in contained if shape.text.strip() == value)
                for value in (long_reason, long_data_reason, long_data_action)
            }
            headings = {
                value: sorted(
                    (shape for shape in contained if shape.text.strip() == value),
                    key=lambda shape: shape.top,
                )
                for value in ("PoCで確かめること", "準備のポイント")
            }
            self.assertLessEqual(
                bodies[long_reason].top + bodies[long_reason].height,
                headings["PoCで確かめること"][0].top,
            )
            self.assertLessEqual(
                bodies[long_data_reason].top + bodies[long_data_reason].height,
                headings["準備のポイント"][0].top,
            )
            self.assertLessEqual(
                bodies[long_data_action].top + bodies[long_data_action].height,
                card.top + card.height,
            )

    def test_fallback_poc_logic_is_specific_to_anomaly_and_priority_optimization(self) -> None:
        assessment = self.detailed_assessment()
        assessment["use_cases"][:3] = [
            {"no": 1, "use_case": "在庫・出荷異常検知", "ai_technology": "機械学習・異常検知",
             "description": "在庫・出荷履歴から例外候補を検出する。"},
            {"no": 2, "use_case": "出荷作業優先順位支援", "ai_technology": "機械学習・スコアリング・最適化",
             "description": "締切と進捗から作業順序を提案する。"},
            {"no": 3, "use_case": "問い合わせ回答支援", "ai_technology": "生成AI・RAG",
             "description": "根拠付き回答案を提示する。"},
        ]
        assessment["poc_recommendations"] = [
            {"theme": case["use_case"], "use_case_no": str(case["no"]),
             "reason": "代表業務で検証する。", "first_step": "対象データを確認する。"}
            for case in assessment["use_cases"][:3]
        ]
        assessment["consulting_front_matter"] = assessment_generator.fallback_consulting_front_matter(assessment)
        candidates = assessment["consulting_front_matter"]["use_case_prioritization"]["candidates"]
        for index, candidate in enumerate(candidates):
            candidate.update(
                value="high" if index < 3 else "medium",
                feasibility="high" if index < 3 else "medium",
                data_readiness="high" if index < 3 else "low",
                scale="high" if index < 3 else "medium",
            )
        assessment.pop("poc_logic_details", None)
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        details = assessment_generator.poc_logic_details_for(assessment)
        anomaly = next(item for item in details if "異常検知" in item["theme"])
        priority = next(item for item in details if "優先順位" in item["theme"])
        query = next(item for item in details if "問い合わせ" in item["theme"])
        anomaly_text = json.dumps(anomaly, ensure_ascii=False)
        priority_text = json.dumps(priority, ensure_ascii=False)

        self.assertIn("異常", anomaly_text)
        self.assertRegex(anomaly_text, r"Oracle Machine Learning|OML")
        self.assertNotIn("金額", anomaly["target_data"])
        for keyword in ("締切", "進捗", "制約"):
            self.assertIn(keyword, priority_text)
        self.assertRegex(priority_text, r"Oracle Machine Learning|OML")
        query_step = assessment_generator._pptx_summary_text(
            query["processing_steps"][1]["description"], 120,
        )
        priority_step = assessment_generator._pptx_summary_text(
            priority["processing_steps"][1]["description"], 120,
        )
        self.assertIn("OCI Generative AI", query_step)
        self.assertIn("回答案", query_step)
        self.assertIn("制約を反映", priority_step)
        self.assertIn("順位候補を算出します。", priority_step)
        self.assertNotIn("抽出・要約", priority_text)

    def test_fallback_poc_logic_distinguishes_reporting_recommendation_and_development(self) -> None:
        assessment = self.detailed_assessment()
        themes = [
            "AIレポート／日報自動作成の高度化",
            "顧客・購買データに基づく販促／接客提案",
            "開発・保守のAIアシスタント",
        ]
        assessment["use_cases"][:3] = [
            {"no": index, "use_case_id": f"UC{index:02d}", "use_case": theme,
             "ai_technology": "生成AI", "description": theme}
            for index, theme in enumerate(themes, 1)
        ]
        assessment["poc_recommendations"] = [
            {"priority": f"P{index}", "use_case_id": f"UC{index:02d}",
             "use_case_no": str(index), "theme": theme,
             "reason": "代表業務で検証する。", "first_step": "対象データを確認する。"}
            for index, theme in enumerate(themes, 1)
        ]
        assessment.pop("poc_logic_details", None)
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment_generator.fallback_consulting_front_matter(assessment),
        )
        details = assessment_generator.poc_logic_details_for(assessment)
        combined = [json.dumps(item, ensure_ascii=False) for item in details]
        self.assertIn("レポート案", combined[0])
        self.assertIn("対象顧客群", combined[1])
        self.assertIn("CI", combined[2])
        self.assertEqual(3, len({tuple(step["label"] for step in item["processing_steps"]) for item in details}))
        proposal = assessment_generator.build_default_technical_proposal(assessment)
        self.assertNotIn("WMS", proposal["nodes"][0]["label"])
        self.assertIn(assessment["service_name"], proposal["nodes"][0]["label"])

    def test_final_poc_logic_generation_uses_confirmed_three_bindings(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        rows = []
        step_sets = [
            [("販売実績受付", "店舗別売上と商品マスタを日次で集計する。"), ("変動要因要約", "前年差から主要因を抽出して報告文を生成する。"), ("店長承認", "日報画面で店長が根拠を確認して配信する。")],
            [("顧客群抽出", "購買履歴と同意状態から施策対象を抽出する。"), ("販促案作成", "商品条件に合う提案と除外理由を生成する。"), ("担当者承認", "販促画面で対象と文面を承認して反応を記録する。")],
            [("コード索引化", "仕様とリポジトリを権限別に検索可能にする。"), ("変更案生成", "関連箇所とテスト観点を根拠付きで提示する。"), ("CI確認", "開発者レビューとCIを通して採否を記録する。")],
        ]
        for index, poc in enumerate(assessment_generator.priority_pocs_for(assessment), 1):
            rows.append({
                "priority": f"P{index}", "use_case_id": poc["use_case_id"], "theme": poc["theme"],
                "business_challenge": f"固有課題{index}", "target_data": f"固有データ{index}",
                "processing_steps": [
                    {"label": label, "description": description}
                    for label, description in step_sets[index - 1]
                ],
                "oci_roles": f"OCIサービスの役割{index}",
                "business_integration": f"既存業務への連携{index}",
                "validation_plan": f"テーマ固有KPI{index}で比較する。",
                "design_notes": f"権限と例外条件{index}を確認する。",
            })
        model_rows = copy.deepcopy(rows)
        for index, row in enumerate(model_rows, 1):
            row["use_case_id"] = f"uc-{index}"
            row["theme"] = f"表記揺れ{index}"
        with patch.object(
            assessment_generator, "_response_json", return_value={"poc_logic_details": model_rows},
        ) as response:
            result = assessment_generator.build_final_poc_logic_details(
                "顧客入力", assessment, client=object(), model_id="test-model",
            )
        self.assertEqual([row["use_case_id"] for row in rows], [row["use_case_id"] for row in result])
        self.assertEqual(1, response.call_count)
        self.assertIn("確定済み優先PoC", response.call_args.kwargs["prompt"])
        self.assertIn("oracle_technologies", response.call_args.kwargs["prompt"])
        self.assertIn("database_objects", response.call_args.kwargs["prompt"])
        self.assertIn("implementation_boundary", response.call_args.kwargs["prompt"])

    def test_poc_technical_design_selects_oracle_pattern_from_use_case_modality(self) -> None:
        base = {
            "target_data": "対象業務の履歴と確定結果",
            "business_integration": "スコアと根拠を既存画面へ返す。",
            "design_notes": "権限と監査ログを管理する。",
            "processing_steps": [{"label": "入力", "description": "準備"}] * 3,
        }
        forecast = assessment_generator.poc_technical_design_for(
            {**base, "theme": "需要予測"}, {"theme": "需要予測"},
        )
        knowledge = assessment_generator.poc_technical_design_for(
            {**base, "theme": "規程の根拠付き検索"}, {"theme": "規程の根拠付き検索"},
        )
        ranking = assessment_generator.poc_technical_design_for(
            {**base, "theme": "案件の優先順位付け"}, {"theme": "案件の優先順位付け"},
        )
        list_shaped = assessment_generator.poc_technical_design_for(
            {
                **base, "theme": "問い合わせ支援",
                "oracle_technologies": ["Oracle AI Vector Search", "OCI Generative AI"],
                "database_objects": ["文書表", "VECTOR索引", "回答監査ログ"],
            },
            {"theme": "問い合わせ支援"},
        )
        self.assertIn("Oracle Machine Learning", forecast["oracle_technologies"])
        self.assertIn("Oracle AI Vector Search", knowledge["oracle_technologies"])
        self.assertIn("制約", ranking["database_objects"])
        self.assertEqual(
            "Oracle AI Vector Search／OCI Generative AI",
            list_shaped["oracle_technologies"],
        )
        self.assertEqual("文書表／VECTOR索引／回答監査ログ", list_shaped["database_objects"])
        self.assertNotIn("[", list_shaped["oracle_technologies"])
        physical_name_fallback = assessment_generator.poc_technical_design_for(
            {
                **base, "theme": "需要予測",
                "database_objects": ["V_ORDER_EVENT", "T_MODEL_SCORE", "T_AUDIT_LOG"],
            },
            {"theme": "需要予測"},
        )
        self.assertNotIn("V_ORDER_EVENT", physical_name_fallback["database_objects"])
        self.assertIn("特徴量ビュー", physical_name_fallback["database_objects"])
        for design in (forecast, knowledge, ranking):
            self.assertIn("既存Oracle Database内", design["implementation_boundary"])
            self.assertIn("Database Linkでの参照", design["implementation_boundary"])
            self.assertIn("API連携", design["implementation_boundary"])

    def test_final_pptx_is_widescreen_and_body_is_at_least_twelve_point(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_assessment_decision_contract(assessment)
        rendering_cost_estimate = assessment_generator.cost_estimate_from_snapshot({
            "schema_version": "1", **TEST_COST_ESTIMATE,
        })
        self.assertIsNotNone(rendering_cost_estimate)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "assessment.pptx"
            assessment_generator.create_pptx(
                assessment, assessment_generator.DEFAULT_ARCHITECTURE_IMAGE,
                output_path, rendering_cost_estimate,
            )
            presentation = Presentation(output_path)

        self.assertEqual(
            round(assessment_generator.PPTX_WIDESCREEN_WIDTH * assessment_generator.PptCanvas.POINT_TO_EMU),
            presentation.slide_width,
        )
        self.assertEqual(
            round(assessment_generator.PPTX_WIDESCREEN_HEIGHT * assessment_generator.PptCanvas.POINT_TO_EMU),
            presentation.slide_height,
        )
        self.assertEqual(9 * presentation.slide_width, 16 * presentation.slide_height)
        self.assertEqual(
            assessment_generator.standard_document_page_count(assessment, include_cost_estimate=True),
            len(presentation.slides),
        )
        # 固定のADB構成図は、編集可能な技術提案図を置き換えず、概算費用の直前に
        # 独立した参照ページとして残す。
        slide_texts = [self.pptx_slide_text(slide) for slide in presentation.slides]
        self.assertNotRegex("\n".join(slide_texts), r"(?:\.{3,}|…)")
        self.assertIn("OCI AI Use Case Assessment", slide_texts[1])
        self.assertNotIn("目次", "\n".join(slide_texts))
        self.assertNotIn("ISV様の新規AIサービス構築をご支援", slide_texts[1])
        fixed_architecture_index = next(
            index for index, text in enumerate(slide_texts)
            if "代表3テーマを実現するOCI構成" in text
        )
        self.assertIn("PoCにおけるOCI概算費用", slide_texts[fixed_architecture_index + 1])
        self.assertTrue(any(
            shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            for shape in presentation.slides[fixed_architecture_index].shapes
        ))
        fixed_architecture_text = slide_texts[fixed_architecture_index]
        self.assertIn("代表技術テーマごとの実現方法", fixed_architecture_text)
        for poc in assessment_generator.priority_pocs_for(assessment):
            self.assertIn(str(poc["theme"]), fixed_architecture_text)
        font_sizes = self.pptx_text_font_sizes(presentation, include_footer=False)
        self.assertTrue(font_sizes)
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in font_sizes
        ))
        for index, slide in enumerate(presentation.slides, 1):
            footer_runs = [
                run for shape in slide.shapes if shape.has_text_frame
                and ("Copyright ©" in shape.text or (
                    shape.text.strip() == str(index)
                    and shape.top > presentation.slide_height * 0.9
                ))
                for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text.strip()
            ]
            # 表紙・クロージングはブランドビジュアルのみでフッターを置かない。
            # 共通フッターはブランド上の補助情報として10ptに固定する。
            if not footer_runs:
                self.assertIn(index, {1, len(presentation.slides)})
                continue
            self.assertTrue(all(run.font.size.pt == assessment_generator.PPTX_FOOTER_FONT_SIZE for run in footer_runs))
        self.assertEqual({"Meiryo UI"}, set(self.pptx_text_font_names(presentation)))
        header_text = "OCI AI USE CASE ASSESSMENT"
        header_runs = [
            run
            for slide in presentation.slides
            for shape in slide.shapes if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.text.strip() == header_text
        ]
        self.assertTrue(header_runs)
        self.assertTrue(all(run.font.size.pt == 10.0 for run in header_runs))
        self.assertNotIn(
            "根拠区分:",
            "\n".join(self.pptx_slide_text(slide) for slide in presentation.slides),
        )
        # スライド背景はPowerPointの背景設定で管理し、編集の妨げになる全画面の
        # 矩形図形を出力しない（写真を全面配置する表紙は対象外）。
        full_slide_rectangles = [
            shape
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and shape.left == 0 and shape.top == 0
            and shape.width == presentation.slide_width and shape.height == presentation.slide_height
        ]
        self.assertEqual([], full_slide_rectangles)

    def test_business_value_pages_share_the_same_three_card_accents(self) -> None:
        self.assertEqual(("#C74634", "#367A9B", "#467653"), assessment_generator.BUSINESS_VALUE_CARD_ACCENT_HEX)

    def test_strategic_logic_rejects_generic_cards_without_causal_fields(self) -> None:
        self.assertEqual([], assessment_generator.normalize_strategic_logic([
            {"label": "要請", "headline": "差別化する", "detail": "AIで効果を測定する。"},
            {"label": "基盤", "headline": "データを活用", "detail": "業務を支援する。"},
            {"label": "展開", "headline": "事業化する", "detail": "PoCで確認する。"},
        ]))

    def test_strategic_logic_builds_explicit_causal_labels(self) -> None:
        cards = assessment_generator.normalize_strategic_logic([
            {"label": "要請", "headline": "成長を支援", "fact": "計画は成長を掲げる。", "decision_implication": "AI投資を優先する。"},
            {"label": "基盤", "headline": "判断を変える", "current_constraint": "人手確認が残る。", "ai_decision_change": "取引データから優先順位を示す。"},
            {"label": "展開", "headline": "事業化を判断", "delay_risk": "対応遅れで機会を失う。", "proof_conditions": "効果と連携負荷を検証する。"},
        ])
        self.assertEqual(3, len(cards))
        self.assertIn("【AIで変える判断】", cards[1]["detail"])
        self.assertIn("【遅延リスク】", cards[2]["detail"])


    def test_non_plan_premises_always_use_distinct_decision_labels(self) -> None:
        premises = assessment_generator.normalize_strategic_premises([
            {"label": "意思決定の前提", "headline": "事業価値", "detail": "業務を支援する。", "basis": "顧客入力"},
            {"label": "意思決定の前提", "headline": "優先業務", "detail": "データを活用する。", "basis": "顧客入力"},
            {"label": "意思決定の前提", "headline": "展開条件", "detail": "PoCで確認する。", "basis": "PoC仮説"},
        ])
        self.assertEqual(["事業機会", "優先テーマ", "事業化条件"], [item["label"] for item in premises])


    def test_use_case_list_keeps_short_fifteen_candidates_on_one_widescreen_slide(self) -> None:
        full_description = "入出荷履歴を分析し、優先対応を提案する。"
        cases = [
            {
                "no": index, "use_case_id": f"UC{index:02d}", "use_case": f"候補{index}",
                "ai_technology": "生成AI・RAG・異常検知", "description": full_description,
            }
            for index in range(1, 16)
        ]
        group = {"service_name": "Example", "service_type": "SaaS", "use_cases": cases}
        chunks = assessment_generator.use_case_list_chunks(group)
        self.assertEqual([15], [len(chunk) for chunk in chunks])
        self.assertEqual(list(range(1, 16)), [case["no"] for chunk in chunks for case in chunk])
        self.assertEqual(1, len(chunks))
        displayed_description = assessment_generator.compact_use_case_description_for_pptx(full_description)
        self.assertEqual(full_description, displayed_description)
        self.assertEqual(full_description, cases[0]["description"])

        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_use_case_list_page(
            ppt, width, height, group, chunks[0], page_number=6,
            total_groups=1, service_index=1, list_page_index=1,
            list_page_count=1, appendix=True,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "use-cases.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)

        self.assertEqual(1, len(presentation.slides))
        slide = presentation.slides[0]
        page_text = self.pptx_slide_text(slide)
        self.assertIn("候補1", page_text)
        self.assertIn("候補15", page_text)
        self.assertNotIn("対象サービス", page_text)
        self.assertNotIn("候補 1–15 / 15", page_text)
        table_shape = next(shape for shape in slide.shapes if shape.has_table)
        self.assertEqual(16, len(table_shape.table.rows))
        headers = [cell.text for cell in table_shape.table.rows[0].cells]
        self.assertEqual(["No.", "ユースケース", "主なAI技術", "業務価値・処理"], headers)
        self.assertNotIn("PoC", headers)
        table_text = "\n".join(cell.text for row in table_shape.table.rows for cell in row.cells)
        self.assertNotIn("…", table_text)
        technology_column = headers.index("主なAI技術")
        technology_column_width = (
            table_shape.table.columns[technology_column].width
            / assessment_generator.PptCanvas.POINT_TO_EMU
        )
        fitted_technology = assessment_generator.fit_use_case_technology_text(
            ppt, "機械学習・スコアリング・生成AI",
            technology_column_width,
        )
        self.assertEqual("機械学習・スコアリング・生成AI", fitted_technology)
        footer = next(shape for shape in slide.shapes if shape.has_text_frame and "Copyright" in shape.text)
        self.assertLess(table_shape.top + table_shape.height, footer.top)
        expected_table_top = round(
            29 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        self.assertGreaterEqual(table_shape.top, expected_table_top)
        minimum_footer_clearance = round(
            9 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        self.assertGreaterEqual(
            footer.top - (table_shape.top + table_shape.height), minimum_footer_clearance,
        )
        expected_row_height = round(
            assessment_generator.USE_CASE_LIST_ROW_HEIGHT
            * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        for row in list(table_shape.table.rows)[1:]:
            self.assertAlmostEqual(row.height, expected_row_height, delta=2)
        expected_table_height = round(
            (
                assessment_generator.USE_CASE_LIST_HEADER_HEIGHT
                + assessment_generator.USE_CASE_LIST_ROW_HEIGHT * 15
            ) * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        self.assertAlmostEqual(table_shape.height, expected_table_height, delta=4)
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in self.pptx_text_font_sizes(presentation, include_footer=False)
        ))

        with tempfile.TemporaryDirectory() as temporary_directory:
            priority_icon_path = Path(temporary_directory) / "priority-insight.png"
            priority_icon_path.write_bytes(assessment_generator.ORACLE_LOGO_IMAGE.read_bytes())
            marked_ppt = assessment_generator.PptCanvas(width, height)
            assessment_generator.draw_use_case_list_page(
                marked_ppt, width, height, group, chunks[0], page_number=6,
                total_groups=1, service_index=1, list_page_index=1,
                list_page_count=1, appendix=False,
                priority_pocs=[
                    {"use_case_no": "1"}, {"use_case_no": "5"}, {"use_case_no": "15"},
                ],
                priority_icon_path=priority_icon_path,
            )
            marked_path = Path(temporary_directory) / "marked-use-cases.pptx"
            marked_ppt.write(marked_path)
            marked_presentation = Presentation(marked_path)
        marked_slide = marked_presentation.slides[0]
        marked_table_shape = next(shape for shape in marked_slide.shapes if shape.has_table)
        marked_headers = [cell.text for cell in marked_table_shape.table.rows[0].cells]
        self.assertEqual(["No.", "ユースケース", "主なAI技術", "業務価値・処理"], marked_headers)
        marked_table_text = "\n".join(
            cell.text for row in marked_table_shape.table.rows for cell in row.cells
        )
        for priority_label in ("P1", "P2", "P3"):
            self.assertNotIn(priority_label, marked_table_text)

        # 表内は全行を同じ中立色にし、優先候補は表外左側のアイコンで示す。
        table_left = marked_table_shape.left
        table_top = marked_table_shape.top
        table_right = table_left + marked_table_shape.width
        table_bottom = table_top + marked_table_shape.height
        table_overlays = [
            shape for shape in marked_slide.shapes
            if not shape.has_table
            and shape.left >= table_left and shape.left + shape.width <= table_right
            and shape.top >= table_top and shape.top + shape.height <= table_bottom
        ]
        self.assertEqual([], table_overlays)

        row_by_number = {
            row.cells[0].text: row
            for row in list(marked_table_shape.table.rows)[1:]
        }
        self.assertTrue(all(
            str(cell.fill.fore_color.rgb).upper() == "F5F7F8"
            for row in row_by_number.values() for cell in row.cells
        ))

        priority_icons = [
            shape for shape in marked_slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        self.assertEqual(3, len(priority_icons))
        expected_icon_size = round(
            assessment_generator.USE_CASE_LIST_PRIORITY_ICON_SIZE
            * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        expected_icon_gap = round(
            assessment_generator.USE_CASE_LIST_PRIORITY_ICON_GAP
            * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        for icon in priority_icons:
            self.assertAlmostEqual(icon.width, expected_icon_size, delta=2)
            self.assertAlmostEqual(icon.height, expected_icon_size, delta=2)
            self.assertAlmostEqual(table_left - (icon.left + icon.width), expected_icon_gap, delta=2)
            self.assertLessEqual(icon.left + icon.width, table_left)

        expected_priority_centers: list[int] = []
        current_top = table_top + marked_table_shape.table.rows[0].height
        for row in list(marked_table_shape.table.rows)[1:]:
            if row.cells[0].text in {"1", "5", "15"}:
                expected_priority_centers.append(current_top + row.height // 2)
            current_top += row.height
        actual_priority_centers = sorted(icon.top + icon.height // 2 for icon in priority_icons)
        self.assertEqual(3, len(expected_priority_centers))
        center_tolerance = round(
            0.25 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        for actual_center, expected_center in zip(actual_priority_centers, expected_priority_centers):
            self.assertAlmostEqual(actual_center, expected_center, delta=center_tolerance)

        marked_footer = next(
            shape for shape in marked_slide.shapes
            if shape.has_text_frame and "Copyright" in shape.text
        )
        self.assertGreaterEqual(marked_footer.top - table_bottom, minimum_footer_clearance)

        with tempfile.TemporaryDirectory() as temporary_directory:
            priority_icon_path = Path(temporary_directory) / "priority-insight.png"
            priority_icon_path.write_bytes(assessment_generator.ORACLE_LOGO_IMAGE.read_bytes())
            second_service_cases = [
                {**case, "use_case_id": f"S02-UC{int(case['no']):02d}"}
                for case in cases
            ]
            second_service_group = {
                "service_name": "Example 2", "service_type": "SaaS",
                "use_cases": second_service_cases,
            }
            multi_service_ppt = assessment_generator.PptCanvas(width, height)
            assessment_generator.draw_use_case_list_page(
                multi_service_ppt, width, height, second_service_group, second_service_cases,
                page_number=7, total_groups=3, service_index=2, list_page_index=1,
                list_page_count=1, appendix=True,
                priority_pocs=[
                    {"use_case_id": "UC01", "use_case_no": "1"},
                    {"use_case_id": "UC05", "use_case_no": "5"},
                    {"use_case_id": "UC15", "use_case_no": "15"},
                ],
                priority_icon_path=priority_icon_path,
            )
            output_path = Path(temporary_directory) / "multi-service-use-cases.pptx"
            multi_service_ppt.write(output_path)
            multi_service_presentation = Presentation(output_path)
        multi_service_text = self.pptx_slide_text(multi_service_presentation.slides[0])
        self.assertIn("サービス 2 / 3", multi_service_text)
        self.assertNotIn("対象サービス 2/3", multi_service_text)
        self.assertEqual([], [
            shape for shape in multi_service_presentation.slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ])

    def test_use_case_list_keeps_borderline_meiryo_rows_on_one_slide(self) -> None:
        """実描画で1行に収まる境界文字列を、粗い文字数推定で2行扱いしない。"""
        boundary_rows = [
            (
                "勤怠・PC稼働ログの確認支援",
                "異常検知・機械学習・データ可視化",
                "勤怠、申請、PCログの差異を採点する。人事の確認優先度と見落としを改善する。",
            ),
            (
                "就業規則変更時の設定影響確認支援",
                "生成AI・RAG・Human-in-the-loop",
                "規則差分と設定値から影響候補を生成する。設定変更の確認漏れを抑制する。",
            ),
            (
                "打刻漏れの個別案内支援",
                "機械学習・生成AI",
                "打刻履歴と勤務予定から漏れを検出する。従業員への適切な申請案内を支援する。",
            ),
            (
                "休暇取得見込みの可視化",
                "時系列予測・データ可視化",
                "休暇履歴と組織属性から取得見込みを予測する。管理者の取得促進計画を支援する。",
            ),
            (
                "初期設定の対話型ガイド",
                "生成AI・AI Agent",
                "就業規則と導入質問から設定候補を作成する。導入担当の設定作業を標準化する。",
            ),
            (
                "労務ルール確認チェックリスト生成",
                "生成AI・RAG",
                "規則文書と勤怠設定から確認項目を生成する。労務担当の点検準備を効率化する。",
            ),
        ]
        rows = (boundary_rows * 3)[:15]
        cases = [
            {
                "no": index, "use_case_id": f"UC{index:02d}",
                "use_case": row[0], "ai_technology": row[1], "description": row[2],
            }
            for index, row in enumerate(rows, 1)
        ]
        group = {
            "service_name": "境界幅テストサービス", "service_type": "勤怠管理",
            "use_cases": cases,
        }
        widths = assessment_generator.use_case_list_column_widths(
            assessment_generator.PPTX_WIDESCREEN_WIDTH,
        )
        heights = assessment_generator.use_case_list_row_heights(cases, widths)
        self.assertEqual([15], [len(chunk) for chunk in assessment_generator.use_case_list_chunks(group)])
        self.assertTrue(all(
            abs(height - assessment_generator.USE_CASE_LIST_ROW_HEIGHT) < 0.01
            for height in heights
        ))

    def test_use_case_list_uses_real_table_region_for_one_wrapped_row(self) -> None:
        """見出しとフッターの間に収まる15行を、14+1へ不要に分割しない。"""
        cases = [
            {
                "no": index,
                "use_case_id": f"UC{index:02d}",
                "use_case": f"業務候補{index}",
                "ai_technology": "生成AI・機械学習",
                "description": "履歴から候補を検出し、管理者の確認を支援する。",
            }
            for index in range(1, 16)
        ]
        # Meiryo UIの実表示では一行に収まる境界文字列。
        cases[1]["description"] = (
            "期限・進捗・人員・ロケーションから作業案を算出する。"
            "管理者の割当判断を標準化する。"
        )
        cases[11]["description"] = (
            "商品・ロケーション属性から不備候補を検出する。"
            "登録修正による現場エラーを抑制する。"
        )
        # ユースケース名だけが実際に二行必要。
        cases[2]["use_case"] = "問い合わせ・障害対応支援と状況説明・レポート作成"
        group = {
            "service_name": "一覧領域テストサービス",
            "service_type": "業務管理",
            "use_cases": cases,
        }
        widths = assessment_generator.use_case_list_column_widths(
            assessment_generator.PPTX_WIDESCREEN_WIDTH,
        )
        heights = assessment_generator.use_case_list_row_heights(cases, widths)
        self.assertEqual(1, sum(
            height > assessment_generator.USE_CASE_LIST_ROW_HEIGHT + 0.01
            for height in heights
        ))
        chunks = assessment_generator.use_case_list_chunks(group)
        self.assertEqual([15], [len(chunk) for chunk in chunks])

        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_use_case_list_page(
            ppt, width, height, group, chunks[0], page_number=6,
            total_groups=1, service_index=1, list_page_index=1,
            list_page_count=1,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "mixed-height-use-cases.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)

        slide = presentation.slides[0]
        table_shape = next(shape for shape in slide.shapes if shape.has_table)
        title_shape = next(
            shape for shape in slide.shapes
            if shape.has_text_frame and "AIユースケース一覧" in shape.text
        )
        footer = next(
            shape for shape in slide.shapes
            if shape.has_text_frame and "Copyright" in shape.text
        )
        expected_table_top = round(
            assessment_generator.USE_CASE_LIST_TABLE_TOP
            * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        self.assertGreaterEqual(table_shape.top, expected_table_top)
        self.assertLessEqual(title_shape.top + title_shape.height, table_shape.top)
        minimum_footer_clearance = round(
            9 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        self.assertGreaterEqual(
            footer.top - (table_shape.top + table_shape.height),
            minimum_footer_clearance,
        )
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in self.pptx_text_font_sizes(presentation, include_footer=False)
        ))

    def test_use_case_list_preserves_long_cells_by_growing_rows_and_paginating(self) -> None:
        exact_name = "就業規則・設定・FAQのナレッジ支援"
        exact_technology = "生成AI・RAG・Vector Search"
        long_description = (
            "勤怠実績、申請履歴、就業規則、設定履歴、FAQ、問い合わせ履歴を権限別に照合し、"
            "根拠付きの確認候補、対応案、次の確認担当を提示して管理者と利用者の判断を支援する。"
        )
        cases = [
            {
                "no": index,
                "use_case_id": f"UC{index:02d}",
                "use_case": exact_name if index == 2 else f"{exact_name}（対象組織{index}）",
                "ai_technology": exact_technology,
                "description": f"{long_description} 対象候補{index}。",
            }
            for index in range(1, 16)
        ]
        group = {"service_name": "長文テストサービス", "service_type": "SaaS", "use_cases": cases}
        chunks = assessment_generator.use_case_list_chunks(group)
        self.assertGreater(len(chunks), 1)
        self.assertEqual(cases, [case for chunk in chunks for case in chunk])

        narrow_width = 8 * assessment_generator.mm
        probe = assessment_generator.PptCanvas(
            assessment_generator.PPTX_WIDESCREEN_WIDTH,
            assessment_generator.PPTX_WIDESCREEN_HEIGHT,
        )
        self.assertEqual(exact_name, assessment_generator.fit_complete_use_case_text(probe, exact_name, narrow_width))
        self.assertEqual(exact_name, assessment_generator.use_case_list_cell_text(probe, exact_name, narrow_width))
        self.assertEqual(
            exact_technology,
            assessment_generator.fit_use_case_technology_text(probe, exact_technology, narrow_width),
        )
        self.assertEqual(
            long_description,
            assessment_generator.compact_use_case_description_for_pptx(long_description),
        )

        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        for page_index, chunk in enumerate(chunks, 1):
            assessment_generator.draw_use_case_list_page(
                ppt, width, height, group, chunk, page_number=page_index,
                total_groups=1, service_index=1, list_page_index=page_index,
                list_page_count=len(chunks), appendix=False,
            )
            if page_index < len(chunks):
                ppt.showPage()
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "long-use-cases.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)

        self.assertEqual(len(chunks), len(presentation.slides))
        rendered_rows: list[list[str]] = []
        minimum_footer_clearance = round(
            9 * assessment_generator.mm * assessment_generator.PptCanvas.POINT_TO_EMU
        )
        for slide, chunk in zip(presentation.slides, chunks):
            table_shape = next(shape for shape in slide.shapes if shape.has_table)
            table = table_shape.table
            widths = [
                column.width / assessment_generator.PptCanvas.POINT_TO_EMU
                for column in table.columns
            ]
            footer = next(shape for shape in slide.shapes if shape.has_text_frame and "Copyright" in shape.text)
            self.assertGreaterEqual(
                footer.top - (table_shape.top + table_shape.height), minimum_footer_clearance,
            )
            expected_heights = assessment_generator.use_case_list_row_heights(chunk, widths)
            for row, expected_height in zip(list(table.rows)[1:], expected_heights):
                values = [cell.text for cell in row.cells]
                rendered_rows.append(values)
                self.assertGreaterEqual(
                    row.height,
                    round(expected_height * assessment_generator.PptCanvas.POINT_TO_EMU) - 2,
                )
        self.assertEqual(
            [assessment_generator.use_case_list_native_row(case) for case in cases],
            rendered_rows,
        )
        self.assertEqual(exact_name, rendered_rows[1][1])
        self.assertEqual(exact_technology, rendered_rows[1][2])
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in self.pptx_text_font_sizes(presentation, include_footer=False)
        ))

    def test_cost_notes_are_present_in_native_pptx_output(self) -> None:
        assessment_generator.register_japanese_font()
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        cost_estimate = assessment_generator.build_poc_cost_estimate()
        assessment_generator.draw_cost_estimate_page(ppt, width, height, cost_estimate, 36)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "cost-estimate.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        slide = presentation.slides[0]
        text = self.pptx_slide_text(slide)
        for heading in ("価格前提", "稼働前提", "再計算条件"):
            self.assertIn(heading, text)
        # API取得が0件の凍結見積もりでは、旧スナップショットの
        # 「リスト価格ベース」を顧客向け表示で事実化しない。
        displayed_first_note = "税抜概算。地域・SKU・価格時点・割引・無料枠・転送費を確認します。"
        self.assertIn(displayed_first_note, text)
        self.assertIn("OCI利用料のみ", text)
        for hidden_text in (
            "価格根拠",
            "設計・構築・運用支援費",
            "構築・支援費",
            "運用支援費",
        ):
            self.assertNotIn(hidden_text, text)

        table_shapes = [shape for shape in slide.shapes if shape.has_table]
        self.assertEqual(1, len(table_shapes))
        table = table_shapes[0].table
        self.assertEqual(3, len(table.columns))
        self.assertEqual(
            ["項目", "月額（円）", "前提・備考"],
            [cell.text for cell in table.rows[0].cells],
        )
        self.assertEqual(len(cost_estimate["lines"]) + 2, len(table.rows))
        for row_index, line in enumerate(cost_estimate["lines"], start=1):
            self.assertEqual(f"{line.monthly_jpy:,}円", table.cell(row_index, 1).text)
            expected_assumption = str(line.assumption)
            if "入力100 token" in expected_assumption:
                expected_assumption = (
                    "月間10,000リクエストの仮置き。RAGの検索文脈を含む入出力tokenを実測し再見積り"
                )
            self.assertEqual(expected_assumption, table.cell(row_index, 2).text)
        total_row = len(cost_estimate["lines"]) + 1
        self.assertEqual("合計", table.cell(total_row, 0).text)
        self.assertEqual(
            f"{cost_estimate['total_monthly_jpy']:,}円",
            table.cell(total_row, 1).text,
        )
        self.assertEqual(
            f"10週間では{round(cost_estimate['total_monthly_jpy'] * 2.5):,}円。OCI利用料のみの単純換算",
            table.cell(total_row, 2).text,
        )
        self.assertNotRegex(text, r"(?:\.{3,}|…+)")
        self.assertTrue(all(str(getattr(line, "source", "")).strip() for line in cost_estimate["lines"]))

        formatted_total = f"{cost_estimate['total_monthly_jpy']:,}円"
        emphasized_total_shapes = [
            shape for shape in slide.shapes
            if shape.has_text_frame and formatted_total in shape.text and "月" in shape.text
        ]
        self.assertTrue(emphasized_total_shapes)
        emphasized_sizes = [
            run.font.size.pt
            for shape in emphasized_total_shapes
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.text.strip() and run.font.size is not None
        ]
        self.assertTrue(any(size >= 16 for size in emphasized_sizes))

    def test_long_story_heading_is_kept_within_a_single_line_title_shape(self) -> None:
        assessment_generator.register_japanese_font()
        width, height = landscape(A4)
        title = "付録A：利用者・担当者・責任者の業務を一つの運用フローで捉える"
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator._executive_evidence_header(ppt, width, height, title, "ASSESSMENT STORY")
        title_shape = next(shape for shape in ppt.slide.shapes if shape.has_text_frame and shape.text == title)
        font_size = title_shape.text_frame.paragraphs[0].runs[0].font.size.pt
        self.assertEqual(20.0, font_size)

    def test_pptx_base_header_keeps_the_same_title_geometry_as_story_pages(self) -> None:
        assessment_generator.register_japanese_font()
        width, height = landscape(A4)
        title = "共通ヘッダーの見出し位置を検証する"
        base = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_pptx_base_header(base, width, height, title, "2.1 業界・サービスの変化")
        story = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_story_header(story, width, height, title, section="2.1 業界・サービスの変化")
        base_shape = next(shape for shape in base.slide.shapes if shape.has_text_frame and shape.text == title)
        story_shape = next(shape for shape in story.slide.shapes if shape.has_text_frame and shape.text == title)
        self.assertEqual((base_shape.left, base_shape.top), (story_shape.left, story_shape.top))
        self.assertEqual(20.0, base_shape.text_frame.paragraphs[0].runs[0].font.size.pt)


    def test_poc_support_bands_and_schedule_cells_are_vertically_centered(self) -> None:
        assessment = self.detailed_assessment()
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator._draw_pptx_poc_support_page(ppt, width, height, assessment, 48)
        expected_labels = {"1  PoC立上げ", "2  体制・役割", "3  スケジュール（目安）", "フェーズ", "主なタスク", "W1"}
        labels = [shape for shape in ppt.slide.shapes if shape.has_text_frame and shape.text in expected_labels]
        self.assertEqual(expected_labels, {shape.text for shape in labels})
        self.assertTrue(all(shape.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE for shape in labels))

    @staticmethod
    def priority_measurement_design_fixture(assessment: dict) -> dict:
        """P1〜P3と一対一の、事業効果と製品先行KPIを持つ定量設計。"""
        priorities = assessment_generator.priority_pocs_for(
            assessment, assessment.get("consulting_front_matter"),
        )
        metrics = ("12〜18%", "18〜24%", "8〜12%")
        product_metrics = ("65〜80%", "70〜85%", "60〜75%")
        display_layers = (
            "provider_business_kpi", "customer_outcome_kpi", "provider_business_kpi",
        )
        items = []
        for index, (poc, metric, product_metric, display_layer) in enumerate(zip(
            priorities, metrics, product_metrics, display_layers,
        ), 1):
            product_owner = {
                "scope": "provider", "role": f"製品KPI責任者P{index}", "status": "confirmed",
            }
            customer_owner = {
                "scope": "customer", "role": f"顧客業務責任者P{index}", "status": "confirmed",
            }
            provider_owner = {
                "scope": "provider", "role": f"事業KPI責任者P{index}", "status": "confirmed",
            }
            product_kpi_name = f"製品先行KPI P{index}"
            business_kpi_name = (
                f"顧客業務成果P{index}"
                if display_layer == "customer_outcome_kpi"
                else f"提供者事業成果P{index}"
            )
            display_owner = customer_owner if display_layer == "customer_outcome_kpi" else provider_owner
            display_attribution = "contributory" if display_layer == "customer_outcome_kpi" else "enabling"
            display_beneficiary = "shared" if display_layer == "customer_outcome_kpi" else "provider"
            business_formula = f"{business_kpi_name}を同一条件で比較し、{metric}を判定する"
            items.append({
                "priority": poc["priority"],
                "use_case_id": poc["use_case_id"],
                "theme": poc["theme"],
                "business_outcome": f"{poc['theme']}による事業効果を高める",
                "headline_metric": metric,
                "kpi": business_kpi_name,
                "detail": "事業効果を本番化判断の主目標とし、製品先行KPIから効果へ至る因果を同時に検証する。",
                "baseline_definition": "表示する事業KPIの直近12週における現行実績",
                "comparison_condition": "同じ対象・期間・判定基準でAI利用有無を比較",
                "formula": business_formula,
                "poc_gate": {
                    "measurement_period": "12週間",
                    "business_go_condition": f"事業効果で{metric}を達成",
                    "leading_kpi_condition": f"製品先行KPIで{product_metric}を達成",
                    "go_condition": f"事業効果{metric}、製品先行KPI{product_metric}をともに達成",
                    "stop_condition": "誤作動、品質低下または運用負荷増大時は停止",
                    "decision_owner": f"PoC承認責任者P{index}",
                },
                "target_rationale": "代表範囲の実測により本番化を判断できるため設定する。",
                "beneficiary": display_beneficiary,
                "metric_owner": display_owner,
                "attribution_level": display_attribution,
                "causal_link": f"因果連鎖P{index}：製品KPI改善が顧客成果と事業成果へ段階的に寄与する。",
                "display_kpi_layer": display_layer,
                "product_kpi": {
                    "kpi": product_kpi_name, "target": product_metric,
                    "baseline_definition": "PoC対象の現行運用実績",
                    "comparison_condition": "同一条件でAI利用有無を比較",
                    "formula": f"{product_kpi_name}を同一条件で比較し、{product_metric}を判定する",
                    "metric_owner": product_owner,
                    "attribution_level": "direct",
                    "causal_link": "AI機能が製品KPIを直接変化させる。",
                },
                "customer_outcome_kpi": {
                    "kpi": f"顧客業務成果P{index}",
                    "target": metric if display_layer == "customer_outcome_kpi" else "PoCで確定",
                    "baseline_definition": "顧客側の現状実績",
                    "comparison_condition": "製品KPI達成群と現行運用を比較",
                    "formula": business_formula if display_layer == "customer_outcome_kpi" else "顧客実績値を用いてPoCで確定",
                    "metric_owner": customer_owner,
                    "attribution_level": "contributory",
                    "causal_link": "製品KPI改善が顧客業務成果へ寄与する。",
                },
                "provider_business_kpi": {
                    "kpi": f"提供者事業成果P{index}",
                    "target": metric if display_layer == "provider_business_kpi" else "PoCで確定",
                    "baseline_definition": "提供者側の現状実績",
                    "comparison_condition": "PoC利用群と現行提供条件を比較",
                    "formula": business_formula if display_layer == "provider_business_kpi" else "提供者実績値を用いてPoCで確定",
                    "metric_owner": provider_owner,
                    "attribution_level": "enabling",
                    "causal_link": "顧客成果と利用定着が提供者事業成果を支える。",
                },
                "basis_type": "decision_threshold",
            })
        return {
            "schema_version": assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
            "status": "decision_thresholds",
            "lead": "優先PoCと同じ3テーマの事業効果、製品先行KPI、責任分界、開始判断条件を一体で確認する。",
            "management_targets": [],
            "items": items,
        }

    def test_p6_contract_requires_the_exact_priority_poc_set_and_responsibility_boundary(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        design = self.priority_measurement_design_fixture(assessment)
        normalized = assessment_generator.normalize_poc_measurement_design(design, assessment)
        expected = assessment_generator.priority_pocs_for(
            assessment, assessment["consulting_front_matter"],
        )

        self.assertEqual(
            [(item["priority"], item["use_case_id"], item["theme"]) for item in expected],
            [(item["priority"], item["use_case_id"], item["theme"]) for item in normalized["items"]],
        )
        self.assertTrue(all(
            item.get("display_kpi_layer") in assessment_generator.BUSINESS_EFFECT_DISPLAY_LAYERS
            and item["headline_metric"] == item[item["display_kpi_layer"]]["target"]
            and item["product_kpi"]["metric_owner"]["scope"] == "provider"
            and item["product_kpi"]["attribution_level"] == "direct"
            and item["customer_outcome_kpi"]["metric_owner"]["scope"] == "customer"
            and item["provider_business_kpi"]["metric_owner"]["scope"] == "provider"
            for item in normalized["items"]
        ))

        duplicate = json.loads(json.dumps(design, ensure_ascii=False))
        duplicate["items"][1]["use_case_id"] = duplicate["items"][0]["use_case_id"]
        self.assertEqual({}, assessment_generator.normalize_poc_measurement_design(duplicate, assessment))

        wrong_theme = json.loads(json.dumps(design, ensure_ascii=False))
        wrong_theme["items"][2]["theme"] = assessment["use_cases"][3]["use_case"]
        self.assertEqual({}, assessment_generator.normalize_poc_measurement_design(wrong_theme, assessment))

        reordered = json.loads(json.dumps(design, ensure_ascii=False))
        reordered["items"][0], reordered["items"][1] = reordered["items"][1], reordered["items"][0]
        self.assertEqual({}, assessment_generator.normalize_poc_measurement_design(reordered, assessment))

    def test_p5_and_p6_share_business_effects_and_product_leading_kpis(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        design = assessment_generator.normalize_poc_measurement_design(
            self.priority_measurement_design_fixture(assessment), assessment,
        )
        story = assessment_generator.assessment_story_for(assessment)
        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        p5 = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_industry_value_story_page(
            p5, width, height, story["industry_value_story"], story["front_matter"], 5, design,
        )
        p6 = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_poc_quantitative_target_page(
            p6, width, height, design, assessment["service_name"], 6,
        )
        p5_text = self.pptx_slide_text(p5.slide)
        p6_text = self.pptx_slide_text(p6.slide)
        self.assertIn("事業効果目標", p5_text)
        self.assertIn("次頁：製品先行KPI", p5_text)
        for item in design["items"]:
            product = item["product_kpi"]
            for value in (item["headline_metric"], item["kpi"]):
                self.assertIn(value, p5_text)
                self.assertIn(value, p6_text)
            for value in (product["target"], product["kpi"]):
                self.assertIn(value, p5_text)
                self.assertIn(value, p6_text)

    def test_invalid_numeric_contract_never_silently_downgrades_to_generic_prepoc_page(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        invalid = self.priority_measurement_design_fixture(assessment)
        invalid["items"][1]["poc_gate"]["go_condition"] = "別の数値10%だけを判定する"
        assessment["poc_measurement_design"] = invalid
        frozen = assessment_generator.poc_measurement_design_for(assessment)
        self.assertEqual("invalid_numeric_contract", frozen["status"])

    def test_decision_thresholds_reject_unverified_downstream_numbers_but_allow_variable_formulas(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        design = self.priority_measurement_design_fixture(assessment)

        unsupported = json.loads(json.dumps(design, ensure_ascii=False))
        unsupported["items"][0]["customer_outcome_kpi"].update({
            "target": "20%削減", "formula": "顧客実績値×20%",
        })
        self.assertEqual({}, assessment_generator.normalize_poc_measurement_design(unsupported, assessment))

        variable = json.loads(json.dumps(design, ensure_ascii=False))
        variable["items"][0]["customer_outcome_kpi"].update({
            "target": "顧客実績値×改善率",
            "formula": "顧客実績値×改善率を用いてPoCで算定する",
        })
        normalized = assessment_generator.normalize_poc_measurement_design(variable, assessment)
        self.assertEqual("decision_thresholds", normalized["status"])
        self.assertEqual("顧客実績値×改善率", normalized["items"][0]["customer_outcome_kpi"]["target"])

    def test_strict_validation_rejects_decision_thresholds_without_generation_source_audit(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        assessment["poc_measurement_design"] = self.priority_measurement_design_fixture(assessment)
        payload = {
            "format": assessment_generator.ASSESSMENT_JSON_FORMAT,
            "assessment": assessment,
            "research": {
                "midterm_plan": {"status": "disabled"}, "industry_sources": [],
                "poc_measurement_design_audit": {
                    "status": "decision_thresholds", "basis_type": "decision_threshold",
                    "reason": "手動記載のため生成元と表示値監査がない。",
                },
            },
        }
        self.assertFalse(any(
            "generation_source=llm_estimate" in error
            for error in assessment_generator.validate_assessment_payload(payload)
        ))
        self.assertTrue(any(
            "generation_source=llm_estimate" in error
            for error in assessment_generator.validate_assessment_payload(payload, strict=True)
        ))

    def test_legacy_poc_measurement_design_upgrades_only_with_exact_theme_mapping(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        selected = assessment_generator.priority_pocs_for(
            assessment, assessment["consulting_front_matter"],
        )
        for schema_version in ("2", "3"):
            with self.subTest(schema_version=schema_version):
                legacy = {
                    "schema_version": schema_version,
                    "status": "decision_thresholds",
                    "items": [{
                        "business_outcome": poc["theme"],
                        "headline_metric": metric,
                        "kpi": f"製品KPI{index}",
                        "detail": "旧契約の製品KPIを同じ優先PoCへ安全に移行する。",
                        "baseline_definition": "対象機能の現行実績",
                        "comparison_condition": "同一条件でAI利用有無を比較",
                        "formula": f"製品KPI改善率を{metric}として算定する",
                        "poc_gate": f"同一条件で{metric}を達成する",
                        "target_rationale": "代表データで直接測定できるため設定する。",
                        "supporting_use_cases": [poc["theme"]],
                        "basis_type": "decision_threshold",
                    } for index, (poc, metric) in enumerate(zip(
                        selected, ("8〜12%", "15〜20%", "10〜16%"),
                    ), 1)],
                }
                upgraded = assessment_generator.normalize_poc_measurement_design(legacy, assessment)
                self.assertEqual(
                    assessment_generator.POC_MEASUREMENT_DESIGN_SCHEMA_VERSION,
                    upgraded["schema_version"],
                )
                self.assertEqual("legacy_migrated", upgraded["migration_status"])
                self.assertEqual(
                    [item["use_case_id"] for item in selected],
                    [item["use_case_id"] for item in upgraded["items"]],
                )
                self.assertTrue(all(
                    item["display_kpi_layer"] == "product_kpi"
                    and item["product_kpi"]["metric_owner"]["scope"] == "provider"
                    for item in upgraded["items"]
                ))

                mismatched = json.loads(json.dumps(legacy, ensure_ascii=False))
                mismatched["items"][2]["business_outcome"] = "別のユースケース"
                mismatched["items"][2]["supporting_use_cases"] = ["別のユースケース"]
                safe = assessment_generator.normalize_poc_measurement_design(mismatched, assessment)
                self.assertEqual({}, safe)

    def test_default_selection_and_details_share_priority_set_and_fit_widescreen_safely(self) -> None:
        assessment = self.detailed_assessment()
        front = assessment["consulting_front_matter"]
        assessment_generator.materialize_assessment_decision_contract(assessment, front)
        assessment["poc_measurement_design"] = self.priority_measurement_design_fixture(assessment)
        rendering_cost_estimate = assessment_generator.cost_estimate_from_snapshot({
            "schema_version": "1", **TEST_COST_ESTIMATE,
        })
        self.assertIsNotNone(rendering_cost_estimate)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "p6-p8-priority-contract.pptx"
            assessment_generator.create_pptx(
                assessment, assessment_generator.DEFAULT_ARCHITECTURE_IMAGE,
                output_path, rendering_cost_estimate,
            )
            presentation = Presentation(output_path)

        self.assertAlmostEqual(
            presentation.slide_width / presentation.slide_height, 16 / 9, places=3,
        )
        slides_with_text = [
            (slide, self.pptx_slide_text(slide)) for slide in presentation.slides
        ]
        p6_slide, p6_text = next(
            pair for pair in slides_with_text if "AI活用で期待できるビジネスインパクト" in pair[1]
        )
        p8_slide, p8_text = next(
            pair for pair in slides_with_text if "15のAIユースケースから整理した AI技術アプローチの代表3テーマ" in pair[1]
        )
        priorities = assessment_generator.priority_pocs_for(assessment, front)
        for item in priorities:
            self.assertIn(item["theme"], p8_text)

        for phrase in ("AIで変わる判断", "確認するKPI", "提供企業側の確認"):
            self.assertIn(phrase, p6_text)
        for phrase in (
            "整理観点", "業務価値", "既存機能との差分", "データ準備",
            "比較検証の成立性", "代表とする理由", "PoCで確かめること", "準備のポイント",
        ):
            self.assertIn(phrase, p8_text)
        for phrase in ("順位指数", "既知平均", "採点率", "開始保留", "未充足", "P1｜", "5候補", "次段階"):
            self.assertNotIn(phrase, p8_text)

        detail_slides = [
            slide for slide, slide_text in slides_with_text
            if "実装対象：データ準備から既存業務への返却まで" in slide_text
            and "Oracle Database／OCIでの実装" in slide_text
            and "DBオブジェクト" in slide_text
        ]
        self.assertEqual(3, len(detail_slides))

        def solid_fill_rgb(shape) -> str | None:
            if not hasattr(shape, "fill"):
                return None
            try:
                rgb = shape.fill.fore_color.rgb
            except TypeError:
                return None
            return str(rgb).upper() if rgb is not None else None

        def visible_shape_colors(shape) -> set[str]:
            colors = {rgb for rgb in (solid_fill_rgb(shape),) if rgb}
            try:
                if shape.line.color.rgb is not None:
                    colors.add(str(shape.line.color.rgb).upper())
            except (AttributeError, TypeError):
                pass
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        try:
                            if cell.fill.fore_color.rgb is not None:
                                colors.add(str(cell.fill.fore_color.rgb).upper())
                        except (AttributeError, TypeError):
                            pass
            return colors

        for expected_accent, detail_slide in zip(assessment_generator.POC_THEME_ACCENT_HEX, detail_slides):
            flow_heading = next(
                shape for shape in detail_slide.shapes
                if shape.has_text_frame
                and shape.text.strip() == "実装対象：データ準備から既存業務への返却まで"
            )
            # 上段のスコープ本文は、次のフロー見出しへ重ならない。必要な高さは
            # 12ptの全文を基に上方向へ確保する。
            upper_text_shapes = [
                shape for shape in detail_slide.shapes
                if shape.has_text_frame and shape.text.strip()
                and shape.shape_id != flow_heading.shape_id and shape.top < flow_heading.top
            ]
            self.assertTrue(all(
                shape.top + shape.height <= flow_heading.top
                for shape in upper_text_shapes
            ))
            self.assertIn(
                expected_accent.lstrip("#"),
                {
                    rgb
                    for shape in detail_slide.shapes
                    for rgb in visible_shape_colors(shape)
                },
            )

        for slide in (p6_slide, p8_slide):
            footer = next(
                shape for shape in slide.shapes
                if shape.has_text_frame and "Copyright ©" in shape.text
            )
            page_number_shapes = {
                shape.shape_id for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip().isdigit()
                and shape.top > presentation.slide_height * 0.9
            }
            self.assertTrue(all(
                shape.top + shape.height <= footer.top
                for shape in slide.shapes
                if shape.shape_id != footer.shape_id and shape.shape_id not in page_number_shapes
            ))
            body_sizes = []
            for shape in slide.shapes:
                if (
                    shape.shape_id == footer.shape_id
                    or shape.shape_id in page_number_shapes
                    or self.is_ten_point_chrome_shape(shape, presentation.slide_height)
                ):
                    continue
                frames = [shape.text_frame] if shape.has_text_frame else []
                if shape.has_table:
                    frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
                for frame in frames:
                    for paragraph in frame.paragraphs:
                        body_sizes.extend(
                            run.font.size.pt if run.font.size else None
                            for run in paragraph.runs if run.text.strip()
                        )
            self.assertTrue(body_sizes)
            self.assertTrue(all(
                size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
                for size in body_sizes
            ))

    def test_p6_max_density_responsibility_and_gate_text_remain_footer_safe(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        design = self.priority_measurement_design_fixture(assessment)
        for item in design["items"]:
            item["causal_link"] = "製品KPIの改善が顧客業務成果へ寄与し、利用定着を介して提供者事業成果へ接続する因果関係を検証する。"
            item["poc_gate"]["stop_condition"] = "品質、安全性、説明可能性または現場運用負荷が合意した許容範囲を外れた場合は停止する。"
            item["poc_gate"]["decision_owner"] = "顧客業務責任者・製品責任者による共同判定会議"
            for layer_name, role in (
                ("product_kpi", "製品・AI機能品質および計測基盤の責任者"),
                ("customer_outcome_kpi", "対象業務プロセスと現場運用の責任者"),
                ("provider_business_kpi", "サービス事業化と継続運用の責任者"),
            ):
                item[layer_name]["metric_owner"]["role"] = role
                if layer_name != "product_kpi":
                    item[layer_name]["metric_owner"]["status"] = "confirm"
            item["metric_owner"] = dict(item[item["display_kpi_layer"]]["metric_owner"])
        normalized = assessment_generator.normalize_poc_measurement_design(design, assessment)
        self.assertEqual("decision_thresholds", normalized["status"])

        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_poc_quantitative_target_page(
            ppt, width, height, normalized, "Example Service", 6,
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "p6-max-density.pptx"
            ppt.write(output_path)
            presentation = Presentation(output_path)
        slide = presentation.slides[0]
        text = self.pptx_slide_text(slide)
        for expected in (
            "製品先行KPI（提供者が直接管理）", "顧客業務成果", "提供者事業成果",
            "停止条件", "判定責任者", "要確認",
        ):
            self.assertIn(expected, text)
        self.assertNotIn("補助確認", text)
        self.assertNotIn("…", text)
        footer = next(shape for shape in slide.shapes if shape.has_text_frame and "Copyright" in shape.text)
        page_number = next(
            shape for shape in slide.shapes if shape.has_text_frame and shape.text.strip() == "6"
            and shape.top > presentation.slide_height * 0.9
        )
        self.assertTrue(all(
            shape.left >= 0 and shape.left + shape.width <= presentation.slide_width
            and shape.top + shape.height <= footer.top
            for shape in slide.shapes
            if shape.shape_id not in {footer.shape_id, page_number.shape_id}
        ))
        self.assertTrue(all(
            size is not None and size >= assessment_generator.MIN_PPTX_FONT_SIZE
            for size in self.pptx_text_font_sizes(presentation, include_footer=False)
        ))

        point_to_emu = assessment_generator.PptCanvas.POINT_TO_EMU
        card_width = round(
            ((assessment_generator.wide_content_width(width) - 12 * assessment_generator.mm) / 3)
            * point_to_emu
        )
        card_height = round(113.5 * assessment_generator.mm * point_to_emu)
        product_width = round(
            (((assessment_generator.wide_content_width(width) - 12 * assessment_generator.mm) / 3)
             - 8 * assessment_generator.mm) * point_to_emu
        )
        product_height = round(19 * assessment_generator.mm * point_to_emu)
        gate_height = round(31 * assessment_generator.mm * point_to_emu)
        cards = sorted(
            (shape for shape in slide.shapes
             if abs(shape.width - card_width) <= 2 and abs(shape.height - card_height) <= 2),
            key=lambda shape: shape.left,
        )
        product_bands = sorted(
            (shape for shape in slide.shapes
             if abs(shape.width - product_width) <= 2 and abs(shape.height - product_height) <= 2),
            key=lambda shape: shape.left,
        )
        gates = sorted(
            (shape for shape in slide.shapes
             if abs(shape.width - card_width) <= 2 and abs(shape.height - gate_height) <= 2),
            key=lambda shape: shape.left,
        )
        self.assertEqual((3, 3, 3), (len(cards), len(product_bands), len(gates)))
        for card, product_band, gate in zip(cards, product_bands, gates):
            card_text = [
                shape for shape in slide.shapes
                if shape.has_text_frame and card.left <= shape.left < card.left + card.width
                and card.top <= shape.top < card.top + card.height
                and shape.shape_id not in {footer.shape_id, page_number.shape_id}
            ]
            causal = next(shape for shape in card_text if shape.text.startswith("因果｜"))
            product_label = next(
                shape for shape in card_text if "製品先行KPI（提供者が直接管理）" in shape.text
            )
            first_gate_line = next(shape for shape in card_text if "事業効果Go" in shape.text)
            decision_line = next(shape for shape in card_text if "判定責任者" in shape.text)
            stop_line = next(shape for shape in card_text if "停止条件" in shape.text)
            self.assertLessEqual(causal.top + causal.height, product_band.top)
            self.assertGreaterEqual(product_label.top, product_band.top)
            self.assertLessEqual(product_label.top + product_label.height, product_band.top + product_band.height)
            self.assertLessEqual(product_band.top + product_band.height, gate.top)
            self.assertGreaterEqual(first_gate_line.top, gate.top)
            self.assertLessEqual(first_gate_line.top + first_gate_line.height, gate.top + gate.height)
            ordered_gate_contents = sorted(
                (first_gate_line, decision_line, stop_line), key=lambda shape: shape.top,
            )
            self.assertTrue(all(
                current.top + current.height <= following.top
                for current, following in zip(ordered_gate_contents, ordered_gate_contents[1:])
            ))

    @staticmethod
    def ai_product_business_impact_fixture(assessment: dict) -> dict:
        assessment["business_model_role"] = "provider"
        raw = {
            "lead": (
                "AI機能で得られる顧客価値を、有償化・標準化・継続利用へ転換します。"
                "売上・利益・継続収益を評価対象企業のKPIとして測定し、事業化を判断します。"
            ),
            "management_targets": [],
            "items": [
                {
                    "category": "revenue_growth", "label": "売上・ARRの拡大",
                    "headline_metric": "5〜10%", "kpi": "AI機能起因の対象サービス売上増加率",
                    "detail": "AI機能を有償提供し、既存契約への付帯と新規契約の獲得を通じて対象サービス売上を拡大する。",
                    "baseline_definition": "AI機能提供前12か月の対象サービス売上高を基準値とする。",
                    "comparison_condition": "同一顧客区分・契約期間でAI機能利用群と未利用群を比較する。",
                    "formula": "売上増加率＝AI機能起因の売上増分÷提供前売上高×100＝5〜10%",
                    "target_rationale": "顧客価値を有償付帯へ転換し、対象契約への展開率を高める。",
                    "customer_value_signal": "顧客業務KPIの達成を有償採用の中間証拠として確認する。",
                    "causal_chain": "AI機能利用 → 顧客価値 → 有償付帯 → 対象サービス売上",
                    "leading_indicator": {"kpi": "AI有償機能付帯率", "target": "15〜25%", "formula": "AI有償機能付帯率＝有償契約数÷提案契約数×100＝15〜25%"},
                    "metric_owner": {"scope": "assessed_company", "role": "サービス事業責任者", "status": "confirm"},
                    "decision_gate": {"measurement_period": "事業化後12か月", "go_condition": "対象サービス売上増加率5〜10%を達成", "stop_condition": "獲得原価が粗利上限を超過した場合", "decision_owner": "サービス事業責任者"},
                },
                {
                    "category": "profitability", "label": "粗利・営業利益の改善",
                    "headline_metric": "3〜6%", "kpi": "AI製品事業の営業利益改善率",
                    "detail": "AI売上と標準運用による提供原価低減を合算し、推論費・開発償却を含む営業利益を改善する。",
                    "baseline_definition": "AI機能提供前12か月の対象サービス営業利益を基準値とする。",
                    "comparison_condition": "同一サービス範囲でAI関連売上・原価を導入前後比較する。",
                    "formula": "営業利益改善率＝AI機能による営業利益増分÷提供前営業利益×100＝3〜6%",
                    "target_rationale": "標準機能化によって個別提供工数を抑え、追加売上を利益へ転換する。",
                    "customer_value_signal": "複数顧客で価値を再現できることを標準提供の中間証拠とする。",
                    "causal_chain": "顧客価値再現 → 標準機能化 → 提供原価低減 → 営業利益",
                    "leading_indicator": {"kpi": "AI機能の標準運用率", "target": "70〜85%", "formula": "標準運用率＝標準設定で提供できた契約数÷AI機能契約数×100＝70〜85%"},
                    "metric_owner": {"scope": "assessed_company", "role": "事業責任者・財務責任者", "status": "confirm"},
                    "decision_gate": {"measurement_period": "事業化後12か月", "go_condition": "AI製品事業の営業利益改善率3〜6%を達成", "stop_condition": "追加原価が売上増分を上回る場合", "decision_owner": "事業責任者・財務責任者"},
                },
                {
                    "category": "recurring_revenue", "label": "継続収益・NRRの向上",
                    "headline_metric": "2〜5%", "kpi": "AI利用契約の継続収益増加率",
                    "detail": "反復利用で顧客成果を継続させ、契約更新・アップセル・横展開を通じて継続収益を高める。",
                    "baseline_definition": "AI機能提供前12か月の更新・拡張・解約を含む継続収益を基準値とする。",
                    "comparison_condition": "同一契約区分でAI利用群と未利用群の更新・拡張を比較する。",
                    "formula": "継続収益増加率＝AI利用契約の継続収益増分÷提供前継続収益×100＝2〜5%",
                    "target_rationale": "反復利用と顧客成果の継続を契約更新・拡張へ接続する。",
                    "customer_value_signal": "顧客KPIの継続達成を更新・拡張の中間証拠として確認する。",
                    "causal_chain": "反復利用 → 顧客成果の継続 → 更新・拡張 → 継続収益",
                    "leading_indicator": {"kpi": "AI機能90日継続利用率", "target": "65〜80%", "formula": "90日継続利用率＝90日後も利用した契約数÷利用開始契約数×100＝65〜80%"},
                    "metric_owner": {"scope": "assessed_company", "role": "事業責任者・カスタマーサクセス責任者", "status": "confirm"},
                    "decision_gate": {"measurement_period": "事業化後12か月", "go_condition": "AI利用契約の継続収益増加率2〜5%を達成", "stop_condition": "解約またはサポート負荷が基準を超過した場合", "decision_owner": "事業責任者・カスタマーサクセス責任者"},
                },
            ],
        }
        return assessment_generator.normalize_ai_product_business_impact_candidate(raw, assessment)

    def test_ai_product_business_impact_is_company_owned_and_not_bound_to_pocs(self) -> None:
        assessment = self.detailed_assessment()
        model = self.ai_product_business_impact_fixture(assessment)
        self.assertEqual([], assessment_generator.ai_product_business_impact_contract_issues(model, assessment))
        self.assertEqual(
            ["revenue_growth", "profitability", "recurring_revenue"],
            [item["category"] for item in model["items"]],
        )
        self.assertTrue(all(item["metric_owner"]["scope"] == "assessed_company" for item in model["items"]))
        self.assertTrue(all(not any(key in item for key in ("priority", "use_case_id", "theme")) for item in model["items"]))

    def test_ai_product_business_impact_builder_serializes_category_guidance(self) -> None:
        """Prompt construction must evaluate the category dictionaries, not a set of dicts."""
        assessment = self.detailed_assessment()
        expected = self.ai_product_business_impact_fixture(assessment)
        # The normal generator uses the original percentage-only decision-
        # threshold contract.  Deterministic low/base/high calculations are not
        # part of the default generation path.
        with (
            patch.object(
                assessment_generator, "_response_json", return_value={"items": []},
            ) as mocked,
            patch.object(
                assessment_generator,
                "normalize_ai_product_business_impact_candidate_with_issues",
                return_value=(expected, []),
            ),
        ):
            result = assessment_generator.build_ai_product_business_impact(
                "会社名: Example株式会社\nサービス名: Example Service",
                assessment,
                client=object(),
                model_id="test-model",
            )

        self.assertEqual(expected, result)
        prompt = mocked.call_args.kwargs["prompt"]
        category_rows = [
            {"category": category, "guidance": guidance}
            for category, guidance in (
                ("revenue_growth", "対象サービスの売上・ARR・MRR・ARPA・有償付帯率に接続する会社KPI"),
                ("profitability", "対象サービスの粗利・営業利益・提供原価に接続する会社KPI"),
                ("recurring_revenue", "NRR・契約更新・解約・アップセル・継続収益に接続する会社KPI"),
            )
        ]
        self.assertIn(json.dumps(category_rows, ensure_ascii=False), prompt)
        self.assertIn(
            '"headline_metric": "割合または割合レンジ。NRRのみ100%超300%以下も可。検証済み中期経営目標に同じ会社KPIの金額がある場合は、その具体的金額も可"',
            prompt,
        )
        self.assertIn('"management_target_claim_id": "金額を使う場合は対応する中期経営目標のclaim_id。割合の場合は空文字"', prompt)
        self.assertIn("検証済み中期経営目標にない金額は作らない", prompt)
        self.assertIn("標準提供", prompt)
        self.assertIn("標準処理", prompt)
        self.assertIn("アクティブ率", prompt)
        self.assertIn("純収益継続率（NRR、ネット収益継続率）", prompt)
        self.assertNotIn('"status": "hypothesis_calculation"', prompt)
        self.assertNotIn('"formula_id":', prompt)
        self.assertEqual("decision_thresholds", result["status"])
        self.assertTrue(all(item["estimate_unit"] == "%" for item in result["items"]))

    def test_ai_product_business_impact_retry_uses_normalized_candidate_and_exact_issues(self) -> None:
        assessment = self.detailed_assessment()
        valid = self.ai_product_business_impact_fixture(assessment)
        invalid = json.loads(json.dumps(valid, ensure_ascii=False))
        invalid["items"][0]["leading_indicator"]["kpi"] = "AI回答速度"
        analysis_log: list[dict] = []

        with patch.object(
            assessment_generator, "_response_json", side_effect=[invalid, valid],
        ) as mocked:
            result = assessment_generator.build_ai_product_business_impact(
                "会社名: Example株式会社\nサービス名: Example Service",
                assessment,
                client=object(),
                model_id="test-model",
                analysis_log=analysis_log,
            )

        self.assertEqual(2, len(analysis_log))
        self.assertEqual("repair_required", analysis_log[0]["status"])
        self.assertTrue(any(
            "会社が直接管理する事業先行KPI" in issue
            for issue in analysis_log[0]["issues"]
        ))
        self.assertFalse(any(
            "schema_version" in issue for issue in analysis_log[0]["issues"]
        ))
        self.assertEqual("accepted", analysis_log[1]["status"])
        self.assertEqual(3, len(result["items"]))
        retry_prompt = mocked.call_args_list[1].kwargs["prompt"]
        self.assertIn("機械的に正規化した前回候補", retry_prompt)
        self.assertIn('"schema_version": "1"', retry_prompt)
        self.assertIn("会社が直接管理する事業先行KPI", retry_prompt)

    def test_ai_product_business_impact_accepts_generic_synonyms_and_nrr_above_100(self) -> None:
        """失敗ログに現れた一般的な同義語とNRRの100%超を受け付ける。"""
        assessment = self.detailed_assessment()
        raw = json.loads(json.dumps(
            self.ai_product_business_impact_fixture(assessment), ensure_ascii=False,
        ))
        profitability = raw["items"][1]
        profitability["kpi"] = "AI提供対象の運用・サポート原価改善率"
        profitability["leading_indicator"]["kpi"] = (
            "AI付帯売上に占める標準提供売上比率"
        )

        recurring = raw["items"][2]
        recurring.update({
            "headline_metric": "105～110％",
            "kpi": "AI付帯契約群のネット収益継続率",
            "formula": (
                "ネット収益継続率＝期末継続収益÷期首継続収益×100＝105〜110%"
            ),
        })
        recurring["leading_indicator"]["kpi"] = "AI付帯契約の月次アクティブ率"
        recurring["decision_gate"]["go_condition"] = (
            "AI付帯契約群のネット収益継続率105〜110%を達成"
        )

        normalized = assessment_generator.normalize_ai_product_business_impact_candidate(
            raw, assessment,
        )
        self.assertTrue(normalized)
        self.assertEqual([], assessment_generator.ai_product_business_impact_contract_issues(
            normalized, assessment,
        ))
        self.assertEqual("105〜110%", normalized["items"][2]["headline_metric"])
        self.assertEqual(105.0, normalized["items"][2]["estimate_low"])
        self.assertEqual(110.0, normalized["items"][2]["estimate_high"])

        invalid_revenue = json.loads(json.dumps(raw, ensure_ascii=False))
        invalid_revenue["items"][0]["headline_metric"] = "105〜110%"
        invalid_revenue["items"][0]["formula"] = (
            "対象サービス売上増加率＝売上増分÷基準売上×100＝105〜110%"
        )
        invalid_revenue["items"][0]["decision_gate"]["go_condition"] = (
            "対象サービス売上増加率105〜110%を達成"
        )
        self.assertEqual({}, assessment_generator.normalize_ai_product_business_impact_candidate(
            invalid_revenue, assessment,
        ))
        _, issues = assessment_generator.normalize_ai_product_business_impact_candidate_with_issues(
            invalid_revenue, assessment,
        )
        self.assertTrue(any(
            "NRRの場合は300%以下" in issue for issue in issues
        ))

    def test_ai_product_business_impact_accepts_only_verified_management_plan_amounts(self) -> None:
        assessment = self.detailed_assessment()
        model = self.ai_product_business_impact_fixture(assessment)
        management_target = {
            "claim_id": "M1-03", "claim_status": "verified_external",
            "label": "クラウドサービス 売上高", "target": "23.9億円",
            "period": "2028年6月期", "comparison": "2025年6月期比+39.1%",
            "source_id": "M1", "source_url": "https://example.com/ir/plan",
            "source_locator": "2028年6月期のクラウドサービス売上高目標",
            "source_metric": "2028年6月期 クラウドサービス 売上高 23.9億円",
            "evidence_excerpt": "クラウドサービス 売上高 23.9億円",
        }
        raw = json.loads(json.dumps(model, ensure_ascii=False))
        raw["management_targets"] = [management_target]
        raw["items"][0].update({
            "headline_metric": "23.9億円",
            "kpi": "クラウドサービス売上高",
            "formula": "中期経営目標へのAI寄与を測定し、クラウドサービス売上高23.9億円への到達を評価する。",
            "management_target_claim_id": "M1-03",
        })
        raw["items"][0]["decision_gate"]["go_condition"] = (
            "AI寄与を分離測定し、クラウドサービス売上高23.9億円への到達を評価する。"
        )

        normalized = assessment_generator.normalize_ai_product_business_impact_candidate(
            raw, assessment,
        )
        self.assertEqual("23.9億円", normalized["items"][0]["headline_metric"])
        self.assertEqual("verified_management_target", normalized["items"][0]["basis_type"])
        self.assertEqual("M1-03", normalized["items"][0]["management_target_claim_id"])
        self.assertEqual("億円", normalized["items"][0]["estimate_unit"])
        self.assertEqual(23.9, normalized["items"][0]["estimate_low"])
        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_ai_product_business_impact_page(
            ppt, width, height, normalized, assessment, assessment["service_name"], 6,
        )
        rendered_text = self.pptx_slide_text(ppt.slide)
        self.assertIn("経営計画接続｜公開目標", rendered_text)
        self.assertIn("経営目標｜23.9億円へのAI寄与確認", rendered_text)
        self.assertNotIn("PoC合格ライン｜23.9億円達成", rendered_text)

        unsupported = json.loads(json.dumps(raw, ensure_ascii=False))
        unsupported["management_targets"] = []
        self.assertEqual(
            {},
            assessment_generator.normalize_ai_product_business_impact_candidate(
                unsupported, assessment,
            ),
        )

    def test_ai_product_business_impact_rejects_customer_metrics_and_owner_laundering(self) -> None:
        assessment = self.detailed_assessment()
        model = self.ai_product_business_impact_fixture(assessment)
        tampered = json.loads(json.dumps(model, ensure_ascii=False))
        tampered["items"][0]["kpi"] = "顧客売上増加率"
        tampered["items"][0]["leading_indicator"]["kpi"] = "顧客作業時間削減率"
        tampered["items"][0]["metric_owner"] = {
            "scope": "customer", "role": "顧客業務責任者", "status": "confirm",
        }

        issues = assessment_generator.ai_product_business_impact_contract_issues(
            tampered, assessment,
        )
        self.assertTrue(any("顧客・現場側" in issue for issue in issues), issues)
        self.assertTrue(any("評価対象企業の責任者" in issue for issue in issues), issues)
        # Normalisation must preserve the invalid customer scope long enough to
        # reject it; it must never rewrite the scope to assessed_company.
        self.assertEqual(
            {},
            assessment_generator.normalize_ai_product_business_impact_candidate(
                tampered, assessment,
            ),
        )

    def test_ai_product_business_impact_audit_detects_target_tampering(self) -> None:
        assessment = self.detailed_assessment()
        model = self.ai_product_business_impact_fixture(assessment)
        items = model["items"]
        audit = {
            "schema_version": assessment_generator.AI_PRODUCT_BUSINESS_IMPACT_SCHEMA_VERSION,
            "status": "decision_thresholds", "generation_source": "llm_decision_threshold",
            "display_contract": "assessed_company_ai_business_impact", "model_id": "test-model",
            "prompt_contract_version": assessment_generator.AI_PRODUCT_BUSINESS_IMPACT_SCHEMA_VERSION,
            "estimate_count": 3, "estimate_ids": [item["estimate_id"] for item in items],
            "business_kpi_ids": [item["business_kpi_id"] for item in items],
            "categories": [item["category"] for item in items],
            "displayed_targets": [item["headline_metric"] for item in items],
            "formulas": [item["formula"] for item in items],
            "assessed_company_actuals_assumed": False, "customer_poc_metrics_reused": False,
            "reason": "会社KPIの事業化判断目標として記録する。",
        }
        research = {"ai_product_business_impact": model, "ai_product_business_impact_audit": audit}
        self.assertTrue(assessment_generator._ai_product_business_impact_is_consistent(model, research, assessment))
        tampered = json.loads(json.dumps(research, ensure_ascii=False))
        tampered["ai_product_business_impact_audit"]["displayed_targets"][0] = "99%"
        self.assertFalse(assessment_generator._ai_product_business_impact_is_consistent(model, tampered, assessment))

    def test_p6_renders_company_business_kpis_without_priority_poc_metrics(self) -> None:
        assessment = self.detailed_assessment()
        model = self.ai_product_business_impact_fixture(assessment)
        width, height = assessment_generator.PPTX_WIDESCREEN_WIDTH, assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_ai_product_business_impact_page(
            ppt, width, height, model, assessment, assessment["service_name"], 6,
        )
        text = self.pptx_slide_text(ppt.slide)
        for expected in (
            "AI導入効果目標", "計画試算", "AIで変わる判断", "算定式", "PoC合格ライン",
            "売上・ARR", "粗利・営業利益", "継続収益・NRR", "5〜10%", "3〜6%", "2〜5%",
        ):
            self.assertIn(expected, text)
        self.assertNotRegex(text, r"\bP[123]\b")
        self.assertNotIn("顧客業務成果", text)
        self.assertFalse(any(shape.has_table for shape in ppt.slide.shapes))
        for poc in assessment_generator.priority_pocs_for(assessment):
            self.assertNotIn(poc["theme"], text)
        for shape in ppt.slide.shapes:
            if (
                not shape.has_text_frame
                or self.is_ten_point_chrome_shape(
                    shape, ppt.presentation.slide_height,
                )
            ):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None:
                        self.assertGreaterEqual(run.font.size.pt, 14.0)

    def test_p5_uses_company_wide_planning_estimates_instead_of_priority_poc_metrics(self) -> None:
        assessment = self.detailed_assessment()
        assessment_generator.materialize_poc_portfolio(
            assessment, assessment["consulting_front_matter"],
        )
        poc_design = assessment_generator.normalize_poc_measurement_design(
            self.priority_measurement_design_fixture(assessment), assessment,
        )
        model = self.ai_product_business_impact_fixture(assessment)
        story = assessment_generator.assessment_story_for(assessment)
        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        ppt = assessment_generator.PptCanvas(width, height)
        assessment_generator.draw_industry_value_story_page(
            ppt, width, height, story["industry_value_story"], story["front_matter"], 5,
            poc_design, model,
        )
        text = self.pptx_slide_text(ppt.slide)
        for expected in (
            "AI導入効果目標", "計画試算", "5〜10%", "3〜6%", "2〜5%",
            "次頁：会社・製品全体のAI導入効果",
        ):
            self.assertIn(expected, text)
        self.assertNotIn("仮説｜", text)
        for poc in assessment_generator.priority_pocs_for(assessment):
            self.assertNotIn(poc["theme"], text)
        for item in poc_design["items"]:
            self.assertNotIn(item["product_kpi"]["target"], text)
            self.assertNotIn(item["product_kpi"]["kpi"], text)

    def test_p6_planning_estimate_layout_keeps_formula_and_gate_separate(self) -> None:
        assessment = self.detailed_assessment()
        base_model = self.ai_product_business_impact_fixture(assessment)
        variants = [base_model, json.loads(json.dumps(base_model, ensure_ascii=False))]
        variants[1]["management_targets"] = [
            {"label": "対象サービス売上", "target": "120億円", "delta": "+20%"},
            {"label": "営業利益率", "target": "18%", "delta": "+3pt"},
        ]
        variants[1]["items"][0]["formula"] = (
            "対象サービス売上増加率＝AI機能の有償付帯と新規契約による売上増分"
            "÷AI機能提供前12か月の対象サービス売上高×100＝5〜10%"
        )

        width = assessment_generator.PPTX_WIDESCREEN_WIDTH
        height = assessment_generator.PPTX_WIDESCREEN_HEIGHT
        for model in variants:
            with self.subTest(management_targets=bool(model["management_targets"])):
                ppt = assessment_generator.PptCanvas(width, height)
                assessment_generator.draw_ai_product_business_impact_page(
                    ppt, width, height, model, assessment, assessment["service_name"], 6,
                )
                slide = ppt.slide
                slide_text = self.pptx_slide_text(slide)
                self.assertNotIn("…", "\n".join(
                    line for line in slide_text.splitlines() if line.startswith("判定｜")
                ))
                planning_labels = [
                    shape for shape in slide.shapes
                    if shape.has_text_frame and shape.text.strip() == "AI導入効果目標｜計画試算"
                ]
                self.assertEqual(3, len(planning_labels))
                card_y = 15.5
                gate_top = round(
                    (height - (card_y + 34) * assessment_generator.mm)
                    * assessment_generator.PptCanvas.POINT_TO_EMU
                )
                gate_bottom = round(
                    (height - card_y * assessment_generator.mm)
                    * assessment_generator.PptCanvas.POINT_TO_EMU
                )
                for item in model["items"]:
                    formula_prefix = str(item["formula"]).split("＝", 1)[0]
                    formula_shape = next(
                        shape for shape in slide.shapes
                        if shape.has_text_frame and formula_prefix in shape.text and "＝" in shape.text
                    )
                    self.assertIn(str(item["headline_metric"]), formula_shape.text)
                    gate_shape = next(
                        shape for shape in slide.shapes
                        if shape.has_text_frame
                        and shape.text.startswith("PoC合格ライン｜")
                        and str(item["headline_metric"]) in shape.text
                    )
                    self.assertIn(str(item["formula"]), formula_shape.text)
                    self.assertGreaterEqual(gate_shape.top, gate_top)
                    self.assertLessEqual(
                        gate_shape.top + gate_shape.height,
                        gate_bottom,
                        "判定帯の本文が色帯の下端からはみ出しています",
                    )
                if model["management_targets"]:
                    target_label = next(
                        shape for shape in slide.shapes
                        if shape.has_text_frame and shape.text.strip() == "中期経営計画との接続"
                    )
                    self.assertLessEqual(
                        target_label.top + target_label.height,
                        min(shape.top for shape in planning_labels),
                    )
                    emu = assessment_generator.PptCanvas.POINT_TO_EMU
                    strategy_band = next(
                        shape for shape in slide.shapes
                        if shape.width >= width * emu * 0.8
                        and abs(shape.height - 11 * assessment_generator.mm * emu) <= 2
                    )
                    card_accents = [
                        shape for shape in slide.shapes
                        if shape.width < width * emu * 0.5
                        and abs(shape.height - 2 * assessment_generator.mm * emu) <= 2
                    ]
                    self.assertEqual(3, len(card_accents))
                    self.assertLessEqual(
                        strategy_band.top + strategy_band.height,
                        min(shape.top for shape in card_accents),
                        "中期経営計画との接続帯と定量効果カードが重なっています",
                    )


if __name__ == "__main__":
    unittest.main()
