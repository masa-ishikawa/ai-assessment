import copy
import unittest
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from reportlab.lib.units import mm
from ai_assess_runtime import presentation as p
from ai_assess_runtime.industry_trends import build_industry_digest, validate_industry_digest
from ai_assess_runtime.pptx_canvas import PPTX_SOURCE_NOTE_SHAPE_NAME
from test import test_generate_assessment as fixtures

class IndustryDigestTests(unittest.TestCase):
    def setUp(self):
        self.scope={'domain_label':'業務利用者の業界','rationale':'入力の顧客像と利用業務に基づく','excluded_vendor_domains':['vendor.example']}
        self.sources=[{'id':f'R{i}','url':f'https://operator{i}.example/ai','fetch_status':'fetched','excerpt':'利用者の確認業務でAIを使用し、関連文書と根拠を表示する仕組みを導入しました。'} for i in range(3)]
        self.digest={'rows':[{'source_id':f'R{i}','source_kind':'industry_ai_adoption','topic':'AIによる業務支援','published_at':'公開日未確認','development':'事業者が確認業務でAIを使用しています。','implication':'利用者の確認作業への適用を検討できます。','evidence_quote':self.sources[i]['excerpt']} for i in range(3)]}
    def test_only_grounded_non_vendor_topics_are_frozen(self):
        self.assertEqual([],validate_industry_digest(self.digest,self.sources,self.scope))
        for mutation in ['vendor','quote','market','duplicate','missing']:
            with self.subTest(mutation=mutation):
                sources=copy.deepcopy(self.sources);d=copy.deepcopy(self.digest)
                if mutation=='vendor':sources[0]['url']='https://corp.vendor.example/ir/plan'
                if mutation=='quote':d['rows'][0]['evidence_quote']='原典には存在しない架空のAI活用の説明です。'
                if mutation=='market':d['rows'][0]['source_kind']='market_forecast'
                if mutation=='duplicate':d['rows'][1]['source_id']='R0'
                if mutation=='missing':sources[0]['fetch_status']='fetch_failed'
                self.assertTrue(validate_industry_digest(d,sources,self.scope))
    def test_digest_keeps_input_scope_urls_and_review_decisions(self):
        audit=[];d=build_industry_digest('入力',{},self.scope,self.sources,respond=lambda _:self.digest,audit_log=audit)
        self.assertEqual(self.sources[0]['url'],d['rows'][0]['url'])
        self.assertEqual(self.scope,d['industry_scope'])
        self.assertEqual('accepted',audit[-1]['status'])
    def test_legacy_vendor_sources_never_become_industry_news(self):
        d=p.industry_trend_digest_for({'industry_sources':[{'id':'R1','fetch_status':'fetched','url':'https://vendor.example/ir','title':'対象製品の中計2034','snippet':'自社製品の紹介'}]}, {'company_name':'対象企業','service_name':'対象製品'})
        self.assertNotIn('自社製品',str(d));self.assertTrue(all(not r['url'] for r in d['rows']))

class LayoutRevisionTests(unittest.TestCase):
    def canvas(self):return p.PptCanvas(p.PPTX_WIDESCREEN_WIDTH,p.PPTX_WIDESCREEN_HEIGHT)
    def test_intro_subtitle_is_smaller_gray_and_below_title(self):
        c=self.canvas();p.draw_default_assessment_intro_page(c,c.page_width,c.page_height,fixtures.DynamicPptxPageTests.detailed_assessment())
        title=next(s for s in c.slide.shapes if s.has_text_frame and s.text=='OCI AI Use Case Assessment')
        sub=next(s for s in c.slide.shapes if s.has_text_frame and s.text=='AIユースケースアセスメント')
        r=sub.text_frame.paragraphs[0].runs[0]
        self.assertEqual(14,r.font.size.pt);self.assertEqual('777777',str(r.font.color.rgb));self.assertGreater(sub.top,title.top+title.height)
    def test_selection_numbers_share_circle_geometry(self):
        c=self.canvas();p.draw_default_poc_selection_page(c,c.page_width,c.page_height,fixtures.DynamicPptxPageTests.detailed_assessment(),7)
        numbers=[s for s in c.slide.shapes if s.has_text_frame and s.text in ['1','2','3']]
        self.assertEqual(3,len(numbers))
        for number in numbers:
            matching=[s for s in c.slide.shapes if s!=number and (s.left,s.top,s.width,s.height)==(number.left,number.top,number.width,number.height)]
            self.assertTrue(matching);self.assertEqual(MSO_ANCHOR.MIDDLE,number.text_frame.vertical_anchor)
            self.assertEqual(PP_ALIGN.CENTER,number.text_frame.paragraphs[0].alignment)
    def test_cost_table_notes_and_totals_have_separate_aligned_regions(self):
        c=self.canvas();p.draw_cost_estimate_page(c,c.page_width,c.page_height,fixtures.assessment_generator.cost_estimate_from_snapshot({"schema_version":"1", **copy.deepcopy(fixtures.TEST_COST_ESTIMATE)}),12)
        table=next(s for s in c.slide.shapes if s.has_table)
        bands=[s for s in c.slide.shapes if s.name.startswith('AI_ASSESS_COST_NOTE_BAND')]
        self.assertEqual(3,len(bands));self.assertLess(table.top+table.height,min(s.top for s in bands))
        for left,right in zip(bands,bands[1:]):self.assertLess(left.top+left.height,right.top)
        totals=[s for s in c.slide.shapes if s.name.startswith('AI_ASSESS_COST_TOTAL_')]
        self.assertEqual(1,len({(s.top,s.height) for s in totals}))
        for s in totals:self.assertEqual(MSO_ANCHOR.MIDDLE,s.text_frame.vertical_anchor)
    def test_midterm_source_matches_industry_source_role(self):
        c=self.canvas();p.draw_midterm_plan_alignment_page(c,c.page_width,c.page_height,{'source':{'title':'中計','url':'https://example.com/plan.pdf'},'ai_alignment':[]},3)
        source=next(s for s in c.slide.shapes if s.name==PPTX_SOURCE_NOTE_SHAPE_NAME)
        self.assertIn('https://example.com/plan.pdf',source.text)
        self.assertEqual(9,source.text_frame.paragraphs[0].runs[0].font.size.pt)
        self.assertLess(source.top+source.height,c.presentation.slide_height-10*mm*12700)

if __name__=='__main__':unittest.main()
