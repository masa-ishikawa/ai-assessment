"""TXT/XLSXから凍結レビューJSONまでの通常CLI経路を固定する統合テスト。"""

import contextlib
import copy
from contextlib import ExitStack
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

from test import create_isv_test_input
import generate_assessment as generator
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from test import test_generate_assessment as fixtures


class GenerationWorkflowIntegrationTests(unittest.TestCase):
    """外部通信をモックし、main -> run_generation の通常経路を検証する。"""

    @staticmethod
    def _analysis_fixture() -> dict:
        assessment = copy.deepcopy(fixtures.DynamicPptxPageTests.detailed_assessment())
        # 入力による正規化を検証するため、LLM応答側は意図的に別名にする。
        assessment["company_name"] = "LLM誤認株式会社"
        assessment["service_name"] = "LLM誤認サービス"
        # 優先候補は15件のカタログへ明示的に結び、統合テスト自体の
        # fixture不整合によってポートフォリオが縮退しないようにする。
        for index, recommendation in enumerate(assessment["poc_recommendations"], 1):
            recommendation["theme"] = f"ユースケース{index}"
        # 通常分析応答へ旧数値フィールドが混入しても、単一定量契約の前で
        # 必ず除去されることを統合テストで固定する。
        assessment["ai_product_business_impact"] = (
            fixtures.DynamicPptxPageTests.ai_product_business_impact_fixture(assessment)
        )
        return assessment

    def _run_json_only_case(
        self, *, input_file: Path, expected_company: str, expected_service: str,
        expects_preprocessing: bool, workspace: Path,
        midterm_plan: dict | None = None,
        expected_analysis_markers: tuple[str, ...] = (),
    ) -> None:
        output_json = workspace / "json" / f"{input_file.stem}.json"
        research_dir = workspace / "research"
        pptx_dir = workspace / "pptx"
        create_pptx = Mock(side_effect=AssertionError("--json-onlyでPPTXを生成してはいけません"))
        write_manifest = Mock(side_effect=AssertionError("--json-onlyでmanifestを生成してはいけません"))
        build_quantitative_hypothesis = Mock(side_effect=AssertionError(
            "通常生成で未検証の定量仮説を生成してはいけません",
        ))
        build_ai_product_business_impact = Mock(side_effect=(
            lambda source_text, assessment, **kwargs:
            fixtures.DynamicPptxPageTests.ai_product_business_impact_fixture(assessment)
        ))

        source_text = generator.load_source_text(input_file)
        expected_source_text_sha = hashlib.sha256(source_text.encode("utf-8")).hexdigest()
        expected_source_file_sha = hashlib.sha256(input_file.read_bytes()).hexdigest()
        analysis_contexts: list[str] = []

        def capture_context(context: str) -> None:
            analysis_contexts.append(context)

        def analyze(context: str, *args, **kwargs) -> dict:
            capture_context(context)
            return self._analysis_fixture()

        def collect_research(context: str, *args, **kwargs) -> list:
            capture_context(context)
            return []

        def build_midterm_analysis(context: str, *args, **kwargs) -> dict:
            capture_context(context)
            return {}

        def build_front_matter(context: str, assessment: dict, **kwargs) -> dict:
            capture_context(context)
            return generator.fallback_consulting_front_matter(assessment)

        def select_sources(context: str, *args, **kwargs) -> dict:
            capture_context(context)
            return {"approved_source_ids": []}

        def build_decision(context: str, *args, **kwargs) -> dict:
            capture_context(context)
            return {}

        with ExitStack() as patches:
            patches.enter_context(patch.object(generator, "RESEARCH_OUTPUT_DIR", research_dir))
            patches.enter_context(patch.object(generator, "JSON_OUTPUT_DIR", workspace / "json-default"))
            patches.enter_context(patch.object(generator, "PPTX_OUTPUT_DIR", pptx_dir))
            patches.enter_context(patch.object(generator, "create_oci_responses_client", return_value=object()))
            patches.enter_context(patch.object(
                generator, "resolve_official_company_name",
                side_effect=lambda company_name, **kwargs: company_name,
            ))
            patches.enter_context(patch.object(
                generator, "collect_midterm_plan",
                return_value=midterm_plan or {
                    "status": "not_found", "reason": "統合テストでは外部調査を行わない",
                },
            ))
            patches.enter_context(patch.object(
                generator, "collect_industry_research", side_effect=collect_research,
            ))
            patches.enter_context(patch.object(
                generator, "build_midterm_plan_analysis", side_effect=build_midterm_analysis,
            ))
            patches.enter_context(patch.object(
                generator, "select_verified_quantitative_sources",
                side_effect=select_sources,
            ))
            patches.enter_context(patch.object(
                generator, "analyze_assessment", side_effect=analyze,
            ))
            patches.enter_context(patch.object(
                generator, "build_consulting_front_matter",
                side_effect=build_front_matter,
            ))
            patches.enter_context(patch.object(
                generator, "build_final_poc_logic_details", return_value=[],
            ))
            patches.enter_context(patch.object(
                generator, "build_quantitative_hypothesis",
                build_quantitative_hypothesis,
            ))
            patches.enter_context(patch.object(
                generator, "build_ai_product_business_impact",
                build_ai_product_business_impact,
            ))
            patches.enter_context(patch.object(
                generator, "build_poc_decision_data", side_effect=build_decision,
            ))
            patches.enter_context(patch.object(
                generator, "build_poc_cost_estimate",
                return_value=copy.deepcopy(fixtures.TEST_COST_ESTIMATE),
            ))
            patches.enter_context(patch.object(generator, "create_pptx", create_pptx))
            patches.enter_context(patch.object(generator, "write_render_manifest", write_manifest))
            patches.enter_context(patch.object(sys, "argv", [
                "generate_assessment.py", str(input_file),
                "--provider", "oci_responses",
                "--model-id", "test-model",
                "--oci-project-ocid", "ocid1.generativeaiproject.oc1..integrationtest",
                "--json-only", "--output-json", str(output_json),
            ]))
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                return_code = generator.main()

        self.assertEqual(0, return_code)
        self.assertTrue(output_json.is_file())
        self.assertFalse(pptx_dir.exists())
        self.assertEqual([], list(workspace.rglob("*.pptx")))
        create_pptx.assert_not_called()
        write_manifest.assert_not_called()
        build_quantitative_hypothesis.assert_not_called()
        build_ai_product_business_impact.assert_called_once()

        payload = json.loads(output_json.read_text(encoding="utf-8"))
        self.assertEqual(generator.ASSESSMENT_JSON_FROZEN_FORMAT, payload["format"])
        self.assertEqual([], generator.validate_assessment_payload(payload, strict=True))
        self.assertEqual(expected_company, payload["assessment"]["company_name"])
        self.assertEqual(expected_service, payload["assessment"]["service_name"])
        self.assertNotEqual("LLM誤認株式会社", payload["assessment"]["company_name"])
        self.assertNotEqual("LLM誤認サービス", payload["assessment"]["service_name"])
        self.assertEqual(15, len(payload["assessment"]["use_cases"]))
        self.assertEqual(3, len(payload["assessment"]["poc_portfolio"]["items"]))
        self.assertEqual(
            3,
            len(payload["assessment"]["ai_product_business_impact"]["items"]),
        )
        self.assertEqual(
            "llm_estimate",
            payload["research"]["ai_product_business_impact_audit"]["evidence_mode"],
        )
        display_contract = payload["research"]["quantitative_display_contract"]
        self.assertEqual("measurement_design_only", display_contract["display_mode"])
        self.assertFalse(display_contract["numeric_claims_allowed"])
        self.assertEqual(
            display_contract["measurement_design"],
            payload["assessment"]["poc_measurement_design"],
        )
        self.assertEqual(
            display_contract["display_mode"],
            payload["research"]["quantitative_analysis"]["final_display_mode"],
        )

        input_record = payload["input"]
        self.assertEqual(str(input_file), input_record["source_file"])
        self.assertEqual(expected_source_text_sha, input_record["source_text_sha256"])
        self.assertEqual(expected_source_file_sha, input_record["source_file_sha256"])
        for marker in expected_analysis_markers:
            self.assertTrue(analysis_contexts, "分析用文脈が後段へ渡されていません")
            self.assertTrue(
                all(marker in context for context in analysis_contexts),
                f"分析用文脈に公開情報がありません: {marker}",
            )
        if expects_preprocessing:
            self.assertIsInstance(input_record["preprocessing"], dict)
            self.assertTrue(input_record["preprocessing"])
        else:
            self.assertIsNone(input_record["preprocessing"])

        research_path = research_dir / f"{generator.safe_filename(expected_service)}_research.json"
        self.assertTrue(research_path.is_file())
        self.assertEqual(
            payload["research"], json.loads(research_path.read_text(encoding="utf-8")),
        )

    def test_txt_and_xlsx_complete_the_normal_json_only_cli_workflow(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            workspace = Path(temporary_directory)
            txt_input = workspace / "source.txt"
            txt_input.write_text(
                "会社名: TXT入力株式会社\n"
                "サービス名: TXT入力サービス\n"
                "サービス概要: 業務データを用いて利用者の判断を支援するクラウドサービス。\n",
                encoding="utf-8",
            )

            xlsx_input = workspace / "source.xlsx"
            answers = create_isv_test_input.build_required_answers(
                "XLSX入力株式会社", "XLSX入力サービス", [],
            )
            create_isv_test_input.fill_template(
                generator.INPUT_DIR / "ISV_AI_Use_Case_Assessment_Input_Template.xlsx",
                xlsx_input,
                answers,
            )

            cases = (
                (txt_input, "TXT入力株式会社", "TXT入力サービス", False),
                (xlsx_input, "XLSX入力株式会社", "XLSX入力サービス", True),
            )
            for input_file, company, service, expects_preprocessing in cases:
                with self.subTest(input_format=input_file.suffix):
                    self._run_json_only_case(
                        input_file=input_file,
                        expected_company=company,
                        expected_service=service,
                        expects_preprocessing=expects_preprocessing,
                        workspace=workspace / input_file.stem,
                    )

    def test_midterm_plan_enriches_analysis_but_not_source_text_hash(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            workspace = Path(temporary_directory)
            source = workspace / "source.txt"
            source.write_text(
                "会社名: Context分離株式会社\n"
                "サービス名: Context分離クラウド\n"
                "サービス概要: 業務判断を支援するサービス。\n",
                encoding="utf-8",
            )
            self._run_json_only_case(
                input_file=source,
                expected_company="Context分離株式会社",
                expected_service="Context分離クラウド",
                expects_preprocessing=False,
                workspace=workspace / "result",
                midterm_plan={
                    "status": "found",
                    "title": "公開中期経営計画2028",
                    "url": "https://example.com/ir/plan",
                    "excerpt": "公開計画固有マーカー：継続収益と品質を向上する。",
                },
                expected_analysis_markers=(
                    "公開中期経営計画2028",
                    "公開計画固有マーカー",
                ),
            )

    @staticmethod
    def _text_frames_for_shape(shape):
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame

    @classmethod
    def _pptx_semantic_signature(cls, path: Path) -> dict:
        """Package metadataを除き、表示内容と配置だけを決定的に比較する。"""
        presentation = Presentation(path)
        slides: list[list[dict]] = []
        for slide in presentation.slides:
            slide_shapes: list[dict] = []
            for shape in slide.shapes:
                shape_record: dict[str, object] = {
                    "type": int(shape.shape_type),
                    "box": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)],
                }
                frames = list(cls._text_frames_for_shape(shape))
                if frames:
                    shape_record["paragraphs"] = [
                        [
                            {
                                "text": run.text,
                                "font_name": run.font.name,
                                "font_size": run.font.size.pt if run.font.size is not None else None,
                                "bold": run.font.bold,
                            }
                            for run in paragraph.runs
                        ]
                        for frame in frames
                        for paragraph in frame.paragraphs
                    ]
                if shape.has_table:
                    shape_record["table"] = [
                        [cell.text for cell in row.cells]
                        for row in shape.table.rows
                    ]
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    shape_record["image_sha256"] = hashlib.sha256(shape.image.blob).hexdigest()
                slide_shapes.append(shape_record)
            slides.append(slide_shapes)
        return {
            "slide_width": int(presentation.slide_width),
            "slide_height": int(presentation.slide_height),
            "slides": slides,
        }

    @classmethod
    def _assert_pptx_contract(cls, test_case: unittest.TestCase, path: Path) -> Presentation:
        """16:9、12pt、スライド領域内、破損なしを一括検証する。"""
        with zipfile.ZipFile(path) as package:
            test_case.assertIsNone(package.testzip(), f"PPTX package is corrupt: {path}")
        presentation = Presentation(path)
        test_case.assertEqual(12.0, generator.MIN_PPTX_FONT_SIZE)
        test_case.assertEqual(10.0, generator.PPTX_FOOTER_FONT_SIZE)
        test_case.assertEqual(
            9 * presentation.slide_width,
            16 * presentation.slide_height,
            "PPTX must be 16:9",
        )
        for slide_index, slide in enumerate(presentation.slides, 1):
            for shape in slide.shapes:
                # Full-bleed photos may intentionally extend beyond the canvas so
                # PowerPoint crops them without distorting the image.  They must
                # still intersect the slide; editable content must stay fully in-bounds.
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    test_case.assertLess(shape.left, presentation.slide_width)
                    test_case.assertLess(shape.top, presentation.slide_height)
                    test_case.assertGreater(shape.left + shape.width, 0)
                    test_case.assertGreater(shape.top + shape.height, 0)
                    continue
                test_case.assertGreaterEqual(
                    shape.left, 0, f"slide {slide_index}: shape starts outside the canvas",
                )
                test_case.assertGreaterEqual(
                    shape.top, 0,
                    f"slide {slide_index}: shape {shape.name!r}/{shape.shape_type} starts outside the canvas",
                )
                test_case.assertLessEqual(
                    shape.left + shape.width,
                    presentation.slide_width,
                    f"slide {slide_index}: shape exceeds the right edge",
                )
                test_case.assertLessEqual(
                    shape.top + shape.height,
                    presentation.slide_height,
                    f"slide {slide_index}: shape exceeds the bottom edge",
                )
                shape_text = shape.text.strip() if shape.has_text_frame else ""
                is_ten_point_chrome = (
                    (shape_text == "OCI AI USE CASE ASSESSMENT"
                     and shape.top < presentation.slide_height * 0.12)
                    or (shape_text.startswith("Copyright ©")
                        and shape.top > presentation.slide_height * 0.85)
                    or (shape_text.isdigit()
                        and shape.top > presentation.slide_height * 0.9)
                )
                is_nine_point_source = getattr(shape, "name", "") == "AI_ASSESS_SOURCE_NOTE"
                for frame in cls._text_frames_for_shape(shape):
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            if not run.text.strip():
                                continue
                            test_case.assertIsNotNone(
                                run.font.size,
                                f"slide {slide_index}: editable text has no explicit font size: {run.text!r}",
                            )
                            if is_ten_point_chrome:
                                test_case.assertEqual(
                                    run.font.size.pt, 10.0,
                                    f"slide {slide_index}: common header/footer must be 10pt: {run.text!r}",
                                )
                            elif is_nine_point_source:
                                test_case.assertEqual(
                                    run.font.size.pt, 9.0,
                                    f"slide {slide_index}: industry source note must be 9pt: {run.text!r}",
                                )
                            else:
                                test_case.assertGreaterEqual(
                                    run.font.size.pt, 12.0,
                                    f"slide {slide_index}: editable content below 12pt: {run.text!r}",
                                )
        return presentation

    @staticmethod
    def _assert_manifest_binding(
        test_case: unittest.TestCase, manifest_path: Path, payload: dict,
    ) -> None:
        test_case.assertEqual([], generator.validate_render_manifest(manifest_path))
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        provenance = payload["provenance"]
        test_case.assertEqual(
            {
                "assessment_run_id": provenance["assessment_run_id"],
                "input_record_sha256": provenance["input_record_sha256"],
                "assessment_sha256": provenance["assessment_sha256"],
                "research_content_sha256": provenance["research_content_sha256"],
                "cost_estimate_sha256": provenance["cost_estimate_sha256"],
                "rendering_sha256": provenance["rendering_sha256"],
            },
            manifest["review"],
        )
        test_case.assertRegex(
            manifest["renderer"]["current_generator_sha256"], r"^[0-9a-f]{64}$",
        )
        test_case.assertEqual(
            payload["rendering"]["architecture_asset"],
            manifest["renderer"]["architecture_asset"],
        )

    def test_normal_generation_and_frozen_json_replay_are_semantically_identical(self) -> None:
        """通常CLIの承認対象と--from-jsonの再描画結果を一本のgoldenで固定する。"""
        real_create_pptx = generator.create_pptx
        rendered_assessments: list[dict] = []

        def create_pptx_without_mutation(assessment, architecture, output, cost, research=None) -> None:
            before = copy.deepcopy(assessment)
            rendered_assessments.append(before)
            real_create_pptx(assessment, architecture, output, cost, research)
            self.assertEqual(before, assessment, "PPTX renderer mutated the frozen assessment")

        with tempfile.TemporaryDirectory() as temporary_directory:
            workspace = Path(temporary_directory)
            source = workspace / "generic_service.txt"
            source.write_text(
                "会社名: Example Software株式会社\n"
                "サービス名: Example Operations Cloud\n"
                "サービス概要: 業務イベントと文書を蓄積し、利用者の判断を支援するクラウドサービス。\n",
                encoding="utf-8",
            )
            json_dir = workspace / "json"
            research_dir = workspace / "research"
            normal_pptx = workspace / "normal.pptx"
            replay_pptx = workspace / "replay.pptx"
            build_quantitative_hypothesis = Mock(side_effect=AssertionError(
                "通常生成で未検証の定量仮説を生成してはいけません",
            ))
            build_ai_product_business_impact = Mock(side_effect=(
                lambda source_text, assessment, **kwargs:
                fixtures.DynamicPptxPageTests.ai_product_business_impact_fixture(assessment)
            ))
            midterm_plan = {
                "status": "found",
                "title": "公開中期経営計画2028",
                "url": "https://example.com/investors/management-plan/2028",
                "excerpt": "継続収益と顧客価値を高める重点方針を推進する。",
            }
            midterm_analysis = {
                "plan_summary": "中期計画の重点方針を、AIの事業価値と実行条件へ接続する。",
                "source": {
                    "title": midterm_plan["title"], "url": midterm_plan["url"],
                },
                "management_targets": [{
                    "label": "継続収益", "target": "120億円",
                    "period": "FY2028", "comparison": "FY2025比+20%",
                }],
                "ai_alignment": [
                    {
                        "plan_priority": f"重点方針{index}",
                        "ai_role": f"AI支援{index}",
                        "why_now": f"優先テーマ{index}の実行条件を確認する。",
                        "related_use_case": f"関連AIテーマ{index}",
                    }
                    for index in range(1, 4)
                ],
                "caveat": "事業効果は同一条件のPoCで測定し、投資判断へ接続する。",
            }

            with ExitStack() as patches:
                patches.enter_context(patch.object(generator, "RESEARCH_OUTPUT_DIR", research_dir))
                patches.enter_context(patch.object(generator, "JSON_OUTPUT_DIR", json_dir))
                patches.enter_context(patch.object(generator, "PPTX_OUTPUT_DIR", workspace / "pptx"))
                patches.enter_context(patch.object(generator, "create_oci_responses_client", return_value=object()))
                patches.enter_context(patch.object(
                    generator, "resolve_official_company_name",
                    side_effect=lambda company_name, **kwargs: company_name,
                ))
                patches.enter_context(patch.object(
                    generator, "collect_midterm_plan",
                    return_value=midterm_plan,
                ))
                patches.enter_context(patch.object(generator, "collect_industry_research", return_value=[]))
                patches.enter_context(patch.object(
                    generator, "build_midterm_plan_analysis",
                    return_value=midterm_analysis,
                ))
                patches.enter_context(patch.object(
                    generator, "select_verified_quantitative_sources",
                    return_value={"approved_source_ids": []},
                ))
                patches.enter_context(patch.object(
                    generator, "analyze_assessment", return_value=self._analysis_fixture(),
                ))
                patches.enter_context(patch.object(
                    generator, "build_consulting_front_matter",
                    side_effect=lambda source_text, assessment, **kwargs:
                        generator.fallback_consulting_front_matter(assessment),
                ))
                patches.enter_context(patch.object(
                    generator, "build_final_poc_logic_details", return_value=[],
                ))
                patches.enter_context(patch.object(
                    generator, "build_quantitative_hypothesis",
                    build_quantitative_hypothesis,
                ))
                patches.enter_context(patch.object(
                    generator, "build_ai_product_business_impact",
                    build_ai_product_business_impact,
                ))
                patches.enter_context(patch.object(generator, "build_poc_decision_data", return_value={}))
                patches.enter_context(patch.object(
                    generator, "build_poc_cost_estimate",
                    return_value=generator.cost_estimate_from_snapshot(
                        generator.cost_estimate_snapshot(
                            copy.deepcopy(fixtures.TEST_COST_ESTIMATE),
                        )
                    ),
                ))
                patches.enter_context(patch.object(
                    generator, "create_pptx", side_effect=create_pptx_without_mutation,
                ))
                patches.enter_context(patch.object(sys, "argv", [
                    "generate_assessment.py", str(source),
                    "--provider", "oci_responses",
                    "--model-id", "test-model",
                    "--oci-project-ocid", "ocid1.generativeaiproject.oc1..goldentest",
                    "--output", str(normal_pptx),
                ]))
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    self.assertEqual(0, generator.main())

            build_quantitative_hypothesis.assert_not_called()
            build_ai_product_business_impact.assert_called_once()

            service_stem = generator.safe_filename("Example Operations Cloud")
            frozen_json = json_dir / f"{service_stem}_assessment.json"
            research_json = research_dir / f"{service_stem}_research.json"
            self.assertTrue(frozen_json.is_file())
            self.assertTrue(research_json.is_file())
            frozen_json_sha = hashlib.sha256(frozen_json.read_bytes()).hexdigest()
            payload = json.loads(frozen_json.read_text(encoding="utf-8"))
            self.assertEqual([], generator.validate_assessment_payload(payload, strict=True))
            self.assertEqual(payload["assessment"], rendered_assessments[0])
            self.assertEqual(payload["research"], json.loads(research_json.read_text(encoding="utf-8")))
            self.assertEqual(
                3,
                len(payload["assessment"]["ai_product_business_impact"]["items"]),
            )
            self.assertEqual(
                "planning_estimate",
                payload["research"]["ai_product_business_impact_audit"]["estimate_classification"],
            )
            display_contract = payload["research"]["quantitative_display_contract"]
            self.assertEqual("measurement_design_only", display_contract["display_mode"])
            self.assertFalse(display_contract["numeric_claims_allowed"])
            self.assertEqual(
                display_contract["measurement_design"],
                payload["assessment"]["poc_measurement_design"],
            )
            self.assertEqual(
                display_contract["display_mode"],
                payload["research"]["quantitative_analysis"]["final_display_mode"],
            )

            normal_manifest = normal_pptx.with_suffix(".pptx.manifest.json")
            self._assert_manifest_binding(self, normal_manifest, payload)
            normal_presentation = self._assert_pptx_contract(self, normal_pptx)
            expected_plan = generator.default_document_slide_plan(
                payload["assessment"], include_cost_estimate=True,
            )
            self.assertEqual(14, len(expected_plan))
            self.assertEqual(len(expected_plan), len(normal_presentation.slides))
            self.assertEqual(
                ["cover", "intro", "midterm", "industry_value", "business_impact"],
                [spec.key for spec in expected_plan[:5]],
            )

            with ExitStack() as patches:
                patches.enter_context(patch.object(generator, "RESEARCH_OUTPUT_DIR", workspace / "replay-research"))
                patches.enter_context(patch.object(generator, "JSON_OUTPUT_DIR", workspace / "replay-json"))
                patches.enter_context(patch.object(generator, "PPTX_OUTPUT_DIR", workspace / "replay-pptx"))
                patches.enter_context(patch.object(
                    generator, "create_pptx", side_effect=create_pptx_without_mutation,
                ))
                patches.enter_context(patch.object(sys, "argv", [
                    "generate_assessment.py", "--from-json", str(frozen_json),
                    "--strict-json", "--output", str(replay_pptx),
                ]))
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    self.assertEqual(0, generator.main())

            self.assertEqual(frozen_json_sha, hashlib.sha256(frozen_json.read_bytes()).hexdigest())
            self.assertEqual(2, len(rendered_assessments))
            self.assertEqual(payload["assessment"], rendered_assessments[1])
            replay_manifest = replay_pptx.with_suffix(".pptx.manifest.json")
            self._assert_manifest_binding(self, replay_manifest, payload)
            replay_presentation = self._assert_pptx_contract(self, replay_pptx)
            self.assertEqual(14, len(replay_presentation.slides))
            self.assertEqual(len(expected_plan), len(replay_presentation.slides))
            self.assertEqual(
                self._pptx_semantic_signature(normal_pptx),
                self._pptx_semantic_signature(replay_pptx),
                "normal generation and frozen JSON replay diverged",
            )


if __name__ == "__main__":
    unittest.main()
